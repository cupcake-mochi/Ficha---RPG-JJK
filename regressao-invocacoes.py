# -*- coding: utf-8 -*-
"""Regressao da aba INVOCAÇÕES: preenche fichas na planilha GERADA, manda o LibreOffice recalcular e compara com o livro
e com a regra, escrita aqui de novo, em Python.

Quem esta sendo julgada e a planilha: as formulas da aba e as contas dela na DADOS_INVOC. Nada daqui le formula nem usa
tabela que o gerador montou; do ficha_invocacoes.py so vem o endereco de cada caixa.

  1. Os exemplos do livro, com o numero que o LIVRO imprime escrito aqui: o Cao de sombra (niveis 2, 4, 5 e 6), o Vigia
     de papel, o Fura de Classe 2, a Onda de Classe 1, a domada de nivel 10, o talisma de nivel 13, o corpo de criacao,
     a Liberacao de Classe 3, a Tecnica Maxima de nivel 17 e a Expansao de nivel 14.
  2. O Buff/Debuff: cada caixinha muda o numero de cima, e so ele.
  3. O que a ficha recusa: Familia Fechada, Forma sem a Familia dela, orcamento estourado, lugar que ainda nao abriu,
     basica com peca que nao e Leve, atributo acima de 3 sem marco, Familia repetida e Familia a mais.
  4. Fichas sorteadas, em varios niveis, tipos e Familias: os numeros da ficha e o resultado de toda carta batem com a
     regra daqui, a certa e a errada (o sorteio monta habilidade ilegal de proposito).

07/10/2026, a grade de 2 x 6: a aba tem doze lugares. As fichas de teste que dividem o mesmo invocador (nivel, Essencia,
Inteligencia, o atributo da Defesa e a Trilha) vao na mesma planilha, cada uma num lugar, e as sorteadas passam pelos
doze. Os enderecos daqui sao os do primeiro lugar; `desloca` os leva ao lugar de cada ficha. A planilha gerada e lida
uma vez so, e cada planilha recalculada e lida sem carregar as outras abas: com as doze fichas cada leitura inteira
leva quatorze segundos, e eram duas por caso.

    python3 regressao-invocacoes.py
"""
import json, math, os, random, re, shutil, subprocess, sys, tempfile
from openpyxl import load_workbook

ARQ = "ficha-v01/ficha-projeto-m-0.1.xlsx"
if not os.path.exists(ARQ):
    print("gere a ficha antes:  python3 ficha-v01/monta.py"); sys.exit(1)
sys.path.insert(0, "ficha-v01")
import indice_ficha as ix, ficha_invocacoes as fi

CAT = json.load(open("catalogo-projeto-m.json", encoding="utf-8"))
LIV = json.load(open("ficha-v01/invocacao-do-livro.json", encoding="utf-8"))
DEC = json.load(open("decisoes-ficha.json", encoding="utf-8"))
ABA, DIV = fi.NOME, fi.DADOS_IV
FALHAS = []


def checa(desc, cond, det=""):
    print(f"  [{'OK' if cond else 'FALHA'}] {desc}" + ("" if cond else f"  <- {det}"))
    if not cond:
        FALHAS.append(desc)


_r = subprocess.run([sys.executable, "ficha-v01/extrair_invocacao.py", "--confere"], capture_output=True, text=True)
print("O LIVRO")
checa("o invocacao-do-livro.json e o que o manual.txt diz hoje", _r.returncode == 0, (_r.stdout + _r.stderr).strip()[-200:])

# ---------------------------------------------------------------------------------------------
# A regra, escrita de novo, do capitulo 17 (Construir invocacoes) e do 16 (Invocacoes em campo).
# ---------------------------------------------------------------------------------------------
MARCOS = [6, 10, 14, 18, 22, 26, 30]
ATR = ["Força", "Destreza", "Constituição", "Inteligência", "Essência"]
maestria = lambda n: 1 + sum(1 for x in (10, 18, 26) if x <= n)                 # Nivel | Maestria: 2 a 9 +1, 10 a 17 +2...
classe_maxima = lambda n: sum(1 for x in (1, 5, 9, 13, 17, 21, 26) if x <= n)   # Progressao da entidade
dados_da_basica = lambda n: 1 if n <= 4 else 2 if n <= 16 else 3
PONTOS = {1: 2, 2: 4, 3: 6, 4: 9, 5: 11, 6: 13, 7: 16}
LIMITE = {1: 2, 2: 2, 3: 3, 4: 3, 5: 4, 6: 4, 7: 4}
PESO = {"Leve": 1, "Media": 2, "Média": 2, "Pesada": 3}
preco = lambda peso, c: math.ceil(c / 2) if peso == 1 else c if peso == 2 else math.ceil(c * 1.5)
preco_livre = lambda peso, c: max(1, preco(peso, c) - max(1, c // 2))           # metade da Classe para baixo, minimo 1
PECAS = {}
for _n, _m in CAT["melhorias"].items():
    if _n == "Condição":
        for _c, _p in CAT["condicoes"].items():
            PECAS[_c] = {"familia": "Controle", "peso": PESO[_p]}
    elif _n == "Efeito Próprio":
        for _p in ("Leve", "Média", "Pesada"):
            PECAS[f"Efeito Próprio ({_p})"] = {"familia": None, "peso": PESO[_p]}
    else:
        PECAS[_n] = {"familia": _m["familia"], "peso": PESO[_m["peso"]]}
RESTRICOES = {}
for _n, _r_ in CAT["restricoes"].items():
    _nv = _r_["devolve"].split(" ou ")
    for _x in _nv:
        RESTRICOES[_n if len(_nv) == 1 else f"{_n} ({'Média' if _x == 'Media' else _x})"] = {"nome": _n, "devolve": PESO[_x]}
FREQUENCIA = ("Uma Vez", "Condicional", "Aquecer", "Dívida")
PARES = [(p["a"], p["b"]) for p in DEC["A3_incompatibilidades"]["pares"]]
FORMAS = {f["nome"]: {"peso": PESO.get(f["custa"], 0), "exige": f["exige"], "embutida": f["embutida"],
                      "tipo": {"Cura": 1, "Onda": 1, "Apoio": 2, "Efeito": 3}.get(f["nome"], 0),
                      "tr": f["nome"] in ("Explosão", "Aura", "Cone", "Linha")} for f in LIV["formas"]}
SOMAM_METADE, REMATE, REDUZIDAS = ("Salto", "Queima", "Estilhaço"), "Remate", LIV["reduzidas"]
ABRE = {"bas": (1, 11), "esp": (1, 4, 8, 12, 16, 20, 24, 28), "ext": (0, 0, 0)}
sinal = lambda x: ("+" if x >= 0 else "−") + str(abs(x))


def regra_da_ficha(nivel, usa, ess, inte, trilha, f):
    """os numeros de uma ficha: `f` traz o que o jogador escreveu"""
    acompanha = f.get("aquis", "Espaço conhecido") in ("Espaço conhecido", "Lista de ritual")
    n = nivel if acompanha else max(1, min(nivel, f.get("nivel_fixo") or nivel))
    M, cl = maestria(nivel), classe_maxima(n)
    pts, buf = f.get("pts", [0] * 5), f.get("buff", [0] * 5)
    T = [p + b for p, b in zip(pts, buf)]
    bs, bt = f.get("bstat", [0] * 5), f.get("btr", [0] * 4)
    marc = sum(1 for m in MARCOS if m <= n)
    va = T[ATR.index(f["acerto"])] if f.get("acerto") else None
    cri = f.get("tipo") == fi.CORPO_CRIACAO
    vida = 5 + T[2] + ((4 if cri else 3) + T[2]) * (n - 1) + bs[3]
    fis = T[1] if f.get("fis") == "Destreza" else T[0]
    tr = [fis, T[2], T[3], T[4]]
    tr = [v + (M if f.get("trT") == nome else 0) + b for v, nome, b in zip(tr, ("Físico", "Vigor", "Intelecto", "Espírito"), bt)]
    fam = [x for x in f.get("fam", []) if x]
    maxfam = 5 if trilha == fi.PRINCIPAL else 4 if trilha in (fi.PARCERIA, fi.MULTIPLAS) else 3
    carregado = f.get("tipo") == fi.SH_CRIACAO and f.get("talisma") == fi.COM_CARGA
    return {"n": n, "cl": cl, "M": M, "T": T, "gastos": sum(pts), "disp": 9 + marc, "acima": sum(max(0, p - 3) for p in pts) > marc,
            "ataque": None if va is None else va + M + bs[0], "cd": None if va is None else 8 + va + M + bs[1],
            "defesa": 10 + T[1] + (inte if usa == "Inteligência" else ess) // 2 + bs[2], "vida": vida, "desl": 9 + bs[4], "tr": tr,
            "nper": 4 + pts[3] // 2, "esp": (2 + n / 2) // 2, "ntal": 1 + marc, "db": dados_da_basica(n),
            "entrada": max(1, math.ceil(cl / 2)) if carregado else cl,
            # o retorno da caida: 2 x Classe, e a carga do talisma abate so o que adiantou (Classe - a entrada carregada)
            "retorno": 2 * cl - ((cl - max(1, math.ceil(cl / 2))) if carregado else 0), "reserva": n * (1 + pts[4] // 3),
            "carga": 5 + T[0], "fam_erro": len(set(fam)) < len(fam) or len(fam) > maxfam, "fam": f.get("fam", [])}


def regra_da_carta(F, tipo, num, c):
    """o resultado de uma carta: `F` e a regra da ficha dela, `c` o que o jogador escolheu"""
    C = 0 if tipo == "bas" else max(0 if tipo == "ext" else 1, min(7, c.get("classe", 1)))
    forma = FORMAS[c.get("forma", "Projétil")]
    mel, res = [m for m in c.get("mel", []) if m], [r for r in c.get("res", []) if r]
    fams = [x for x in F["fam"] if x]
    livres = [x for x in F["fam"][:2] if x]
    erros = []
    if ABRE[tipo][num] > F["n"]:
        erros.append("abre")
    if C > F["cl"]:
        erros.append("classe")
    if forma["exige"] and forma["exige"] not in fams:
        erros.append("forma")
    if C == 0 and c.get("forma") in ("Cura", "Onda"):
        erros.append("forma0")
    erros += ["fechada"] * sum(1 for m in mel if PECAS[m]["familia"] and PECAS[m]["familia"] not in fams)
    gasto = 0 if C == 0 else (preco(forma["peso"], C) if forma["peso"] else 0)
    if C > 0:
        gasto += sum(preco_livre(PECAS[m]["peso"], C) if PECAS[m]["familia"] in livres else preco(PECAS[m]["peso"], C) for m in mel)
        if len(mel) > LIMITE[C]:
            erros.append("limite")
        if len(res) + forma["embutida"] > 2:
            erros.append("restricoes")
    else:
        if len(mel) > 1:
            erros.append("limite")
        if any(PECAS[m]["peso"] > 1 for m in mel):
            erros.append("leve")
        if len(res) > 1:
            erros.append("restricoes")
        if any(RESTRICOES[r]["devolve"] > 1 for r in res):
            erros.append("restricao leve")
    dev = 0 if C == 0 else sum(math.ceil(C / 2) if RESTRICOES[r]["devolve"] == 1 else C for r in res) + (C if forma["embutida"] else 0)
    usa = min(dev, 2 * C, gasto)
    saldo = (PONTOS[C] if C else 0) - gasto + usa
    if C > 0 and saldo < 0:
        erros.append("orçamento")
    bases = [RESTRICOES[r]["nome"] for r in res]
    if sum(1 for b in bases if b in FREQUENCIA) > 1:
        erros.append("frequência")
    tem = lambda x: x in mel or x in bases
    erros += ["par"] * sum(1 for a, b in PARES if tem(a) and tem(b))
    if "Corpo a Corpo" in bases:
        erros.append("corpo a corpo")                     # ou a Forma ja traz, ou e Cone ou Linha, ou vira outra Forma
        if not forma["embutida"] and c.get("forma") not in ("Cone", "Linha"):
            erros.pop()
    if "Tudo ou Nada" in bases and not forma["tr"]:
        erros.append("tudo ou nada")
    if "Inescapável" in mel and (len(mel) > 1 or len(res) + forma["embutida"] > 0):
        erros.append("inescapável")
    d = max(0, F["db"] - (1 if mel else 0)) if C == 0 else max(0, saldo)
    dep = 0 if C == 0 else sum(1 for x in SOMAM_METADE if x in mel) * (d // 2)
    if C > 0 and d * (1.25 if REMATE in mel else 1) + dep > 4 * C:
        erros.append("teto")
    tp, td = forma["tipo"], (" " + c["tdano"] if c.get("tdano") else "")
    if C == 0:
        dano = "Sem dano" if tp >= 1 or d == 0 else f"{d}d6{td}"
    elif tp == 2:
        dano = f"{3 * d} de vida temp."
    elif tp == 3:
        dano = "Sem dano"
    elif d == 0:
        dano = "Sem cura" if tp == 1 else "Sem dano"
    elif tp == 1:
        dano = f"Cura {d}d8"
    else:
        dano = f"{d}d8{td}" + (f" · +{dep}d8" if dep else "")
    # 07/10/2026, pelo arnes-invocacoes.py: a linha da CONTA não era conferida nas sorteadas, e a devolução sem o teto de
    # 2 × Classe, ou somada além do gasto, passava calada (o dano da carta sai de outra coluna, que refaz a conta)
    dv = min(dev, 2 * C)
    conta = None if C == 0 else (f"{PONTOS[C]} − {gasto} + {usa} = {saldo} · teto {4 * C}d8" + (f" · a Forma devolve {C}" if forma["embutida"] else "")
                                 + (f" · perde {dv - usa}" if dv > usa else ""))
    # a D43 do livro: Condição, Prende e Cerca só entram na falha de um TR, mesmo numa habilidade de ataque
    pede_tr = any(m in CAT["condicoes"] or m in ("Prende", "Cerca") for m in mel)
    return {"C": C, "saldo": saldo, "erro": bool(erros), "erros": erros, "dano": dano, "pe": "Sem PE" if C == 0 else f"{3 * C} PE",
            "reduzida": C > 0 and any(m in REDUZIDAS for m in mel), "conta": conta, "pede_tr": pede_tr}


# ---------------------------------------------------------------------------------------------
# Preencher, recalcular e ler.
# ---------------------------------------------------------------------------------------------
G = fi.celulas_da_ficha(fi.L0, 0)          # os enderecos do PRIMEIRO lugar; `desloca` os leva aos outros
CJ = fi.celulas_do_conjunto()
LUGARES = fi.lugares()                     # [(numero, fileira, coluna de fichas)]
PASTA = tempfile.mkdtemp(prefix="reg-inv-")
PLANILHAS, CASOS = [], {}                  # cada planilha: um invocador, e ate doze fichas


def desloca(cel, lugar):
    """o endereco do primeiro lugar, levado ao lugar da ficha. O conjunto, a esquerda da primeira lombada, nao anda."""
    lin, col = ix._lc(cel)
    if lugar == 0 or col < fi.lombada(0):
        return cel
    _, j, k = LUGARES[lugar]
    return fi._a1(col + fi.PASSO_COL * k, lin + fi.PASSO_LIN * j)


BARRAS = {desloca(G["barra"], i) for i in range(len(LUGARES))}      # a barra e uma SPARKLINE, que so o Sheets desenha


def indice(wb):
    dd, out = wb["DADOS"], {}
    for r in range(5, 200):
        k, v = dd.cell(row=r, column=53).value, ix.endereco(dd.cell(row=r, column=54).value)
        if k and v:
            out[k] = v
    return out


def monta(nome, nivel, ficha, ess=2, inte=0, usa="Essência", trilha=None, sozinha=False, lugar=None):
    """guarda a ficha para ser escrita. As fichas do mesmo invocador dividem a planilha; `sozinha` da uma planilha so
    para ela (o teste que olha o conjunto ou a aba inteira), no primeiro lugar."""
    chave = (nivel, ess, inte, usa, trilha)
    pl = None if sozinha else next((p for p in PLANILHAS if p["chave"] == chave and not p["fechada"]
                                    and len(p["casos"]) < len(LUGARES) and (lugar is None or lugar not in p["lugares"])), None)
    if pl is None:
        pl = {"chave": chave, "casos": [], "lugares": set(), "fechada": sozinha, "n": len(PLANILHAS)}
        PLANILHAS.append(pl)
    if lugar is None:
        # cada planilha comeca num lugar diferente e anda de cinco em cinco: os exemplos do livro tambem saem do primeiro
        lugar = 0 if sozinha else next(x for x in ((pl["n"] * 5 + 5 * i) % len(LUGARES) for i in range(len(LUGARES))) if x not in pl["lugares"])
    pl["casos"].append(nome)
    pl["lugares"].add(lugar)
    CASOS[nome] = {"planilha": pl, "lugar": lugar, "nivel": nivel, "ess": ess, "inte": inte, "usa": usa, "trilha": trilha, "ficha": ficha}


def escreve_tudo():
    """a planilha gerada e lida uma vez; cada planilha de teste e ela com as fichas escritas, e depois volta ao que era"""
    wb = load_workbook(ARQ)
    idx, f, a = indice(wb), wb["FICHA"], wb[ABA]
    # a arte fica de fora: o openpyxl so consegue gravar a mesma imagem uma vez, e a conta nao depende dela
    for ws in wb:
        ws._images = []
    # o IFS e o TEXTJOIN ficam crus na ficha, porque ela vive no Sheets; o LibreOffice so os reconhece com o prefixo do Excel
    for ws in wb:
        for linha in ws.iter_rows():
            for c in linha:
                if isinstance(c.value, str) and c.value.startswith("="):
                    c.value = re.sub(r"(?<![A-Z_.])(IFS|TEXTJOIN)\(", r"_xlfn.\1(", c.value)
    for pl in PLANILHAS:
        antes = []

        def poe(ws, cel, valor):
            antes.append((ws, cel, ws[cel].value))
            ws[cel] = valor
        nivel, ess, inte, usa, trilha = pl["chave"]
        poe(f, idx["nivel"], nivel); poe(f, idx["atr_base_Essência"], ess); poe(f, idx["atr_base_Inteligência"], inte)
        if trilha:
            poe(f, idx["caminho"], "Evocador"); poe(f, idx["trilha"], trilha)
        poe(a, CJ["usa"], usa)
        for nome in pl["casos"]:
            lug = CASOS[nome]["lugar"]
            la = lambda cel: desloca(cel, lug)
            for k, v in CASOS[nome]["ficha"].items():
                if k == "cartas":
                    for (tipo, n), c in v.items():
                        cel = fi.celulas_da_carta(*G["cartas"][(tipo, n)])
                        for kk, vv in c.items():
                            if isinstance(vv, list):
                                for i, x in enumerate(vv):
                                    poe(a, la(cel[kk][i]), x)
                            else:
                                poe(a, la(cel[kk]), vv)
                elif k == "lib":
                    for i, c in enumerate(v):
                        poe(a, la(G["lib"][i]["classe"]), c)
                elif k == "exp":
                    poe(a, la(G["exp"]["degrau"]), v)
                elif isinstance(v, list):
                    for i, x in enumerate(v):
                        poe(a, la(G[k][i]), x)
                else:
                    poe(a, la(G[k]), v)
        pl["arquivo"] = os.path.join(PASTA, f"{pl['n']:03d}.xlsx")
        wb.save(pl["arquivo"])
        for ws, cel, valor in reversed(antes):
            ws[cel] = valor


def recalcula_tudo():
    saida = os.path.join(PASTA, "saida")
    arquivos = [pl["arquivo"] for pl in PLANILHAS]
    for i in range(0, len(arquivos), 20):            # em lotes: o LibreOffice engasga com cem arquivos numa chamada
        subprocess.run(["libreoffice", "--headless", "--convert-to", "xlsx", "--outdir", saida] + arquivos[i:i + 20], capture_output=True, timeout=1800)
    for pl in PLANILHAS:
        feito = os.path.join(saida, os.path.basename(pl["arquivo"]))
        if not os.path.exists(feito):
            print(f"o LibreOffice nao recalculou {pl['arquivo']}"); sys.exit(1)
        # so as duas abas da invocacao, e so o valor: a leitura inteira da planilha leva quatorze segundos
        wb = load_workbook(feito, data_only=True, read_only=True)
        pl["v"] = {}
        for nome_aba in (ABA, DIV):
            pl["v"][nome_aba] = {f"{ix._letras(c + 1)}{r + 1}": x for r, linha in enumerate(wb[nome_aba].iter_rows(values_only=True))
                                 for c, x in enumerate(linha) if x is not None}
        wb.close()


def le(nome):
    caso = CASOS[nome]
    v, lug = caso["planilha"]["v"][ABA], caso["lugar"]
    return lambda cel: v.get(desloca(cel, lug))


def carta(nome, tipo, n):
    v, cel = le(nome), fi.celulas_da_carta(*G["cartas"][(tipo, n)])
    return {k: ([v(x) for x in c] if isinstance(c, list) else v(c)) for k, c in cel.items()}


# ---------------------------------------------------------------------------------------------
# 1. Os exemplos do livro.
# ---------------------------------------------------------------------------------------------
CAO = {"nome": "Cão de sombra", "pts": [3, 2, 2, 1, 1], "acerto": "Força", "fis": "Força", "trT": "Físico",
       "fam": ["Mira", None, "Alcance", "Controle", None], "per": ["Atletismo", "Furtividade", "Percepção", "Sobrevivência"],
       "tal": ["Farejador"],
       "cartas": {("bas", 0): {"nome": "Mordida", "forma": "Toque", "tdano": "Perfurante"},
                  ("esp", 0): {"nome": "Mordida precisa", "classe": 1, "forma": "Toque", "tdano": "Perfurante", "mel": ["Precisão"]}}}
VIGIA = {"nome": "Vigia de papel", "pts": [0, 2, 2, 3, 2], "acerto": "Inteligência", "fis": "Destreza", "trT": "Intelecto",
         "fam": ["Amparo", None, "Auxiliares", "Alcance", None], "per": ["Acrobacia", "Furtividade", "Investigação", "Percepção", "Sobrevivência"],
         "tal": ["Talento Próprio (CE 1)"],
         "cartas": {("bas", 0): {"nome": "Orientação", "forma": "Apoio", "mel": ["Impulso"]},
                    ("esp", 0): {"nome": "Tiras de resgate", "classe": 2, "forma": "Apoio", "mel": ["Guarda", "Empurrão"]},
                    ("esp", 1): {"nome": "Remendo de papel", "classe": 1, "forma": "Cura"}}}
com = lambda base, **k: {**base, **k}
monta("cão 2", 2, CAO)
monta("cão 4", 4, CAO)
monta("cão 5", 5, CAO, sozinha=True)                 # o teste do conjunto e da lista olha esta planilha
monta("cão 6", 6, com(CAO, pts=[3, 2, 3, 1, 1]))
monta("vigia 5", 5, VIGIA, ess=2, inte=4, usa="Inteligência")
monta("fura", 5, com(CAO, cartas={("esp", 0): {"nome": "Espinho", "classe": 2, "forma": "Projétil", "mel": ["Fura"]}}))
monta("fura sem livre", 5, com(CAO, fam=["Alcance", None, "Mira", "Controle", None],
                               cartas={("esp", 0): {"nome": "Espinho", "classe": 2, "forma": "Projétil", "mel": ["Fura"]}}))
monta("onda", 2, com(VIGIA, cartas={("esp", 0): {"nome": "Roda de papel", "classe": 1, "forma": "Onda", "mel": ["Impulso"], "res": ["Gesto"]}}))
monta("domada", 10, com(CAO, tipo=fi.DOMADA, aquis="Maldição domada", nivel_fixo=10, pts=[2, 2, 2, 2, 3]))
monta("domada menor", 10, com(CAO, tipo=fi.DOMADA, aquis="Maldição domada", nivel_fixo=5))
monta("talismã", 13, com(CAO, tipo=fi.SH_CRIACAO, aquis="Criação", nivel_fixo=13, talisma=fi.COM_CARGA))
monta("corpo de criação", 5, com(CAO, tipo=fi.CORPO_CRIACAO, aquis="Criação", nivel_fixo=5))
monta("trunfos 17", 17, com(CAO, tipo=fi.DOMADA, aquis="Maldição domada", nivel_fixo=17, lib=[3, 4, 7]))
monta("expansão 14", 14, com(CAO, tipo=fi.DOMADA, aquis="Maldição domada", nivel_fixo=14, exp="Completa"))
monta("vazia", 5, {}, sozinha=True)                  # nenhuma caixa da aba inteira acende
# 2. o Buff/Debuff
monta("buff", 5, com(CAO, buff=[1, 0, 0, 0, 0], bstat=[2, 1, -1, 5, 3], btr=[0, 2, 0, -1]))
# 3. o que a ficha recusa
monta("recusa", 2, com(CAO, pts=[4, 2, 2, 1, 0], fam=["Mira", None, "Mira", "Controle", None],
                       cartas={("bas", 0): {"nome": "Mordida", "forma": "Toque", "mel": ["Fura"]},
                               ("esp", 0): {"nome": "Sopro", "classe": 1, "forma": "Cone", "mel": ["Guarda", "Lento"]},
                               ("esp", 1): {"nome": "Cedo demais", "classe": 1, "forma": "Projétil"},
                               ("ext", 0): {"nome": "Estouro", "classe": 1, "forma": "Projétil", "mel": ["Cego", "Prende"]},
                               ("ext", 1): {"nome": "Laço", "classe": 1, "forma": "Projétil", "mel": ["Prende"]},
                               # o Toque já traz uma devolução Média, e com mais duas a soma (3) passa do teto de 2 × Classe
                               ("ext", 2): {"nome": "Garra", "classe": 1, "forma": "Toque", "mel": ["Fura"], "res": ["Sangra", "Sem Volta"]}}))
monta("quatro famílias", 5, com(CAO, fam=["Mira", None, "Alcance", "Controle", "Castigo"]))
monta("quatro famílias, Parceria", 5, com(CAO, fam=["Mira", None, "Alcance", "Controle", "Castigo"]), trilha=fi.PARCERIA)
monta("Múltiplas", 5, CAO, trilha=fi.MULTIPLAS, sozinha=True)
# o inventário pequeno: o Volume do que ela empunha, veste e guarda, contra o limite de 5 + Força (8, com Força 3)
monta("equipada", 5, com(CAO, eq=["Katana", None, "Traje 1"], eq_vol=[1, None, 1], guarda=["Kit de primeiros socorros"], guarda_vol=[0.5]))
monta("carregada demais", 5, com(CAO, eq_vol=[3, 2, 2], guarda_vol=[0.5, None, None, None, None, 1]))
# 08/10/2026 (B41): a Integridade da entidade. O cão de nível 5 tem 27 de vida, e 13 de Integridade (metade, para baixo)
for _nome, _int in (("alma 10", 10), ("alma 9", 9), ("alma 6", 6), ("alma 3", 3), ("alma 0", 0), ("alma demais", 40)):
    monta(_nome, 5, com(CAO, integ=_int))
monta("sem alma", 5, com(CAO, alma=fi.SEM_ALMA, integ=5))
monta("talismã sem carga", 13, com(CAO, tipo=fi.SH_CRIACAO, aquis="Criação", nivel_fixo=13, talisma=fi.SEM_CARGA))

# 4. as fichas sorteadas
random.seed(20261006)
FAMILIAS = list(CAT["familias"])


def sorteia(k, nivel):
    tipo = random.choice([t["nome"] for t in LIV["tipos"]])
    aquis = random.choice([a["nome"] for a in LIV["aquisicoes"]])
    total = 9 + sum(1 for m in MARCOS if m <= nivel) + random.choice([0, 0, 0, 1, -1])
    pts = [0] * 5
    for _ in range(max(0, total)):
        livres = [i for i in range(5) if pts[i] < 6]
        pts[random.choice(livres)] += 1
    fam = random.sample(FAMILIAS, random.choice([3, 3, 3, 4, 5]))
    fam = [fam[0], fam[3] if len(fam) > 4 else None, fam[1], fam[2], fam[4] if len(fam) > 4 else fam[3] if len(fam) > 3 else None]
    f = {"nome": f"Sorteada {k}", "tipo": tipo, "aquis": aquis, "nivel_fixo": random.choice([None, max(1, nivel - 3), nivel]),
         "pts": pts, "buff": [random.choice([0, 0, 0, 1, -1]) for _ in range(5)], "acerto": random.choice(ATR),
         "fis": random.choice(["Força", "Destreza"]), "trT": random.choice(["Físico", "Vigor", "Intelecto", "Espírito"]),
         "bstat": [random.choice([0, 0, 1, -2]) for _ in range(5)], "btr": [random.choice([0, 0, 1]) for _ in range(4)],
         "talisma": random.choice([fi.SEM_CARGA, fi.COM_CARGA]), "fam": fam, "cartas": {}}
    abertas = [x for x in fam if x]
    n_ficha = nivel if aquis in ("Espaço conhecido", "Lista de ritual") else max(1, min(nivel, f["nivel_fixo"] or nivel))
    cl = classe_maxima(n_ficha)
    for tipo_c, n in [("bas", 0), ("bas", 1)] + [("esp", i) for i in range(8)] + [("ext", i) for i in range(3)]:
        # quase sempre dentro da regra, e de vez em quando fora, de proposito: lugar que nao abriu, Classe alta demais,
        # Familia Fechada, peca demais
        certo = random.random() < 0.7
        if ABRE[tipo_c][n] > n_ficha and (certo or random.random() < 0.6):
            continue
        classe = random.randint(0 if tipo_c == "ext" else 1, cl if certo else 7)
        formas = [x for x, d in FORMAS.items() if not certo or ((not d["exige"] or d["exige"] in abertas) and not (tipo_c == "bas" and x in ("Cura", "Onda")))]
        pecas = [p for p, d in PECAS.items() if (d["familia"] in abertas or d["familia"] is None) or (not certo and random.random() < 0.3)]
        if tipo_c == "bas" and (certo or random.random() < 0.5):
            pecas = [p for p in pecas if PECAS[p]["peso"] == 1]
        teto = 1 if tipo_c == "bas" or classe == 0 else LIMITE[classe]
        mel = random.sample(pecas, min(len(pecas), random.randint(0, teto) if certo else random.choice([1, 2, 3, 4])))
        res = random.sample(list(RESTRICOES), random.choice([0, 0, 0, 1]) if certo else random.choice([0, 1, 2]))
        f["cartas"][(tipo_c, n)] = {"nome": f"H{tipo_c}{n}", "classe": classe, "forma": random.choice(formas),
                                    "tdano": random.choice(LIV["tipos_de_dano"]), "mel": mel + [None] * (4 - len(mel)), "res": res + [None] * (2 - len(res))}
        if tipo_c == "bas":
            del f["cartas"][(tipo_c, n)]["classe"]
    return f


# seis invocadores sorteados, cada um com tres fichas; os lugares de cada trio andam pela grade, e os seis juntos passam
# pelos doze lugares
SORTEADAS = []
NIVEIS = [2, 3, 4, 5, 8, 9, 11, 13, 16, 17, 21, 26, 30]
for g_ in range(6):
    nivel = random.choice(NIVEIS)
    usa = random.choice(["Essência", "Inteligência"])
    trilha = random.choice([None, None, fi.PRINCIPAL, fi.PARCERIA, fi.MULTIPLAS])
    ess, inte = random.randint(0, 6), random.randint(0, 6)
    for lug in ((g_ * 2) % len(LUGARES), (g_ * 2 + 7) % len(LUGARES), (g_ * 2 + 5) % len(LUGARES)):
        k = len(SORTEADAS)
        monta(f"sorteada {k}", nivel, sorteia(k, nivel), ess=ess, inte=inte, usa=usa, trilha=trilha, lugar=lug)
        SORTEADAS.append(f"sorteada {k}")

print(f"\nescrevendo {len(CASOS)} fichas em {len(PLANILHAS)} planilhas e recalculando no LibreOffice...")
escreve_tudo()
recalcula_tudo()
_usados = {c["lugar"] for c in CASOS.values()}
checa(f"as {len(CASOS)} fichas de teste passam pelos {len(LUGARES)} lugares da grade (as duas colunas de fichas e as seis fileiras)",
      len(LUGARES) == fi.N_COLUNAS * fi.N_FILEIRAS and _usados == set(range(len(LUGARES)))
      and {CASOS[n]["lugar"] for n in SORTEADAS} == set(range(len(LUGARES))), str(sorted(_usados)))

print("\n1. OS EXEMPLOS DO LIVRO")
v = le("cão 2")
checa("Cão de sombra, nível 2: ataque +4 e CD 12, Defesa 13, vida 12 (o quadro do livro)",
      [v(c) for c in G["stat"]][:4] == ["+4", 12, 13, 12], str([v(c) for c in G["stat"]]))
checa("Cão de sombra: Físico +4 treinado, Vigor +2, Intelecto +1, Espírito +1", [v(c) for c in G["tr"]] == ["+4", "+2", "+1", "+1"], str([v(c) for c in G["tr"]]))
checa("Cão de sombra: 9 de 9 pontos, quatro perícias, um espaço de especial, um talento",
      v(G["pontos_rot"]) == "ATRIBUTOS · 9 DE 9 PONTOS" and v(G["per_rot"]) == "PERÍCIAS · 4 DE 4" and "1 espaço de especial até a Classe 1 · 1 talento" in v(G["hab"]),
      f'{v(G["pontos_rot"])} | {v(G["per_rot"])} | {v(G["hab"])}')
c = carta("cão 2", "bas", 0)
checa("Mordida: básica de Classe 0, ataque +4, 1d6 de Perfurante, sem PE", (c["dano"], c["resolve"], c["pe"], c["estado"]) == ("1d6 Perfurante", "Ataque +4", "Sem PE", "Na regra"), str(c))
c = carta("cão 2", "esp", 0)
checa("Mordida precisa: 2 − 1 + 1 = 2d8, ataque +6 com a Precisão, 3 PE, a 1,5 m",
      (c["dano"], c["resolve"], c["pe"], c["estado"], c["alcance"]) == ("2d8 Perfurante", "Ataque +6", "3 PE", "Na regra", "1,5 m, um alvo") and c["conta"].startswith("2 − 1 + 1 = 2"), str(c))
checa("a Precisão custa 1 mesmo com a Mira Livre, pelo preço mínimo", c["preco"][0] == "−1 · Leve · Livre", str(c["preco"]))
# 07/10/2026, a D43 do livro: Condição, Prende e Cerca só entram na falha de um TR, mesmo numa habilidade de ataque
_c, _l = carta("recusa", "ext", 0), carta("recusa", "ext", 1)
checa("o ataque com Condição ou Prende mostra também o TR e a CD dela (Estouro: Projétil com Cego e Prende; Laço: só com Prende); sem essas peças, só o ataque",
      all(re.fullmatch(r"Ataque [+−-]\d+ · TR CD \d+", x["resolve"] or "") is not None for x in (_c, _l)) and " · TR" not in c["resolve"],
      f'{_c["resolve"]} | {_l["resolve"]} | {c["resolve"]}')
# pelo arnes-invocacoes.py: nenhuma carta sorteada tinha devolução acima do teto, e tirar o teto passava calado
_k = CASOS["recusa"]
_rg = regra_da_carta(regra_da_ficha(_k["nivel"], _k["usa"], _k["ess"], _k["inte"], _k["trilha"], _k["ficha"]), "ext", 2, _k["ficha"]["cartas"][("ext", 2)])
_lg = carta("recusa", "ext", 2)
checa("a devolução para em 2 × Classe: o Toque com duas Restrições Médias na Classe 1 devolveria 3, e a conta trabalha com 2",
      _lg["conta"] == _rg["conta"] and str(_rg["conta"]).endswith(" · a Forma devolve 1 · perde 1"), f'{_lg["conta"]} | {_rg["conta"]}')
v = le("cão 4")
checa("no nível 4 o cão tem 22 de vida e um segundo espaço de especial", v(G["stat"][3]) == 22 and "2 espaços de especial" in v(G["hab"]), f'{v(G["stat"][3])} | {v(G["hab"])}')
v = le("cão 5")
c = carta("cão 5", "esp", 0)
checa("no nível 5: 27 de vida, básica de 2d6, Classe 2, e a Mordida precisa amplia para 4d8 por 6 PE",
      v(G["stat"][3]) == 27 and carta("cão 5", "bas", 0)["dano"] == "2d6 Perfurante" and v(G["status"][1]) == 2 and c["ampliar"] == "2 → 4d8, 6 PE",
      f'{v(G["stat"][3])} | {carta("cão 5", "bas", 0)["dano"]} | {c["ampliar"]}')
v = le("cão 6")
checa("no nível 6, com a Constituição em 3: 8 + 6 × 5 = 38 de vida, e o segundo talento", v(G["stat"][3]) == 38 and "2 talentos" in v(G["hab"]) and v(G["pontos_rot"]) == "ATRIBUTOS · 10 DE 10 PONTOS",
      f'{v(G["stat"][3])} | {v(G["hab"])} | {v(G["pontos_rot"])}')
v = le("vigia 5")
checa("Vigia de papel, nível 5, com a Inteligência 4 da Bruna: ataque +4 e CD 12, Defesa 14, vida 27",
      [v(c) for c in G["stat"]][:4] == ["+4", 12, 14, 27], str([v(c) for c in G["stat"]]))
checa("Vigia: Intelecto +4 treinado, Físico com Destreza +2, Vigor +2, Espírito +2, e cinco perícias",
      [v(c) for c in G["tr"]] == ["+2", "+2", "+4", "+2"] and v(G["per_rot"]) == "PERÍCIAS · 5 DE 5", f'{[v(c) for c in G["tr"]]} | {v(G["per_rot"])}')
c0, c1, c2 = carta("vigia 5", "bas", 0), carta("vigia 5", "esp", 0), carta("vigia 5", "esp", 1)
checa("Orientação: básica de Apoio com Impulso, sem dano", c0["dano"] == "Sem dano" and c0["estado"] == "Na regra" and c0["preco"][0] == "−1 dado", str(c0))
checa("Tiras de resgate: 4 − 2 − 1 = 1 ponto, 3 de vida temporária, 6 PE, aliado a 9 m",
      (c1["dano"], c1["pe"], c1["estado"], c1["alcance"]) == ("3 de vida temp.", "6 PE", "Na regra", "Um aliado a 9 m") and c1["preco"][:2] == ["−2 · Média", "−1 · Leve"], str(c1))
checa("Remendo de papel: 2 − 1 = 1d8 de cura, 3 PE", (c2["dano"], c2["pe"], c2["estado"]) == ("Cura 1d8", "3 PE", "Na regra"), str(c2))
c = carta("fura", "esp", 0)
checa("Fura de Classe 2 em Projétil, com a Mira Livre: custa 1 e fica com 3d8, e o efeito usa a Classe 1",
      c["dano"].startswith("3d8") and c["preco"][0] == "−1 · Média · Livre" and "Efeito de Fura pela Classe 1" in c["avisos"], str(c))
c = carta("fura sem livre", "esp", 0)
checa("o mesmo Fura sem a Mira Livre custa 2 e fica com 2d8", c["dano"].startswith("2d8") and c["preco"][0] == "−2 · Média", str(c))
c = carta("onda", "esp", 0)
checa("Onda de Classe 1 com Impulso e Gesto cabe nos 2 pontos, com saldo zero", c["estado"].startswith(fi.T_AVISO) is False and fi.T_ERRO not in c["estado"] and c["conta"].startswith("2 − 3 + 1 = 0"), str(c))
v = le("domada")
checa("a domada de nível 10 e Essência 3 tem 20 PE de reserva, entra por 3 PE e retorna por 6",
      v(G["reserva_rot"]) == "RESERVA · MÁX. 20" and v(G["status"][3]) == "3 PE" and v(G["status"][4]) == "6 PE", f'{v(G["reserva_rot"])} {v(G["status"][3])} {v(G["status"][4])}')
v = le("domada menor")
checa("a domada de nível 5 de um invocador de nível 10: o nível é o dela, a maestria é a dele (ataque 3 + 2)",
      v(G["status"][0]) == 5 and v(G["status"][2]) == "+2" and v(G["stat"][0]) == "+5" and v(G["stat"][3]) == 27, f'{[v(c) for c in G["status"]]} {v(G["stat"][0])}')
v = le("talismã")
checa("a entidade de nível 13 (Classe 4) com o talismã carregado entra por 2 PE", v(G["status"][1]) == 4 and v(G["status"][3]) == "2 PE", str([v(c) for c in G["status"]]))
# o exemplo do livro, em Talismãs: "Se ela estivesse caída, o retorno custaria 8 PE ao todo: os 2 adiantados e mais 6 no retorno"
checa("e, caída, volta por mais 6 PE (os 8 do retorno, menos os 2 que a carga adiantou); sem carga, entra por 4 e volta por 8",
      v(G["status"][4]) == "6 PE" and (le("talismã sem carga")(G["status"][3]), le("talismã sem carga")(G["status"][4])) == ("4 PE", "8 PE"),
      f'{v(G["status"][4])} | {le("talismã sem carga")(G["status"][3])} {le("talismã sem carga")(G["status"][4])}')
v = le("corpo de criação")
checa("o corpo de criação de nível 5 com Constituição 2 tem 7 + 6 × 4 = 31 de vida", v(G["stat"][3]) == 31, str(v(G["stat"][3])))
v = le("trunfos 17")
checa("Liberação de Classe 3: 6 pontos + 3d8, por 14 PE; a de Classe 7, 16 pontos + 7d8, por 32 PE",
      (v(G["lib"][0]["pontos"]), v(G["lib"][0]["pe"]), v(G["lib"][2]["pontos"]), v(G["lib"][2]["pe"])) == ("6 pontos + 3d8", "14 PE", "16 pontos + 7d8", "32 PE"),
      str([(v(l["pontos"]), v(l["pe"])) for l in G["lib"]]))
checa("Técnica Máxima da domada de nível 17: 19d8 fixos, 8 pontos, 25 PE", (v(G["tm"]["dados"]), v(G["tm"]["pe"])) == ("19d8 fixos · 8 pontos", "25 PE"), f'{v(G["tm"]["dados"])} {v(G["tm"]["pe"])}')
checa("antes do nível 17 a Técnica Máxima diz quando abre", le("cão 5")(G["tm"]["dados"]) == "Abre no nível 17" and le("cão 5")(G["tm"]["pe"]) == "—")
v = le("expansão 14")
checa("Expansão completa da domada de nível 14 (Classe 4): 24 PE e Acerto de 4d8", (v(G["exp"]["pe"]), v(G["exp"]["acerto"])) == ("24 PE", "Acerto 4d8") and "nível 14" in v(G["exp"]["pede"]),
      f'{v(G["exp"]["pe"])} {v(G["exp"]["acerto"])} {v(G["exp"]["pede"])}')
v = le("vazia")
checa("a ficha vazia não mostra erro nem número de ataque", "lugar vazio" in v(G["titulo"]) and v(G["stat"][0]) in (None, "") and
      all(fi.T_ERRO not in str(x) for x in CASOS["vazia"]["planilha"]["v"][ABA].values()))
v = le("cão 5")
checa("o conjunto: nível 5, maestria +1, até duas ativas, e a lista traz a ficha no lugar dela e as outras vazias",
      (v(CJ["nivel"]), v(CJ["maestria"]), v(CJ["ativas"])) == (5, "+1", 2)
      and v(CJ["rol"][0]) == "1 · Cão de sombra · nv 5 · 27/27" and [v(c) for c in CJ["rol"][1:]] == [f"{i} · vazia" for i in range(2, len(LUGARES) + 1)]
      and v(CJ["rol_rot"]) == f"AS INVOCAÇÕES · 1 DE {len(LUGARES)}",
      f'{v(CJ["nivel"])} {v(CJ["maestria"])} {v(CJ["ativas"])} {v(CJ["rol"][0])} | {v(CJ["rol"][1])} | {v(CJ["rol_rot"])}')
checa("com Múltiplas Invocações o limite de ativas vai a quatro", le("Múltiplas")(CJ["ativas"]) == 4, str(le("Múltiplas")(CJ["ativas"])))
_eq, _dm = le("equipada")(G["equip_rot"]), le("carregada demais")(G["equip_rot"])
checa("o equipamento soma o Volume do que ela empunha, veste e guarda (1 + 1 + 0,5 de 8), e acende quando passa do limite (8,5 de 8)",
      re.fullmatch(r"EQUIPAMENTO · 2[.,]5 DE 8 DE VOLUME", str(_eq)) is not None and str(_dm).startswith(fi.T_ERRO)
      and re.search(r"EQUIPAMENTO · 8[.,]5 DE 8 DE VOLUME · PASSOU DO LIMITE$", str(_dm)) is not None
      and le("equipada")(G["guarda_rot"]) == "GUARDADO · SÓ COM A CARACTERÍSTICA DE TRANSPORTE", f"{_eq} | {_dm}")
# 08/10/2026 (B41): a Integridade da entidade com alma. Metade da vida máxima, para baixo (27 -> 13); nasce cheia; os
# estágios pelo que falta (um quarto, metade, três quartos, toda), com frações exatas: de 13, perder 3 ainda não é um
# quarto (3,25), perder 4 é; e o corpo sem alma não tem Integridade
_i = lambda nome: (le(nome)(G["integ_max"]), le(nome)(G["estagio"]))
checa("a Integridade da entidade é metade da vida máxima (13 de 27), nasce cheia e sem estágio, e a ficha vazia não mostra nada",
      (v(G["integ"]), v(G["integ_max"]), v(G["estagio"]), v(G["alma"])) == (13, 13, fi.ESTAGIOS[0], fi.COM_ALMA)
      and [le("vazia")(G[k]) or "" for k in ("integ", "integ_max", "estagio")] == ["", "", ""],
      f'{v(G["integ"])} {v(G["integ_max"])} {v(G["estagio"])} | {[le("vazia")(G[k]) for k in ("integ", "integ_max", "estagio")]}')
_esp = {"alma 10": fi.ESTAGIOS[0], "alma 9": fi.ESTAGIOS[1], "alma 6": fi.ESTAGIOS[2], "alma 3": fi.ESTAGIOS[3], "alma 0": fi.ESTAGIOS[4],
        "alma demais": fi.ESTAGIOS[0]}
checa("os estágios da entidade seguem o que falta da Integridade, em fração exata: com 10 de 13 inteira, 9 estágio 1, 6 estágio 2, 3 estágio 3, 0 estágio 4",
      all(_i(n) == (13, e) for n, e in _esp.items()), str({n: _i(n) for n in _esp}))
checa("a entidade sem alma não tem Integridade: a máxima fica em traço e o estágio diz que ela é imune ao dano de Alma",
      _i("sem alma") == ("—", fi.ESTAGIOS[5]), str(_i("sem alma")))
# 07/10/2026: a vida como na FICHA do jogador, e o que saiu da mesa
checa("a VIDA MÁXIMA aparece ao lado da VIDA ATUAL, com o mesmo número da caixa dos números (27), e o título da ficha diz a vida",
      v(G["vida_max"]) == 27 == v(G["stat"][3]) and v(G["titulo"]).endswith("vida 27 de 27"), f'{v(G["vida_max"])} {v(G["stat"][3])} {v(G["titulo"])}')
# 07/10/2026, depois de ele montar no Sheets: "seria bom ao preencher a vida máxima, a vida atual preencher tbm, na criação da ficha"
checa("a VIDA ATUAL nasce cheia: na ficha com nome em que ninguém escreveu a vida ela mostra a máxima (27), e na ficha vazia fica em branco",
      v(G["vida"]) == 27 and le("vazia")(G["vida"]) in (None, ""), f'{v(G["vida"])} | {le("vazia")(G["vida"])!r}')
_rotulos = {str(x).split(" · ")[0] for x in CASOS["cão 5"]["planilha"]["v"][ABA].values() if isinstance(x, str)}
_fora = ("CARGA MÁXIMA", "TAREFA", "MOVIMENTO", "BÁSICA DO CICLO", "ESTADO", "ORDEM PENDENTE", "REAÇÃO COLETIVA", "DANO NO TURNO", "VÍNCULO",
         "APRIMORAMENTO DE VÍNCULO", "EM QUEM", "OFENSIVO", "PROTEÇÃO", "PERÍCIA")
checa("a aba não traz mais as caixas de turno (tarefa, movimento, básica do ciclo, estado, ordem, reação, dano no turno, Vínculo, usos da rodada)",
      not [x for x in _fora if x in _rotulos] and {"VIDA ATUAL", "VIDA MÁXIMA", "TEMPORÁRIA", "± PERDA / GANHO", "CONDIÇÕES E USOS GASTOS", "ANOTAÇÕES"} <= _rotulos,
      str([x for x in _fora if x in _rotulos]))
checa("o rótulo de cada talento é sempre o nível e a Categoria, sem o ABRE NO",
      [v(c) for c in G["tal_nv"]][:3] == ["NV 1 · CE 1", "NV 6 · CE 1", "NV 10 · CE 2"] and not any(str(x).startswith("ABRE NO") for x in CASOS["cão 5"]["planilha"]["v"][ABA].values()),
      str([v(c) for c in G["tal_nv"]]))

print("\n2. O BUFF/DEBUFF")
v = le("buff")
checa("Força +1 no Buff/Debuff: o total vai a 4, e o ataque (+2 no dele) a +7, a CD (+1) a 14, a carga a 9",
      v(G["total"][0]) == 4 and v(G["stat"][0]) == "+7" and v(G["stat"][1]) == 14 and v(G["equip_rot"]) == "EQUIPAMENTO · 0 DE 9 DE VOLUME" and v(G["pontos_rot"]) == "ATRIBUTOS · 9 DE 9 PONTOS",
      f'{v(G["total"][0])} {v(G["stat"][0])} {v(G["stat"][1])} {v(G["equip_rot"])} {v(G["pontos_rot"])}')
checa("Defesa −1, vida máxima +5 e deslocamento +3 m", (v(G["stat"][2]), v(G["stat"][3]), v(G["stat"][4])) == (12, 32, "12 m"), str([v(c) for c in G["stat"]]))
checa("o Buff/Debuff de cada Teste de Resistência", [v(c) for c in G["tr"]] == ["+5", "+4", "+1", "+0"], str([v(c) for c in G["tr"]]))

print("\n3. O QUE A FICHA RECUSA")
v = le("recusa")
checa("atributo em 4 no nível 2, sem marco: a caixa dos pontos acende", v(G["pontos_rot"]).startswith(fi.T_ERRO) and "ACIMA DE 3" in v(G["pontos_rot"]), str(v(G["pontos_rot"])))
checa("Família repetida: a caixa das Famílias acende", v(G["fam_rot"]).startswith(fi.T_ERRO) and "REPETIDA" in v(G["fam_rot"]), str(v(G["fam_rot"])))
c = carta("recusa", "bas", 0)
checa("básica com Melhoria Média: recusada", fi.T_ERRO in c["estado"] and "só aceita Melhoria Leve" in c["avisos"], str(c["avisos"]))
c = carta("recusa", "esp", 0)
checa("Cone sem Área aberta, e Guarda de Família Fechada: os dois erros", "pede Área aberta" in c["avisos"] and "Família Fechada: Guarda" in c["avisos"], str(c["avisos"]))
c = carta("recusa", "esp", 1)
checa("habilidade num lugar que só abre no nível 4", "Este lugar abre no nível 4" in c["avisos"] and c["titulo"] == "ESPECIAL 2 · ABRE NO NÍVEL 4", f'{c["titulo"]} | {c["avisos"]}')
c = carta("recusa", "ext", 0)
checa("orçamento estourado: uma condição Pesada e uma Melhoria Média de Controle numa Classe 1 (2 + 1 contra 2 pontos)",
      "Orçamento estourado: faltam 1 ponto" in c["avisos"], str(c["avisos"]))
checa("quatro Famílias sem Trilha: acende; com Parceria, não", le("quatro famílias")(G["fam_rot"]).startswith(fi.T_ERRO) and
      le("quatro famílias, Parceria")(G["fam_rot"]) == "FAMÍLIAS · 4 ABERTAS", f'{le("quatro famílias")(G["fam_rot"])} | {le("quatro famílias, Parceria")(G["fam_rot"])}')

print("\n4. AS FICHAS SORTEADAS")
dif, cartas_vistas, com_erro, contas_vistas, ataques_vistos, com_tr = [], 0, 0, 0, 0, 0
for nome in SORTEADAS:
    caso, v = CASOS[nome], le(nome)
    f = caso["ficha"]
    F = regra_da_ficha(caso["nivel"], caso["usa"], caso["ess"], caso["inte"], caso["trilha"], f)
    esperado = {"nível": F["n"], "classe": F["cl"], "maestria": sinal(F["M"]), "entrada": f"{F['entrada']} PE", "retorno": f"{F['retorno']} PE",
                "totais": F["T"], "ataque": sinal(F["ataque"]), "cd": F["cd"], "defesa": F["defesa"], "vida": F["vida"],
                "desl": f"{F['desl']} m", "tr": [sinal(x) for x in F["tr"]], "carga": F["carga"],
                "pontos acendem": F["gastos"] > F["disp"] or F["acima"], "famílias acendem": F["fam_erro"]}
    lido = {"nível": v(G["status"][0]), "classe": v(G["status"][1]), "maestria": v(G["status"][2]), "entrada": v(G["status"][3]),
            "retorno": v(G["status"][4]), "totais": [v(c) for c in G["total"]], "ataque": v(G["stat"][0]), "cd": v(G["stat"][1]),
            "defesa": v(G["stat"][2]), "vida": v(G["stat"][3]), "desl": v(G["stat"][4]), "tr": [v(c) for c in G["tr"]],
            "carga": int(re.search(r" DE (\d+) DE VOLUME", str(v(G["equip_rot"]))).group(1)),
            "pontos acendem": str(v(G["pontos_rot"])).startswith(fi.T_ERRO), "famílias acendem": str(v(G["fam_rot"])).startswith(fi.T_ERRO)}
    dif += [f"{nome}: {k} = {lido[k]!r}, e a regra diz {esperado[k]!r}" for k in esperado if lido[k] != esperado[k]]
    for (tipo, n), c in f["cartas"].items():
        r, lc = regra_da_carta(F, tipo, n, c), carta(nome, tipo, n)
        cartas_vistas += 1
        com_erro += r["erro"]
        if lc["dano"] != r["dano"] or lc["pe"] != r["pe"] or (fi.T_ERRO in str(lc["estado"])) != r["erro"] or ("pela Classe" in str(lc["avisos"])) != r["reduzida"]:
            dif.append(f"{nome} {tipo}{n} {c}: a planilha diz {lc['dano']!r}, {lc['pe']!r}, {lc['estado']!r} ({lc['avisos']!r}); a regra, {r['dano']!r}, {r['pe']!r}, erros {r['erros']}")
        if r["conta"] is not None and lc["conta"] != r["conta"]:
            dif.append(f"{nome} {tipo}{n} {c}: a CONTA da planilha é {lc['conta']!r}; a regra, {r['conta']!r}")
            contas_vistas -= 1
        contas_vistas += r["conta"] is not None
        if str(lc["resolve"]).startswith("Ataque"):
            ataques_vistos += 1
            com_tr += r["pede_tr"]
            if (" · TR" in str(lc["resolve"])) != r["pede_tr"]:
                dif.append(f"{nome} {tipo}{n} {c}: a carta de ataque diz {lc['resolve']!r}, e {'devia' if r['pede_tr'] else 'não devia'} pedir o TR")
checa(f"{len(SORTEADAS)} fichas sorteadas: os números de cada uma e as {cartas_vistas} cartas ({com_erro} fora da regra) batem com a regra escrita aqui, "
      f"a linha da conta de {contas_vistas} e o TR de {com_tr} das {ataques_vistos} de ataque inclusive",
      not dif and cartas_vistas > 100 and 20 < com_erro < cartas_vistas - 20 and contas_vistas > 60 and ataques_vistos > 20 and 3 < com_tr < ataques_vistos,
      " || ".join(dif[:4]) or f"{cartas_vistas} cartas, {com_erro} com erro, {contas_vistas} contas, {com_tr} de {ataques_vistos} ataques com TR")
erros_de_formula = [(pl["casos"][0], aba_, cel, x) for pl in PLANILHAS for aba_ in (ABA, DIV) for cel, x in pl["v"][aba_].items()
                    if isinstance(x, str) and re.match(r"^(#[A-Z/0!?]+|Err:\d+)$", x)
                    and not (aba_ == ABA and cel in BARRAS)]     # a barra é uma SPARKLINE, que só o Sheets desenha
checa(f"nenhuma fórmula da aba ou da DADOS_INVOC dá erro em nenhuma das {len(PLANILHAS)} planilhas ({len(CASOS)} fichas)", not erros_de_formula, str(erros_de_formula[:4]))

shutil.rmtree(PASTA, ignore_errors=True)
print()
print("=" * 74)
if FALHAS:
    print(f"REGRESSAO DA INVOCAÇÕES: {len(FALHAS)} FALHA(S)")
    for f_ in FALHAS:
        print("  - " + f_)
    sys.exit(1)
print("REGRESSAO DA INVOCAÇÕES: tudo passou")
