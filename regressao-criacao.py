#!/usr/bin/env python3
"""Recalcula as Famílias da 3c com expectativas aritméticas independentes.

Arquivos locais são apenas calculadora de apoio, nunca importação para Sheets.
Comparação cobre os limites, catálogo, estados, avisos e códigos de preços do
livro. Menus, visual e operação no Google têm conferência nativa separada.
"""
import argparse
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import range_boundaries, get_column_letter

ROOT = Path(__file__).resolve().parent
sys.path[:0] = [str(ROOT/'ficha-v01'), str(ROOT/'ficha')]
import custom_criacao as custom
import ficha_amaldicoada as fa
import indice_ficha as ix


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--modelo', type=Path, required=True)
    p.add_argument('--base', type=Path, required=True)
    args = p.parse_args()
    model = json.loads(args.modelo.read_text())
    raw = {a['nome']: {v[0]: v[1] for v in a['celulas']} for a in model['layout']['abas']}
    base_wb = load_workbook(args.base, read_only=True, data_only=False)
    base = {s: {c.coordinate: c.value for row in base_wb[s] for c in row if c.value is not None}
            for s in [fa.NOME, fa.DADOS_AM]}
    base_wb.close()
    checks = []

    def check(name, ok, detail=None):
        print(('OK ' if ok else 'FALHA ')+name+(' '+str(detail) if not ok else ''), flush=True)
        checks.append(ok)

    check('capacidade de Famílias deriva das peças da técnica e nove aberturas de cada entidade',
          custom.FAMILIAS == 45+165+12*9 and len(model['familias']) == custom.FAMILIAS)
    check('referências escritas dentro de texto não são reescritas',
          all(custom.troca_refs('="DADOS_AM!$H$2:$K$10"', 'DADOS_AM', {('DADOS_AM','H2:K10'):'OTHER'}) ==
              '="DADOS_AM!$H$2:$K$10"' for _ in [0]))
    # Legacy ANALISE sources, values and ordering remain; only the format rises.
    old = load_workbook(args.base, read_only=True, data_only=False)
    prior = [tuple(c.value for c in row) for row in old['ANALISE'].iter_rows(min_row=2, max_col=4)]
    old.close()
    now = raw['ANALISE']
    differences = []
    for r, row in enumerate(prior, 2):
        expected = ('formato.versao', '8', *row[2:]) if row[0] == 'formato.versao' else row
        actual = tuple(now.get(c+str(r)) for c in 'ABCD')
        if tuple(None if v=='' else v for v in actual) != tuple(None if v=='' else v for v in expected):
            differences.append((r, actual, expected))
    check('todas as 7485 chaves, fontes e posições anteriores preservadas', len(prior)==7485 and not differences, differences[:2])
    all_keys = [v for c, v in now.items() if re.fullmatch(r'A\d+', c) and c != 'A1']
    check('1597 chaves novas sem duplicatas', len(all_keys)==7485+1597 and len(set(all_keys))==len(all_keys))
    inv_sources = [m for a in model['layout']['abas'] if a['nome']=='INVOCAÇÕES'
                   for m in a['menus'] if m['formula']==model['catalogo_familias']]
    inv_cells = set()
    for m in inv_sources:
        for area in m['onde'].split():
            c,r,cc,rr=range_boundaries(area)
            inv_cells.update(get_column_letter(x)+str(y) for x in range(c,cc+1) for y in range(r,rr+1))
    expected_inv = {get_column_letter(c+19*k)+str(r+198*j) for k in range(2) for j in range(6)
                    for r,c in [(48,c) for c in range(15,20)]+[(50,c) for c in range(15,19)]}
    check('os 108 menus usam o catálogo livro + custom', inv_cells==expected_inv, len(inv_cells))
    check('auxiliar dos catálogos permanece oculto', next(a for a in model['layout']['abas'] if a['nome']==custom.NOME)['estado']=='hidden')

    g = fa.geometria()
    state_cells = [col+str(g['familias']+1) for col, _ in g['cols_fam']]
    book = json.loads((ROOT/'catalogo-projeto-m.json').read_text())
    book_names = list(book['familias'])
    counters = {}
    for coord, value in raw[fa.DADOS_AM].items():
        if value in ('famílias livres', 'famílias fechadas'):
            r, c = ix._lc(coord)
            counters[value] = ix._letras(c+1)+str(r)
    summary = g['resumo_familias']
    base_summary_link = re.fullmatch(r"=DADOS_AM!(\$?[A-Z]+\$?\d+)", base[fa.NOME][summary])
    summary_cache = base_summary_link[1].replace('$', '') if base_summary_link else None
    first, second = model['familias'][:2]
    # Slots are inputs, independent of the formulas under test.
    def inputs(name='', state='Neutra', name2='', state2='Neutra', seen=False):
        return {first['nome']: name, first['estado']: state, first['mestre_viu']: seen,
                second['nome']: name2, second['estado']: state2}
    cases = [
        ('vazio', {}, {}, (9,4,2,3,False,False)),
        ('neutra', inputs('Ressonância'), {}, (10,5,2,3,False,False)),
        ('livre', inputs('Ressonância','Livre'), {157:3}, (10,4,3,3,False,False)),
        ('fechada', inputs('Ressonância','Fechada'), {158:4}, (10,4,2,4,False,False)),
        ('excesso_estado', inputs('Ressonância','Livre'), {}, (10,5,3,3,True,False)),
        ('duplicata', inputs('Ressonância','Neutra','ressonância'), {}, (9,4,2,3,False,True)),
        ('livro_repetido', inputs(book_names[0]), {}, (9,4,2,3,False,True)),
        ('asterisco_literal', inputs('Lua*','Neutra','Lua X'), {}, (11,6,2,3,False,False)),
        ('estado_invalido', inputs('Ressonância','Aberta'), {}, (9,4,2,3,False,True)),
        ('checkbox_invalido', inputs('Ressonância',seen=1), {}, (9,4,2,3,False,True)),
        ('limites_excessivos', {}, {157:9,158:1}, (9,-1,2,3,True,False)),
    ]
    # Only independent family lookup columns needed to test price codes.
    def selected_data(source):
        a = source[fa.DADOS_AM]
        headers = {ix._lc(c)[1]: v for c, v in a.items() if ix._lc(c)[0]==1}
        starts = [c for c, v in headers.items() if v in ('forma', 'peça')]
        cols = {c+i for c in starts for i in range(5)}
        cols.update(c for c, v in headers.items() if v in ('família da técnica','estado da família','família livre','família fechada'))
        return {c:v for c,v in a.items() if ix._lc(c)[1] in cols or c in counters.values() or c==summary_cache}
    normalize = lambda v: None if v == '' else v
    results = {}
    with tempfile.TemporaryDirectory(prefix='b47-criacao-') as tmp:
        tmp = Path(tmp)
        out = tmp/'out'; out.mkdir()
        paths = []
        for name, edits, limits, expected in cases + [('base_anterior', {}, {}, None)]:
            previous = name=='base_anterior'
            wb = Workbook(); wb.remove(wb.active)
            for s in [fa.NOME, fa.DADOS_AM, 'CONFIGURAÇÃO', custom.NOME, 'RESULTADO']:
                wb.create_sheet(s)
            sources = {fa.DADOS_AM:selected_data(base if previous else raw)}
            if not previous:
                sources[custom.NOME] = raw[custom.NOME]
                sources['CONFIGURAÇÃO'] = {c:v for c,v in raw['CONFIGURAÇÃO'].items() if
                    (156<=ix._lc(c)[0]<=164 and 2<=ix._lc(c)[1]<=7) or
                    (ix._lc(c)[0]>=custom.INICIO and v is not None and not (isinstance(v,str) and v.startswith('=')))}
            for s, vals in sources.items():
                for c,v in vals.items():
                    if v is None: continue
                    if isinstance(v,str) and v.startswith('='):
                        v=re.sub(r'(?<![A-Z_.])TEXTJOIN\(', '_xlfn.TEXTJOIN(',v)
                    wb[s][c]=v
            wb[fa.NOME][summary]=(base if previous else raw)[fa.NOME][summary]
            for i,c in enumerate(state_cells):
                wb[fa.NOME][c]='Livre' if i<2 else 'Fechada' if i<5 else 'Neutra'
            for c,v in edits.items():wb['CONFIGURAÇÃO'][c]=v
            for r,v in limits.items():wb['CONFIGURAÇÃO']['F'+str(r)]=v
            for i,(s,c) in enumerate([('CONFIGURAÇÃO','G160'),('CONFIGURAÇÃO','G159'),
                  (fa.DADOS_AM,counters['famílias livres']), (fa.DADOS_AM,counters['famílias fechadas']),
                  (fa.NOME,summary),(custom.NOME,'F2'),('CONFIGURAÇÃO','B162')],1):
                wb['RESULTADO']['A'+str(i)]='='+custom.ref(c,s)
            path=tmp/(name+'.xlsx'); wb.save(path);paths.append(path)
        run=subprocess.run([shutil.which('libreoffice') or 'soffice',
             '-env:UserInstallation='+(tmp/'profile').as_uri(), '--headless','--convert-to','xlsx',
             '--outdir',str(out),*map(str,paths)],capture_output=True,text=True,timeout=180)
        if run.returncode:
            raise RuntimeError(run.stdout+run.stderr)
        for path in paths:
            wb=load_workbook(out/path.name,data_only=True)
            results[path.stem] = dict(values=[normalize(wb['RESULTADO']['A'+str(i)].value) for i in range(1,8)],
                prices={c:normalize(wb[fa.DADOS_AM][c].value) for c in selected_data(base)})
            wb.close()
    for name, _, _, expected in cases:
        v=results[name]['values'];total,neutral,nl,nf,warning,invalid=expected
        check(name+': total, saldo e contagens', v[:4]==[total,neutral,nl,nf],v[:4])
        check(name+': aviso de limite e cadastro',str(v[4]).startswith('⚠')==warning and bool(v[5])==invalid,v[4:6])
    check('custom/campanha vazios preservam todos os códigos e flags de preço antigos',
          results['vazio']['prices']==results['base_anterior']['prices'])
    check('resumo anterior idêntico quando nada foi configurado',results['vazio']['values'][4]==results['base_anterior']['values'][4],
          (results['vazio']['values'][4],results['base_anterior']['values'][4]))
    # Arithmetic overshoot remains visible instead of silently granting capacity.
    check('excesso dos limites gera aviso na configuração',bool(results['limites_excessivos']['values'][6]))
    print('CRIAÇÃO CONFERE' if all(checks) else 'CRIAÇÃO COM FALHAS',flush=True)
    return 0 if all(checks) else 1


if __name__=='__main__':
    sys.exit(main())
