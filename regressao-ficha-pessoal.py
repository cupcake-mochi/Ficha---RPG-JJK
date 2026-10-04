# -*- coding: utf-8 -*-
"""Regressao da FICHA PESSOAL: preenche casos na ficha GERADA, manda o LibreOffice recalcular e compara
com a regra, montada aqui a partir do catalogo.

Quem esta sendo julgada e a planilha: as formulas da aba, as contas dela na DADOS e as tres caixas da
FICHA que ela mexe (o EQUIPAMENTO, o XP e o DESLOCAMENTO). O Codigo.gs fica com o regressao-pessoal.js.

O que o LibreOffice nao tem e fica de fora: a SPARKLINE (as duas barras), que cai no IFERROR e sai vazia.
"""
import json, os, shutil, subprocess, sys, tempfile, itertools, math, re
from openpyxl import load_workbook

ARQ = "ficha-v01/ficha-projeto-m-0.1.xlsx"
if not os.path.exists(ARQ):
    print("gere a ficha antes:  python3 ficha-v01/monta.py"); sys.exit(1)
sys.path.insert(0, "ficha-v01")
import indice_ficha as ix, ficha_pessoal as fp

CAT = json.load(open("catalogo-projeto-m.json", encoding="utf-8"))
R = fp.regras(CAT)
G = fp.geometria(R)
ABA = fp.NOME
FALHAS = []


def checa(desc, cond, det=""):
    print(f"  [{'OK' if cond else 'FALHA'}] {desc}" + ("" if cond else f"  <- {det}"))
    if not cond:
        FALHAS.append(desc)


def indice_da_ficha(wb):
    dd, out = wb["DADOS"], {}
    for r in range(5, 200):
        k, v = dd.cell(row=r, column=53).value, ix.endereco(dd.cell(row=r, column=54).value)
        if k and v:
            out[k] = v
    return out


PASTA = tempfile.mkdtemp(prefix="pessoal-")
FILA = []


def prepara(preenche, nome):
    """uma copia da ficha, preenchida por `preenche(wb)`, na fila do LibreOffice"""
    copia = os.path.join(PASTA, nome + ".xlsx")
    shutil.copy(ARQ, copia)
    wb = load_workbook(copia)
    preenche(wb)
    # o IFS fica cru na ficha, porque ela vive no Sheets; o LibreOffice so o reconhece com o prefixo do Excel
    for ws in wb:
        for linha in ws.iter_rows():
            for c in linha:
                if isinstance(c.value, str) and c.value.startswith("=") and "IFS(" in c.value and "_xlfn.IFS(" not in c.value:
                    c.value = re.sub(r"(?<![A-Z_.])IFS\(", "_xlfn.IFS(", c.value)
    wb.save(copia)
    FILA.append(copia)


def recalcula_tudo():
    """o LibreOffice recalcula a fila inteira numa passada so; devolve {nome: pasta de trabalho so com valores}"""
    saida = os.path.join(PASTA, "saida")
    subprocess.run(["libreoffice", "--headless", "--convert-to", "xlsx", "--outdir", saida] + FILA,
                   capture_output=True, timeout=600)
    out = {}
    for copia in FILA:
        feito = os.path.join(saida, os.path.basename(copia))
        if not os.path.exists(feito):
            print("o LibreOffice nao converteu; sem ele esta checagem NAO roda"); sys.exit(1)
        out[os.path.basename(copia)[:-5]] = load_workbook(feito, data_only=True)
    shutil.rmtree(PASTA, ignore_errors=True)
    return out


def txt(v):
    """o valor como texto, com o ponto decimal do LibreOffice virando o da ficha e sem o .0 do inteiro"""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return "" if v is None else str(v).replace(".", ",") if isinstance(v, (int, float)) else str(v)


def num(s):
    return str(s).replace(".", ",")


WB0 = load_workbook(ARQ)
IDX = indice_da_ficha(WB0)
ATR = {n: IDX["atr_base_" + n] for n in ("Força", "Destreza", "Constituição", "Inteligência", "Essência")}
ARMAS = {a["nome"]: a for a in R["armas"] + R["escudos"]}
UNIF = {u["nome"]: u for u in R["uniformes"]}
E_INI = G["equip_ini"]


def monta(forca=0, destreza=0, principal=fp.SOCO, secundaria=fp.MAO_LIVRE, vestindo="Traje 1", grau="Grau 4",
          equip=(), itens=(), livres=(), treinos=(), grupos=(), caminho=None, nivel=2, missoes=()):
    def preenche(wb):
        f, p = wb["FICHA"], wb[ABA]
        f[ATR["Força"]], f[ATR["Destreza"]] = forca, destreza
        f[IDX["nivel"]] = nivel
        if caminho:
            f[IDX["caminho"]] = caminho
        p[G["principal"]], p[G["secundaria"]], p[G["vestindo"]], p[G["grau"]] = principal, secundaria, vestindo, grau
        for i, (nome, qtd) in enumerate(equip):
            p[f"D{E_INI + i}"] = nome
            if qtd is not None:
                p[f"K{E_INI + i}"] = qtd
        for i, linha in enumerate(livres):                       # a linha livre e digitada inteira
            lin = E_INI + fp.EQUIP_MENU + i
            for (rot, a, b), v in zip(fp.COLS_EQUIP, linha):
                if v is not None:
                    p[f"{a}{lin}"] = v
        for i, (nome, qtd) in enumerate(itens):
            cols = fp.COLS_ITENS[i // fp.ITENS_LINHAS]
            lin = G["itens_ini"] + i % fp.ITENS_LINHAS
            p[f"{cols[0][1]}{lin}"] = nome
            if qtd is not None:
                p[f"{cols[1][1]}{lin}"] = qtd
        for nome in treinos:
            col, lin, _ = G["armas"][nome]
            p.cell(row=lin, column=col).value = True
        for cat in grupos:
            col, lin, _ = G["grupos"][cat]
            p.cell(row=lin, column=col).value = True
        _n = G["missao_fim"] - G["missao_ini"] + 1
        for i, (tipo, mult, desc) in enumerate(missoes):               # enche um bloco e passa para o seguinte
            b = fp.BLOCOS[i // _n]
            lin = G["missao_ini"] + i % _n
            p.cell(row=lin, column=b).value = f"missão {i + 1}"
            p.cell(row=lin, column=b + fp.C_XP).value = tipo
            p.cell(row=lin, column=b + fp.C_ADIC).value = mult
            p.cell(row=lin, column=b + fp.C_DESC).value = desc
    return preenche


# o Kanabō e a Faca da ficha da Kaori do estudo, com a Kunai em dobro e o Broquel
KIT = [("Kanabō", 1), ("Faca", None), ("Kunai", 2), ("Broquel", 1)]
ITENS = [("Corda", 1), ("Lanterna", None), ("Kit de escalada", 1)]
LEVE = R["leve"]


def carga_do(equip, itens, vestindo, extra=0):
    v = sum(ARMAS[n]["volume"] * (q or 1) for n, q in equip) + sum(LEVE * (q or 1) for _, q in itens) + extra
    return round(v + (UNIF[vestindo]["volume"] if vestindo in UNIF else 0), 2)


# O que cada propriedade faz, para a nota da arma em uso: lido do arquivo que o extrair_equipamento.py tira do livro, e
# montado aqui de novo, sem passar pelo gerador. A ordem é a das duas tabelas do livro.
LIVRO = json.load(open("ficha-v01/equipamento-do-livro.json", encoding="utf-8"))
# a frase do livro sobre a coluna Mãos não serve na ficha: o gerador corta, e aqui também
DUAS_MAOS = next(p["faz"] for p in LIVRO["propriedades"] if p["nome"] == "Duas mãos").replace(" Aparece como 2 na coluna Mãos.", "")


def nota_da_arma(nome, a, principal):
    props = a["propriedades"].split(" · ") if a["propriedades"] else []
    linhas = [nome]
    for pr in LIVRO["propriedades"] + LIVRO["restricoes"]:
        if pr["nome"] == "Duas mãos":
            if not (principal and a["mao"] == 2):
                continue
            t = DUAS_MAOS
        elif pr["nome"] in props:
            t = " ".join([pr["faz"]] + (LIVRO["a_regra_da_secao"][pr["ver"]] if pr["ver"] else []))
        else:
            continue
        if pr["nome"] == "Longo Alcance" and a["alcance"]:
            t += f" Nesta arma: {a['alcance']}."
        if pr["nome"] == "Munição" and a["recarga"]:
            t += f" Nesta arma: {a['recarga']} {'ataque' if a['recarga'] == 1 else 'ataques'} por carga."
        if pr["nome"] == "Alcance":                 # decisão dele de 01/10/2026: toda arma com a propriedade chega a 3 m
            t += " Nesta arma: 3 m."
        linhas.append(f"{pr['nome']}: {t}")
    return "\n".join(linhas)


def cobertura():
    """as armas que, juntas, têm todas as propriedades do catálogo: a de mais propriedades novas primeiro"""
    falta = {x for a in CAT["equipamento"]["armas"].values() for x in a["propriedades"]} | {"Duas mãos"}
    tem = lambda n: set(CAT["equipamento"]["armas"][n]["propriedades"]) | ({"Duas mãos"} if CAT["equipamento"]["armas"][n]["mao"] == 2 else set())
    out = []
    while falta:
        n = max(sorted(CAT["equipamento"]["armas"]), key=lambda k: len(tem(k) & falta))
        out.append(n)
        falta -= tem(n)
    return out


MENU4 = [("Katana", 1), ("Faca", 1), ("Kanabō", 1), ("Broquel", 1)]
CASOS = {
    "de fábrica": dict(),
    "Kaori": dict(forca=3, destreza=2, principal="Faca", secundaria="Broquel", equip=KIT, itens=ITENS, treinos=["Faca"],
                  missoes=[("Padrão", None, None)]),
    "sem a Força": dict(forca=2, destreza=2, principal="Kanabō", secundaria="Broquel", vestindo="Revestimento 1",
                        equip=KIT, itens=ITENS, grupos=["Massa"]),
    "carga estourada": dict(forca=3, principal="Faca", secundaria="Broquel", equip=KIT, itens=ITENS,
                            livres=[("Cofre", 1, "Carga", None, None, None, None, 9, None, None, None)]),
    "arremesso": dict(forca=3, principal="Kunai", equip=KIT),
    "Versátil": dict(forca=3, destreza=3, principal="Katana", secundaria=fp.NAS_DUAS, equip=MENU4, treinos=["Katana"],
                     grau="Grau 3", vestindo="Traje 2", nivel=7, missoes=[("Final de arco", "2x", None)] * 3),
    "arma de fogo": dict(forca=3, destreza=4, principal="Pistola", secundaria="Médio", vestindo="Revestimento 1",
                         equip=[("Pistola", 1), ("Médio", 1)]),
    "Grau baixo": dict(forca=4, vestindo="Revestimento 2", grau="Grau 4", nivel=20),
    "Grau 3": dict(forca=4, vestindo="Revestimento 2", grau="Grau 3", nivel=10),
    "escudo sem a Força": dict(forca=1, destreza=3, principal="Faca", secundaria="Torre", vestindo=fp.SEM_UNIFORME,
                               equip=[("Faca", 1), ("Torre", 1)]),
    "linha livre": dict(forca=3, principal="Lança-foguete", secundaria="Faca", equip=[("Faca", 1)], grupos=["Lâmina Curta"],
                        livres=[("Lança-foguete", 1, "Massa", 1, "3d6", "Rompe", 5, 3, None, None, None)], nivel=30),
    "duas mãos": dict(forca=6, destreza=2, principal="Kanabō", vestindo="Traje 3", equip=MENU4[1:]),
}


COBRE = cobertura()
_uma = [n for n in COBRE if ARMAS[n]["mao"] == 1] + ["Torre"]
for _i, _n in enumerate(COBRE):
    _s = _uma[_i % len(_uma)] if ARMAS[_n]["mao"] == 1 else fp.MAO_LIVRE
    CASOS[f"nota · {_n}"] = dict(forca=6, principal=_n, secundaria=_s, equip=[(x, 1) for x in dict.fromkeys([_n, _s]) if x in ARMAS])
# a arma digitada na mão sem estar guardada: a linha de baixo avisa, e a nota fica vazia
CASOS["não guardada"] = dict(forca=3, principal="Katana", secundaria="Faca")
# o aviso de cada mão conta só o que serve a ela: o escudo não entra na principal, e a arma de duas mãos não entra na secundária
CASOS["só o escudo guardado"] = dict(forca=3, secundaria="Broquel", equip=[("Broquel", 1)])
CASOS["só arma de duas mãos guardada"] = dict(forca=6, equip=[("Kanabō", 1)])


def esperado(c):
    forca, pri, sec, ves = c.get("forca", 0), c.get("principal", fp.SOCO), c.get("secundaria", fp.MAO_LIVRE), c.get("vestindo", "Traje 1")
    grau, equip, itens, livres = c.get("grau", "Grau 4"), c.get("equip", ()), c.get("itens", ()), c.get("livres", ())
    nivel = c.get("nivel", 2)
    maestria = 1 + sum(1 for m in (10, 18, 26) if nivel >= m)
    tab = {n: dict(ARMAS[n]) for n, _ in equip}
    for l in livres:
        tab[l[0]] = {"categoria": l[2], "mao": l[3] or 0, "dado": l[4] or "", "propriedades": l[5] or "", "forca": l[6] or "—",
                     "volume": l[7] or 0, "alcance": None, "recarga": None, "livre": True}
    out = {}

    def forca_de(a):
        return 0 if a["forca"] == "—" else a["forca"]

    def marca(pede):
        if pede == 0:
            return "Sem requisito de Força"
        return f"{fp.T_FALTA_FORCA}: {forca} de {pede}" if pede > forca else f"Força {forca} de {pede}"

    def detalhe(a, dado, extra=""):
        return dado + (f" · {a['propriedades']}" if a["propriedades"] else "") + extra + \
            (f" · {a['alcance']}" if a["alcance"] else "") + (f" · recarga a cada {a['recarga']}" if a["recarga"] else "")
    # B30: sem equipável guardado que sirva à mão, a linha de baixo diz como pôr um no menu
    serve_p = [n for n, a in tab.items() if a["categoria"] != "Escudo"]
    serve_s = [n for n, a in tab.items() if a["categoria"] == "Escudo" or a["mao"] == 1]
    if pri == fp.SOCO:
        mao_p, falta_p = 1, 0
        out["det_principal"] = f"{R['soco'][maestria - 1]} · " + ("sem propriedade · o dado sobe com a maestria" if serve_p else
                                                                 "para outra arma, guarde ela nos Equipáveis guardados")
        out["marca_treino"], out["marca_forca_p"] = "Treinada", "Sem requisito de Força"
        versatil = False
    elif pri not in tab:
        mao_p, falta_p, versatil = 0, 0, False
        out["det_principal"] = "Não está nos equipáveis guardados"
        out["marca_treino"] = "Treinada" if pri in c.get("treinos", ()) else f"{fp.T_SEM_TREINO}: desvantagem"
        out["marca_forca_p"] = "Sem requisito de Força"
    else:
        a = tab[pri]
        mao_p, falta_p = a["mao"], forca_de(a) > forca
        versatil = "Versátil" in a["propriedades"]
        duas = versatil and sec == fp.NAS_DUAS
        dado = {"d6": "d8", "d8": "d10", "d10": "d12"}.get(a["dado"], a["dado"]) if duas else a["dado"]
        out["det_principal"] = detalhe(a, dado, " · nas duas mãos" if duas else " · duas mãos" if mao_p == 2 else "")
        treinada = (pri in c.get("treinos", ())) if not a.get("livre") else (a["categoria"] in c.get("grupos", ()))
        if not a.get("livre") and a["categoria"] in c.get("grupos", ()) and pri not in c.get("treinos", ()):
            treinada = False                    # a caixa da arma manda: a do grupo so vale para a arma digitada
        out["marca_treino"] = "Treinada" if treinada else f"{fp.T_SEM_TREINO}: desvantagem"
        out["marca_forca_p"] = marca(forca_de(a))
    vale = mao_p != 2 and sec not in ("", fp.MAO_LIVRE, fp.NAS_DUAS)
    guardada = vale and sec in tab
    escudo = guardada and tab[sec]["categoria"] == "Escudo"
    falta_s = guardada and forca_de(tab[sec]) > forca
    if mao_p == 2:
        out["det_secundaria"] = "Ocupada: a arma da outra mão é de duas mãos"
    elif sec == fp.NAS_DUAS:
        out["det_secundaria"] = "A mesma arma, nas duas mãos" if versatil else "A arma da outra mão não é Versátil"
    elif not vale:
        out["det_secundaria"] = "Mão livre" if serve_s else "Mão livre · arma de uma mão ou escudo guardado aparece aqui"
    elif not guardada:
        out["det_secundaria"] = "Não está nos equipáveis guardados"
    elif escudo:
        out["det_secundaria"] = "Escudo · " + tab[sec]["propriedades"]
    else:
        out["det_secundaria"] = detalhe(tab[sec], tab[sec]["dado"])
    out["marca_forca_s"] = (marca(forca_de(tab[sec])) if guardada else "Sem requisito de Força") if vale else "—"
    vale = guardada                                           # daqui para baixo, o que não está guardado não conta
    out["marca_selo"] = "Trava Selo de gesto" if escudo else "Sem escudo"
    u = UNIF.get(ves)
    falta_u = bool(u) and u["forca"] > forca
    ordem = [p[0] for p in R["patentes"]].index(grau) + 1
    if u:
        out["det_vestindo"] = f"Proteção {u['protecao']} · " + ("sem teto de Destreza" if u["teto"] == "—" else f"teto de Destreza {u['teto']}")
        out["marca_grau"] = f"Liberado no {grau}" if ordem >= u["ordem"] else f"{fp.T_PEDE_GRAU}{u['grau']}"
        out["marca_forca_v"] = marca(u["forca"])
    else:
        out["det_vestindo"], out["marca_grau"], out["marca_forca_v"] = "Sem uniforme: vale a proteção do cobrir-se", "—", "—"
    out["nota:arma da principal"] = (f"{fp.SOCO}\n{LIVRO['soco']}" if pri == fp.SOCO else
                                     nota_da_arma(pri, tab[pri], True) if pri in tab else "")
    out["nota:arma da secundária"] = ("" if not vale else f"{sec}\n{LIVRO['escudo']}" if escudo else nota_da_arma(sec, tab[sec], False))
    carga = carga_do(equip, itens, ves, sum(l[7] or 0 for l in livres))
    limite = R["limite_base"] + forca
    out["carga"] = f"{num(carga if carga % 1 else int(carga))} de {limite}" + (fp.T_ACIMA if carga > limite else "")
    out["salario"] = dict((p[0], p[1]) for p in R["patentes"])[grau]
    out["requisito"] = fp.T_NAO_CUMPRIDO if (falta_p or falta_s or falta_u) else "Cumprido"
    out["__equipamento"] = " + ".join(x for x in (ves if u else "", sec if escudo else "") if x)
    # 04/10/2026, o livro reconstruído: arma empunhada sem a Força corta o deslocamento pela metade; carga acima do limite
    # não deixa andar. Uniforme e escudo sem a Força só perdem a proteção (a Defesa, mais abaixo)
    meia = falta_p or (falta_s and not escudo)
    out["__deslocamento"] = "0 m" if carga > limite else "4,5 m" if meia else "9 m"
    # a Defesa da FICHA continua a mesma conta, agora pelo espelho: 10 + Destreza (com o teto) + proteção
    return out


tipos, mults, descs = dict(R["tipos"]), dict(R["multiplicadores"]), dict(R["descontos"])
combos = list(itertools.product(tipos, [None] + list(mults), [None] + list(descs)))
por_bloco = G["missao_fim"] - G["missao_ini"] + 1
_cabem = len(fp.BLOCOS) * por_bloco                                  # as quatro tabelas: as duas do painel e as duas da extensão
LOTES = [combos[i:i + _cabem] for i in range(0, len(combos), _cabem)]


def total(tipo, mult, desc):
    # o livro: arredonda para baixo só o XP final, e o positivo menor que 1 vira 1
    bruto = round(tipos[tipo] * (mults[mult] if mult else 1) * (descs[desc] if desc else 1), 4)
    return max(1, math.floor(bruto)) if bruto > 0 else 0


for nome, c in CASOS.items():
    prepara(monta(**c), "caso-" + nome)
for n, lote in enumerate(LOTES):
    prepara(monta(missoes=lote), f"xp-{n}")
LIDO = recalcula_tudo()
_ED = CAT["equipamento_defesa"]
acum, ac = {}, 0
for n in sorted(int(x) for x in R["niveis"] if int(x) >= 2):
    acum[n] = ac
    custo = str(R["niveis"][str(n)]["xp"])
    ac += int(custo.replace(".", "")) if custo != "—" else 0


def coluna(d, titulo):
    """a coluna da DADOS sob `titulo`, sem as células vazias (que o menu suspenso não mostra)"""
    for linha in d.iter_rows(min_row=1, max_row=8):
        for c in linha:
            if c.value == titulo:
                return [v for v in (d.cell(row=r, column=c.column).value for r in range(c.row + 1, c.row + 40)) if v not in (None, "")]
    return None


def notas_vivas(d):
    """as notas que mudam com a ficha: {nome: (texto, caixa)}, da tabela `nota viva` da DADOS"""
    for linha in d.iter_rows(min_row=1, max_row=8):
        for c in linha:
            if c.value == "nota viva":
                out = {}
                for r in range(c.row + 1, c.row + 20):
                    if d.cell(row=r, column=c.column).value:
                        out[d.cell(row=r, column=c.column).value] = (d.cell(row=r, column=c.column + 1).value, d.cell(row=r, column=c.column + 2).value)
                return out
    return {}


print("O QUE A ABA LÊ DO LIVRO")
_r = subprocess.run([sys.executable, "ficha-v01/extrair_equipamento.py", "--confere"], capture_output=True, text=True)
if "nao esta nesta maquina" in _r.stdout:
    print("  [--] o livro não está nesta máquina: o equipamento-do-livro.json não foi comparado com ele")
else:
    checa(f"o equipamento-do-livro.json é o que o extrair_equipamento.py lê do livro hoje ({LIVRO['_meta']['versao_do_livro']})",
          _r.returncode == 0, _r.stdout.strip()[-200:])
_duas = next(pr for pr in LIVRO["propriedades"] if pr["nome"] == "Duas mãos")
checa("o texto das Duas mãos daqui abre a frase do livro", _duas["faz"].startswith(DUAS_MAOS), _duas["faz"])
checa("o livro escreve 1,5 m para a arma de mão e 3 m para a que chega mais longe",
      LIVRO["alcance_no_corpo_a_corpo"] == {"padrao": "1,5 m", "o_que_chega_mais_longe": "3 m"}, str(LIVRO["alcance_no_corpo_a_corpo"]))
_usadas = {x for a in CAT["equipamento"]["armas"].values() for x in a["propriedades"]}
_com_texto = {pr["nome"] for pr in LIVRO["propriedades"] + LIVRO["restricoes"]}
checa(f"as {len(_usadas)} propriedades que as 52 armas do catálogo usam têm texto no livro", _usadas <= _com_texto, str(sorted(_usadas - _com_texto)))
_cobertas = {x for n in COBRE for x in CAT["equipamento"]["armas"][n]["propriedades"]}
checa(f"os casos de nota ({', '.join(COBRE)}) passam por todas elas, e por arma de duas mãos",
      _cobertas == _usadas and any(ARMAS[n]["mao"] == 2 for n in COBRE), str(sorted(_usadas - _cobertas)))

print("\nAS CAIXAS DA ABA, CASO A CASO")
for nome, c in CASOS.items():
    wb = LIDO["caso-" + nome]
    p, f, d = wb[ABA], wb["FICHA"], wb["DADOS"]
    esp = esperado(c)
    erros = []
    vivas = notas_vivas(d)
    for k, v in esp.items():
        lido = (txt(vivas.get(k[5:], ("<a DADOS não tem esta nota>",))[0]) if k.startswith("nota:") else
                txt(f[IDX[k[2:]]].value) if k.startswith("__") else txt(p[G[k]].value))
        if k in ("carga", "__deslocamento"):
            lido = lido.replace(".", ",")          # o LibreOffice junta o número com o ponto ou a vírgula do sistema
        if lido != v:
            erros.append(f"{k}: a ficha diz {lido!r}, a regra diz {v!r}")
    checa(f"{nome}: as {len(esp)} caixas e notas batem com a regra", not erros, " · ".join(erros))
_cx = notas_vivas(LIDO["caso-de fábrica"]["DADOS"])
checa("a nota da arma mora na linha embaixo de cada mão",
      (_cx.get("arma da principal", (0, 0))[1], _cx.get("arma da secundária", (0, 0))[1]) == (G["det_principal"], G["det_secundaria"]), str(_cx))

print("\nA DEFESA DA FICHA, PELO ESPELHO DO EQUIPAMENTO")
for nome, c in CASOS.items():
    ves, sec, des = c.get("vestindo", "Traje 1"), c.get("secundaria", fp.MAO_LIVRE), c.get("destreza", 0)
    duas = c.get("principal") in ARMAS and ARMAS[c["principal"]]["mao"] == 2
    u, e = _ED["uniformes"].get(ves), (None if duas else _ED["escudos"].get(sec))
    tetos = [x["teto_de_destreza"] for x in (u, e) if x and x["teto_de_destreza"] is not None]
    prot = (u["protecao"] if u else 1) + (e["protecao"] if e else 0)       # sem uniforme, o cobrir-se do refino 1: 1/3 + 1
    # 04/10/2026: a peça usada sem a Força não dá a proteção dela, e a arma empunhada sem a Força tira a Destreza
    fc = c.get("forca", 0)
    req = lambda x: x["requer_forca"] or 0
    prot -= (u["protecao"] if u and req(u) > fc else 0) + (e["protecao"] if e and req(e) > fc else 0)
    armas_em_uso = [n for n in (c.get("principal"), None if e else sec) if n in ARMAS and ARMAS[n]["categoria"] != "Escudo"]
    sem_forca = any((0 if ARMAS[n]["forca"] == "—" else ARMAS[n]["forca"]) > fc for n in armas_em_uso)
    esp = 10 + (0 if sem_forca else min([des] + tetos)) + prot
    lido = LIDO["caso-" + nome]["FICHA"][IDX["defesa"]].value
    checa(f"{nome}: {ves}, {sec if e else 'sem escudo'}, Destreza {des}: Defesa {esp}", lido == esp, f"a ficha diz {lido}")

print("\nOS MENUS QUE MUDAM COM A FICHA")
d = LIDO["caso-Versátil"]["DADOS"]
checa("a mão principal lista o Soco e as armas guardadas, sem o escudo",
      coluna(d, "menu da mão principal") == [fp.SOCO, "Katana", "Faca", "Kanabō"], str(coluna(d, "menu da mão principal")))
checa("com arma Versátil na principal, a secundária oferece as duas mãos, as armas de uma mão e o escudo",
      coluna(d, "menu da mão secundária") == [fp.MAO_LIVRE, fp.NAS_DUAS, "Katana", "Faca", "Broquel"], str(coluna(d, "menu da mão secundária")))
_g3 = [u["nome"] for u in R["uniformes"] if u["ordem"] <= 2]
checa(f"no Grau 3 o menu de vestir libera {len(_g3)} uniformes e segura o resto",
      coluna(d, "menu do vestindo") == [fp.SEM_UNIFORME] + _g3, str(coluna(d, "menu do vestindo")))
d = LIDO["caso-duas mãos"]["DADOS"]
checa("com arma de duas mãos na principal, a secundária só oferece a mão livre",
      coluna(d, "menu da mão secundária") == [fp.MAO_LIVRE], str(coluna(d, "menu da mão secundária")))
checa("no Grau 4 o menu de vestir não tem o Revestimento 2 nem o 3",
      coluna(d, "menu do vestindo") == [fp.SEM_UNIFORME] + [u["nome"] for u in R["uniformes"] if u["ordem"] == 1],
      str(coluna(d, "menu do vestindo")))

print("\nO XP DAS MISSÕES")
errados, soma_ok = [], True
for n, lote in enumerate(LOTES):
    wb = LIDO[f"xp-{n}"]
    p = wb[ABA]
    soma = 0
    for i, (t, m, de) in enumerate(lote):
        b = fp.BLOCOS[i // por_bloco]
        lin = G["missao_ini"] + i % por_bloco
        lido, esp = p.cell(row=lin, column=b + fp.C_TOTAL).value, total(t, m, de)
        soma += esp
        if lido != esp:
            errados.append(f"{t} {m} {de}: {lido} != {esp}")
    xp_ficha, xp_painel = wb["FICHA"][IDX["xp"]].value, p[G["xp_total"]].value
    if not (xp_ficha == xp_painel == soma):
        soma_ok = False
        errados.append(f"lote {n}: XP da FICHA {xp_ficha}, do painel {xp_painel}, a regra {soma}")
checa(f"as {len(combos)} combinações de tipo, adicional e desconto dão o Total da regra (para baixo no fim, e no mínimo 1)",
      not errados, " · ".join(errados[:6]))
checa("o XP total do painel e o XP da FICHA são a soma da coluna Total das quatro tabelas, as duas da extensão inclusive",
      soma_ok and len(combos) > 2 * por_bloco, f"{len(combos)} missões para {por_bloco} linhas por tabela")
# os exemplos do livro: "Uma quinta missão longa paga 200 ÷ 8 = 25 XP, e não 24", e a coluna da missão padrão da tabela
checa("os exemplos do livro: a quinta longa paga 25, a padrão paga 12 na quinta e 6 na sexta, e o positivo pequeno paga 1",
      total("Longa", None, "5ª · 12,5%") == 25 and total("Padrão", None, "5ª · 12,5%") == 12 and total("Padrão", None, "6ª · 6,25%") == 6
      and total("Padrão", None, "7ª · 3,125%") == 3 and total("Curta", None, "7ª · 3,125%") == 1)

print("\nO QUE FALTA PARA O PRÓXIMO NÍVEL")
for nome, c in CASOS.items():
    nivel, xp = c.get("nivel", 2), sum(total(*m) for m in c.get("missoes", ()))
    p = LIDO["caso-" + nome][ABA]
    rot = "NÍVEL MÁXIMO" if nivel >= 30 else f"FALTA PARA O NÍVEL {nivel + 1}"
    falta = "—" if nivel >= 30 else txt(max(0, acum[nivel + 1] - xp))
    lido_r, lido_f = p[G["falta_rot"]].value, txt(p[G["falta"]].value)
    setas = [c_.value for linha in p.iter_rows(min_row=G["niveis_ini"], max_row=G["niveis_ini"] + 15, min_col=fp.B1) for c_ in linha
             if isinstance(c_.value, str) and c_.value.startswith("▸")]
    checa(f"{nome}: nível {nivel} com {txt(xp)} de XP: {rot.lower()}, {falta}, e a seta no {nivel}",
          lido_r == rot and lido_f == falta and setas == [f"▸ {nivel}"], f"{lido_r!r}, {lido_f!r}, setas {setas}")

print()
print("=" * 70)
if FALHAS:
    print(f">>> {len(FALHAS)} FALHA(S) NA FICHA PESSOAL")
    sys.exit(1)
print(">>> A FICHA PESSOAL CALCULA O QUE A REGRA MANDA")
