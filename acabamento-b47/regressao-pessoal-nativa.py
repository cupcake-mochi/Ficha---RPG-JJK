"""Reutiliza cenários e oráculos originais; só o recálculo é feito no Sheets.

preparar emite entradas/faixas, sem rede. conferir relê respostas CellData
coletadas pelo integrador e executa as asserções originais, sem Calc.
Os recibos/valores da planilha permanecem no diretório privado informado.
"""
import ast
import hashlib
import json
import os
from pathlib import Path
import re
import sys
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
os.chdir(ROOT)
SOURCE = ROOT / 'regressao-ficha-pessoal.py'


def contexto(baseline=None):
    source = SOURCE.read_text(encoding='utf-8')
    tree = ast.parse(source)
    stop = next(n.lineno for n in tree.body if isinstance(n, ast.For)
                and isinstance(n.iter, ast.Call) and ast.unparse(n.iter) == 'CASOS.items()')
    env = {'__file__': str(SOURCE), '__name__': 'regressao_ficha_pessoal_nativa'}
    prefix = '\n'.join(source.splitlines()[:stop-1])
    if baseline is not None and baseline.exists():
        # Usa o estado nativo atual para coordenadas/constantes, sem confiar
        # num XLSX histórico nem importar arquivo Office na planilha.
        wb = Workbook(); wb.remove(wb.active)
        obj = json.loads(baseline.read_text(encoding='utf-8'))
        for sheet in obj['sheets']:
            ws = wb.create_sheet(sheet['properties']['title'])
            for block in sheet.get('data', []):
                for row, rd in enumerate(block.get('rowData', []), block.get('startRow', 0)+1):
                    for col, cell in enumerate(rd.get('values', []), block.get('startColumn', 0)+1):
                        entered = cell.get('userEnteredValue', {})
                        if entered:
                            target = ws.cell(row, col)
                            target.value = next(iter(entered.values()))
                            if 'stringValue' in entered: target.data_type = 's'
        env['_baseline_nativa'] = wb
        prefix = prefix.replace('if not os.path.exists(ARQ):', 'if False:  # baseline nativa disponível')
        prefix = prefix.replace('WB0 = load_workbook(ARQ)', 'WB0 = _baseline_nativa')
    exec(compile(prefix, str(SOURCE), 'exec'), env)
    return source, stop, env


class Entradas:
    def __init__(self):
        self.cells = {}
    def __getitem__(self, name):
        outer = self
        class Aba:
            def __setitem__(self, address, value):
                outer.cells[(name, address)] = value
            def cell(self, row, column):
                class Cel:
                    @property
                    def value(self):
                        return outer.cells.get((name, get_column_letter(column)+str(row)))
                    @value.setter
                    def value(self, value):
                        outer.cells[(name, get_column_letter(column)+str(row))] = value
                return Cel()
        return Aba()


def ranges(cells):
    grouped = {}
    for sheet, coord in cells:
        m = re.fullmatch(r'([A-Z]+)(\d+)', coord)
        col = 0
        for c in m[1]: col = col*26 + ord(c)-64
        grouped.setdefault((sheet, int(m[2])), set()).add(col)
    result = []
    for (sheet, row), cols in sorted(grouped.items()):
        cols = sorted(cols)
        a = b = cols[0]
        for c in cols[1:]+[10**6]:
            if c-b <= 3:
                b = c
            else:
                result.append(f"'{sheet}'!{get_column_letter(a)}{row}:{get_column_letter(b)}{row}")
                a = b = c
    return result


def preparar(folder):
    source, stop, e = contexto(folder/'baseline.json')
    scenarios = []
    for name, args in list(e['CASOS'].items()) + [(f'xp-{i}', {'missoes': lote}) for i,lote in enumerate(e['LOTES'])]:
        inputs = Entradas()
        e['monta'](**args)(inputs)
        scenarios.append({'nome': name if name.startswith('xp-') else 'caso-'+name,
                          'entradas': [[s,c,v] for (s,c),v in inputs.cells.items()]})
    all_inputs = {(s,c) for x in scenarios for s,c,v in x['entradas']}
    wb = e['WB0']
    outputs = {(ws.title,c.coordinate) for ws in wb if ws.title in ('FICHA','FICHA PESSOAL','DADOS')
               for c in ws._cells.values() if c.data_type == 'f'}
    # FILTER pode derramar em células vazias: os mesmos menus lidos pelo teste.
    for row in wb['DADOS'].iter_rows(min_row=1,max_row=8):
        for c in row:
            if c.value in ('menu da mão principal','menu da mão secundária','menu do vestindo'):
                outputs.update(('DADOS', f'{c.column_letter}{r}') for r in range(c.row+1,c.row+40))
            if c.value == 'nota viva':
                outputs.update(('DADOS',f'{get_column_letter(k)}{r}') for r in range(c.row+1,c.row+20)
                               for k in range(c.column,c.column+3))
    plan = {'fonte': str(SOURCE.name), 'sha256_teste': hashlib.sha256(source.encode()).hexdigest(),
            'cenarios': scenarios, 'entradas': [[s,c,wb[s][c].value] for s,c in sorted(all_inputs)],
            'faixas_entrada': ranges(all_inputs), 'faixas_saida': ranges(outputs)}
    folder.mkdir(parents=True,exist_ok=True)
    (folder/'plano-pessoal.json').write_text(json.dumps(plan,ensure_ascii=False),encoding='utf-8')
    print(json.dumps({'cenarios':len(scenarios),'entradas':len(all_inputs),
                      'faixas_saida':len(plan['faixas_saida'])},ensure_ascii=False))


def conferir(folder):
    source, stop, e = contexto(folder/'baseline.json')
    if '--sem-salario-legado' in sys.argv:
        # A extensão salarial substituiu este campo por número/rotina de
        # eventos. Escritas via API não disparam onEdit; não certificar esse
        # comportamento com o oráculo antigo que esperava texto imediato.
        original = e['esperado']
        e['esperado'] = lambda case: {k:v for k,v in original(case).items() if k != 'salario'}
        print('FORA DO ESCOPO: salário legado (requer ensaio da rotina atual).')
    plan = json.loads((folder/'plano-pessoal.json').read_text(encoding='utf-8'))
    assert hashlib.sha256(source.encode()).hexdigest() == plan['sha256_teste'], 'teste mudou'
    lidos = {}
    for i,case in enumerate(plan['cenarios']):
        wb = Workbook(); wb.remove(wb.active)
        # Constantes vêm do artefato gerado; toda fórmula é substituída pela
        # resposta nativa. Nunca se usa cache do XLSX como resultado calculado.
        for old in e['WB0']:
            if old.title not in ('FICHA','FICHA PESSOAL','DADOS'): continue
            ws=wb.create_sheet(old.title)
            for cell in old._cells.values():
                if cell.data_type != 'f' and cell.value is not None: ws[cell.coordinate]=cell.value
        for s,c,v in case['entradas']: wb[s][c]=v
        response=json.loads((folder/f'pessoal-{i:02}.json').read_text(encoding='utf-8'))
        obj=response.get('structuredContent',response)
        assert not response.get('isError'), response
        for sheet in obj['sheets']:
            for block in sheet.get('data',[]):
                for row,rd in enumerate(block.get('rowData',[]),block.get('startRow',0)+1):
                    for col,cell in enumerate(rd.get('values',[]),block.get('startColumn',0)+1):
                        value=cell.get('effectiveValue',{})
                        val=value.get('numberValue',value.get('boolValue',value.get('stringValue')))
                        if 'errorValue' in value: val=value['errorValue']['type']
                        wb[sheet['properties']['title']].cell(row,col).value=val
        lidos[case['nome']]=wb
    # O restante é literalmente o teste original, incluindo notas, menus, XP,
    # Defesa, carga e leitura independente do livro. Só a preparação é evitada.
    e['prepara']=lambda *args: None
    e['recalcula_tudo']=lambda: lidos
    exec(compile('\n'.join(source.splitlines()[stop-1:]),str(SOURCE),'exec'),e)


if __name__ == '__main__':
    {'preparar':preparar,'conferir':conferir}[sys.argv[1]](Path(sys.argv[2]))
