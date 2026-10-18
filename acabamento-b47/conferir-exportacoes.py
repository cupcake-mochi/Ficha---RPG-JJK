"""Confere exportações XLSX somente-leitura da revisão, sem rede e sem escrever no Google.

Uso quando a sessão só alcança o Drive (exportar), não a API Sheets nem o Apps Script.
As exportações ficam fora do Git: têm os dados do personagem. O resumo gravado aqui
leva somente contagens, hashes e endereços estruturais.

XLSX aproxima o modelo nativo: não traz sheetId, regras condicionais, proteções nem
acionadores, e "sem formatação explícita" aparece como Arial 10 preto sem preenchimento.
Serve de fotografia e de plano; aplicar no Google exige releitura nativa e guarda.
"""
import argparse
import hashlib
import json
from pathlib import Path

import openpyxl
from openpyxl.worksheet.formula import ArrayFormula

STUDIES = {'ESTUDO CAMPANHA', 'ESTUDO CUSTOM', 'ESTUDO 3B', 'ESTUDO 3C'}
FAIXA = (1733, 2812, 1, 16)  # CONFIGURAÇÃO!A1733:P2812
SEM_FORMATO = (None, 'Arial', 10.0, 'FF000000')
BASE = ('FF1B111D', 'Roboto', 10.0, 'FFFBE4EF')  # Mizuki Escuro, formato-base canônico


def L(x):
    return openpyxl.utils.get_column_letter(x)


def texto(v):
    if isinstance(v, ArrayFormula):
        return '=' + str(v.text).lstrip('=') + ' @' + str(v.ref)
    return str(v)


def fundo(c):
    f, fo = c.fill, c.font
    return (f.fgColor.rgb if f and f.fill_type else None, fo.name, fo.sz,
            fo.color.rgb if fo.color is not None else None)


def estilo(c):
    a, b = c.alignment, c.border
    return fundo(c) + (c.font.b, c.font.i, c.number_format, a.horizontal, a.vertical, a.wrap_text,
                       *[getattr(getattr(b, k), 'style', None) for k in ('left', 'right', 'top', 'bottom')])


def foto(path):
    wf = openpyxl.load_workbook(path)
    wv = openpyxl.load_workbook(path, data_only=True)
    abas = {}
    for nome in wf.sheetnames:
        a, v = wf[nome], wv[nome]
        F = {c.coordinate: texto(c.value) for row in a.iter_rows() for c in row if c.value is not None}
        V = {c.coordinate: str(c.value) for row in v.iter_rows() for c in row if c.value is not None}
        N = {c.coordinate: c.comment.text for row in a.iter_rows() for c in row if c.comment}
        abas[nome] = {
            'F': F, 'V': V, 'N': N,
            'dv': sorted(str(r) for x in a.data_validations.dataValidation for r in x.sqref.ranges),
            'merged': sorted(map(str, a.merged_cells.ranges)),
            'rowgroups': sorted((r, d.outlineLevel) for r, d in a.row_dimensions.items() if d.outlineLevel),
            'colgroups': sorted((k, d.outlineLevel) for k, d in a.column_dimensions.items() if d.outlineLevel),
        }
    return wf, {'abas': wf.sheetnames, 'nomes': sorted(wf.defined_names.keys()), 'sheets': abas}


def resumo(d):
    s = d['sheets'].values()
    analise = d['sheets']['ANALISE']['V']
    chaves = [v for c, v in analise.items() if c[:1] == 'A' and c[1:].isdigit() and int(c[1:]) > 1]
    erros = [x for a in s for x in a['V'].values() if x[:1] == '#' and (x.endswith('!') or x.endswith('?') or '/' in x or x == '#N/A')]
    return {'abas': len(d['abas']), 'formulas': sum(x.startswith('=') for a in s for x in a['F'].values()),
            'celulas': sum(len(a['F']) for a in s), 'notas': sum(len(a['N']) for a in s),
            'validacoes': sum(len(a['dv']) for a in s), 'intervalos_nomeados': len(d['nomes']),
            'erros_formula': len(erros), 'analise_chaves': len(chaves), 'analise_unicas': len(set(chaves))}


def diff(a, b, campos=('F', 'V', 'N', 'dv', 'merged', 'rowgroups', 'colgroups')):
    out = {}
    for nome in b['abas']:
        for k in campos:
            x, y = a['sheets'][nome][k], b['sheets'][nome][k]
            if x == y:
                continue
            if isinstance(x, dict):
                out.setdefault(nome, {})[k] = sorted(c for c in set(x) | set(y) if x.get(c) != y.get(c))
            else:
                out.setdefault(nome, {})[k] = 'diferente'
    return out


def diff_estilo(wa, wb):
    out = {}
    # style_id só tem sentido dentro do próprio workbook. Resolver a tupla
    # semântica uma vez por estilo mantém o mesmo oráculo e evita milhões de
    # acessos a fontes/bordas em exportações com quase meio milhão de células.
    def resolve():
        memo = {}
        def get(cell):
            key = cell.style_id
            if key not in memo:
                memo[key] = estilo(cell)
            return memo[key]
        return get
    estilo_a, estilo_b = resolve(), resolve()
    for nome in wa.sheetnames:
        a, b = wa[nome], wb[nome]
        # max_column percorre todas as células do openpyxl. Dentro do laço
        # de linhas isso era quadrático, sobretudo nas 8.193 linhas de ANALISE.
        max_row, max_col = max(a.max_row, b.max_row), max(a.max_column, b.max_column)
        cells = [L(x) + str(y) for y in range(1, max_row + 1)
                 for x in range(1, max_col + 1)
                 if estilo_a(a.cell(y, x)) != estilo_b(b.cell(y, x))]
        if cells:
            out[nome] = {'celulas': len(cells), 'primeira': cells[0], 'ultima': cells[-1]}
    return out


def plano_fundo(wr, wd):
    """Células vazias e sem formato na revisão que a descartável corrigida tem no formato-base."""
    r, d = wr['CONFIGURAÇÃO'], wd['CONFIGURAÇÃO']
    y0, y1, x0, x1 = FAIXA
    sel, outras = set(), []
    for y in range(y0, y1 + 1):
        for x in range(x0, x1 + 1):
            a, b = r.cell(y, x), d.cell(y, x)
            if fundo(a) == SEM_FORMATO and fundo(b) == BASE:
                assert a.value in (None, ''), f'{L(x)}{y} tem valor'
                sel.add((y, x))
            elif estilo(a) != estilo(b):
                outras.append(L(x) + str(y))
    spans = {}
    for y in range(y0, y1 + 1):
        x = x0
        while x <= x1:
            if (y, x) not in sel:
                x += 1
                continue
            ini = x
            while x <= x1 and (y, x) in sel:
                x += 1
            spans.setdefault((ini, x - 1), []).append(y)
    faixas = []
    for (xa, xb), rows in sorted(spans.items()):
        ini = ant = rows[0]
        for y in rows[1:] + [None]:
            if y is not None and y == ant + 1:
                ant = y
                continue
            faixas.append(f'{L(xa)}{ini}:{L(xb)}{ant}')
            ini = ant = y
    coberto = set()
    for f in faixas:
        (c0, r0), (c1, r1) = (openpyxl.utils.cell.coordinate_from_string(p) for p in f.split(':'))
        c0, c1 = openpyxl.utils.column_index_from_string(c0), openpyxl.utils.column_index_from_string(c1)
        coberto |= {(y, x) for y in range(r0, r1 + 1) for x in range(c0, c1 + 1)}
    assert coberto == sel
    return {'aba': 'CONFIGURAÇÃO', 'faixa': 'A1733:P2812', 'celulas': len(sel), 'intervalos': faixas,
            'formato': {'backgroundColor': '#1B111D', 'textFormat': {'fontFamily': 'Roboto', 'fontSize': 10,
                        'foregroundColor': '#FBE4EF'}},
            'outras_diferencas_de_estilo': outras,
            'sha256_intervalos': hashlib.sha256('\n'.join(faixas).encode()).hexdigest()}


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


if __name__ == '__main__':
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--revisao', required=True, type=Path)
    p.add_argument('--backup', required=True, type=Path)
    p.add_argument('--snapshot', required=True, type=Path)
    p.add_argument('--descartavel', required=True, type=Path)
    p.add_argument('--saida', required=True, type=Path)
    p.add_argument('--plano', required=True, type=Path)
    a = p.parse_args()
    w, f = {}, {}
    for n in ('revisao', 'backup', 'snapshot', 'descartavel'):
        w[n], f[n] = foto(getattr(a, n))
    comuns = [s for s in f['backup']['abas'] if s not in STUDIES]
    assert comuns == f['revisao']['abas'], 'Abas da revisão diferem do backup sem os estudos'
    plano = plano_fundo(w['revisao'], w['descartavel'])
    rel = {
        'exportacoes_sha256': {n: sha(getattr(a, n)) for n in w},
        'resumo': {n: resumo(f[n]) for n in f},
        'backup_menos_revisao_abas': sorted(set(f['backup']['abas']) - set(f['revisao']['abas'])),
        'snapshot_x_revisao': diff(f['snapshot'], f['revisao']),
        'backup_x_revisao_abas_comuns': diff(f['backup'], f['revisao']),
        'revisao_x_descartavel_conteudo': diff(f['revisao'], f['descartavel']),
        'estilo_snapshot_x_revisao': diff_estilo(w['snapshot'], w['revisao']),
        'estilo_revisao_x_descartavel': diff_estilo(w['revisao'], w['descartavel']),
        'plano_fundo': {k: plano[k] for k in ('celulas', 'sha256_intervalos', 'outras_diferencas_de_estilo')},
    }
    rel['plano_fundo']['intervalos'] = len(plano['intervalos'])
    a.saida.write_text(json.dumps(rel, ensure_ascii=False, indent=1) + '\n', encoding='utf8')
    a.plano.write_text(json.dumps(plano, ensure_ascii=False, indent=1) + '\n', encoding='utf8')
    print(json.dumps({'resumo': rel['resumo'], 'plano': rel['plano_fundo']}, ensure_ascii=False))
