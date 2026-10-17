"""Mesmas entradas e asserções da regressão Configuração, sem recalcular XLSX."""
import ast
import hashlib
import json
import os
from pathlib import Path
import runpy
import sys
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
os.chdir(ROOT)
sys.path.insert(0,str(ROOT/'ficha-v01'))
SOURCE = ROOT/'regressao-configuracao.py'
helpers=runpy.run_path(str(ROOT/'acabamento-b47/regressao-pessoal-nativa.py'))


def preparar(folder):
    source=SOURCE.read_text(encoding='utf-8'); tree=ast.parse(source)
    e={'__file__':str(SOURCE),'__name__':'configuracao_nativa'}
    split=next(i for i,n in enumerate(tree.body) if isinstance(n,ast.Assign) and ast.unparse(n.targets[0])=='wb')
    exec(compile(ast.Module(body=tree.body[:split],type_ignores=[]),str(SOURCE),'exec'),e)
    initial={}
    def carregar(*args,**kwargs):
        wb=load_workbook(*args,**kwargs)
        initial.update({(s.title,c.coordinate):c.value for s in wb for c in s._cells.values() if c.value is not None})
        return wb
    e['load_workbook']=carregar
    end=next(i for i,n in enumerate(tree.body) if isinstance(n,ast.If) and 'shutil.which' in ast.unparse(n.test))
    exec(compile(ast.Module(body=tree.body[split:end],type_ignores=[]),str(SOURCE),'exec'),e)
    cases=[]
    def salvar(path):
        edits=[]
        for s in e['wb']:
            for c in s._cells.values():
                before=initial.get((s.title,c.coordinate)); after=c.value
                if before==after:continue
                if isinstance(before,str) and before.startswith('=') and isinstance(after,str) and after.replace('_xlfn.','')==before:continue
                edits.append([s.title,c.coordinate,after])
        cases.append({'nome':Path(path).stem,'entradas':edits})
    e['wb'].save=salvar
    block=next(n for n in tree.body if isinstance(n,ast.With))
    stop=next(i for i,n in enumerate(block.body) if isinstance(n,ast.Assign) and ast.unparse(n.targets[0])=='saida')
    block.body=block.body[:stop]
    exec(compile(ast.Module(body=[block],type_ignores=[]),str(SOURCE),'exec'),e)
    inputs={(s,c) for case in cases for s,c,v in case['entradas']}
    outputs={(s,c) for _,s,c,_ in e['alvos']}|{('DADOS_CONFIG','G'+str(r)) for *_,r in e['alvos']}|{('DADOS_CONFIG',c) for c in ['E5']+['M'+str(r) for r in range(4,11)]}
    plan={'fonte':SOURCE.name,'sha256_teste':hashlib.sha256(source.encode()).hexdigest(),
          'cenarios':cases,'entradas':[[s,c,initial.get((s,c))] for s,c in sorted(inputs)],
          'faixas_entrada':helpers['ranges'](inputs),'faixas_saida':helpers['ranges'](outputs),
          'alvos':e['alvos'],'mudancas':e['casos']}
    folder.mkdir(parents=True,exist_ok=True)
    (folder/'plano-configuracao.json').write_text(json.dumps(plan,ensure_ascii=False),encoding='utf-8')
    print(json.dumps({'cenarios':len(cases),'entradas':len(inputs),'alvos':len(e['alvos'])}))


def conferir(folder):
    source=SOURCE.read_text(encoding='utf-8')
    plan=json.loads((folder/'plano-configuracao.json').read_text(encoding='utf-8'))
    assert hashlib.sha256(source.encode()).hexdigest()==plan['sha256_teste']
    books={}
    for i,case in enumerate(plan['cenarios']):
        data=json.loads((folder/f'configuracao-{i:02}.json').read_text(encoding='utf-8'))
        assert not data.get('isError')
        wb=Workbook(); wb.remove(wb.active)
        for s in data['structuredContent']['sheets']:
            ws=wb.create_sheet(s['properties']['title'])
            for b in s.get('data',[]):
                for r,row in enumerate(b.get('rowData',[]),b.get('startRow',0)+1):
                    for c,cell in enumerate(row.get('values',[]),b.get('startColumn',0)+1):
                        v=cell.get('effectiveValue',{})
                        value=v.get('numberValue',v.get('boolValue',v.get('stringValue')))
                        if 'errorValue' in v:
                            value={'N_A':'#N/A','REF':'#REF!','VALUE':'#VALUE!','NAME':'#NAME?',
                                   'DIVIDE_BY_ZERO':'#DIV/0!'}.get(v['errorValue']['type'],v['errorValue']['type'])
                        ws.cell(r,c).value=value
        books[case['nome']]=wb
    tree=ast.parse(source)
    # As funções checa/numero e o laço de asserções são os originais.
    nodes=[n for n in tree.body if isinstance(n,(ast.Import,ast.ImportFrom,ast.FunctionDef)) and
           (not isinstance(n,ast.FunctionDef) or n.name in ('checa','numero'))]
    e={'FALHAS':[],'alvos':plan['alvos'],'casos':plan['mudancas'],'__file__':str(SOURCE)}
    exec(compile(ast.Module(body=nodes,type_ignores=[]),str(SOURCE),'exec'),e)
    e['load_workbook']=lambda path,**kwargs:books[Path(path).stem]
    e['saida']=Path('.')
    block=next(n for n in tree.body if isinstance(n,ast.With))
    assertions=next(n for n in block.body if isinstance(n,ast.For) and ast.unparse(n.iter)=='casos.items()')
    exec(compile(ast.Module(body=[assertions],type_ignores=[]),str(SOURCE),'exec'),e)
    print('CONFIGURAÇÃO CONFERE' if not e['FALHAS'] else f"{len(e['FALHAS'])} FALHAS")
    sys.exit(bool(e['FALHAS']))


if __name__=='__main__':
    {'preparar':preparar,'conferir':conferir}[sys.argv[1]](Path(sys.argv[2]))
