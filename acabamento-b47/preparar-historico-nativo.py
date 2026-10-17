"""Prepara referência histórica a partir dos GS do Git, nunca importando XLSX.

O XLSX congelado é apenas conferido offline, célula por célula, contra o GS
da mesma época. Só se forem equivalentes o modelo JSON pode ser ensaiado em
descartável nativa. Não acessa Google nem escreve no checkout histórico.
"""
import hashlib
import io
import json
from pathlib import Path
import subprocess
import sys
import zipfile
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as L
from openpyxl.formula.translate import Translator

ROOT=Path(__file__).resolve().parents[1]
sys.path[:0]=[str(ROOT/'ficha'),str(ROOT/'capacidade-b47'),str(ROOT)]
import emitir_gs
from migracao import equivalent


def preparar(kind,folder):
    refs={'campanha':('d5a1c38','base-parte1.xlsx'),
          'custom':('d0427b6','base-parte2-local.xlsx')}
    revision,fixture=refs[kind]
    sha=subprocess.check_output(['git','rev-parse',revision],cwd=ROOT,text=True).strip()
    target=(folder/'historicos'/sha).resolve();target.mkdir(parents=True,exist_ok=True)
    archive=subprocess.check_output(['git','archive','--format=zip',sha,'apps-script'],cwd=ROOT)
    with zipfile.ZipFile(io.BytesIO(archive)) as z:
        for member in z.infolist():
            if member.is_dir():continue
            dest=(target/member.filename).resolve()
            assert dest.is_relative_to(target)
            dest.parent.mkdir(parents=True,exist_ok=True)
            data=z.read(member)
            if dest.exists():assert dest.read_bytes()==data
            else:dest.write_bytes(data)
    sheets=emitir_gs.abas_do_script(str(target/'apps-script/Ficha.gs'))
    # abas_do_script expande as cópias de estilo; criarAba_ também preenche
    # `abaixo` por copyTo relativo (ficha/modelo.gs.js). Reproduzir esse passo.
    for sheet in sheets:
        values={(v[0],v[1]):list(v) for v in sheet['vals']}
        for r1,c1,r2,c2 in sheet.get('abaixo',[]):
            for r in range(r1+1,r2+1):
                for c in range(c1,c2+1):
                    seed=values.get((r1,c))
                    if seed is None:values.pop((r,c),None);continue
                    v=seed[2]
                    if isinstance(v,str) and v.startswith('='):
                        v=Translator(v,origin=L(c)+str(r1)).translate_formula(L(c)+str(r))
                    values[(r,c)]=[r,c,v,*seed[3:]]
        sheet['vals']=list(values.values())
    generated={s['nome']:{L(c)+str(r):v for r,c,v,*_ in s['vals'] if v not in (None,'')} for s in sheets}
    file=ROOT/'nuvem-b47/fixtures'/fixture
    wb=load_workbook(file,read_only=False)
    dif=[]
    for ws in wb:
        a={c.coordinate:c.value for c in ws._cells.values() if c.value not in (None,'')}
        b=generated.get(ws.title,{})
        for c in sorted(set(a)|set(b)):
            x,y=a.get(c),b.get(c)
            # Mesma regra de transporte versionada: ArrayFormula, SPARKLINE
            # embrulhada pela exportação Excel e texto numérico com apóstrofo.
            if not equivalent(emitir_gs._valor(x),y):dif.append([ws.title,c,x,y])
    model={'referencia':kind,'git_sha':sha,'fixture_sha256':hashlib.sha256(file.read_bytes()).hexdigest(),
           'diferencas':dif,'abas':sheets}
    (folder/f'historico-{kind}.json').write_text(json.dumps(model,ensure_ascii=False),encoding='utf-8')
    print(json.dumps({'referencia':kind,'git_sha':sha,'abas':len(sheets),'diferencas':len(dif),
                      'primeiras':dif[:5]},ensure_ascii=False))
    if dif:sys.exit(1)


def lotes(kind,folder):
    model=json.loads((folder/f'historico-{kind}.json').read_text(encoding='utf-8'))
    assert not model['diferencas']
    meta=json.loads((folder/f'historico-{kind}-metadata.json').read_text(encoding='utf-8'))
    assert meta['properties']['title'].startswith('TESTE DESCARTÁVEL B47 — referência histórica')
    assert meta['spreadsheetId'] not in ['1EdEXzFRL8n21dLb60lxE2wbP4PFxFUiBdw7BiTTkNg0','1u0zWQFVaRTTc_VlrAJsabgAiJ1D4rTYhzG7QQWof1ro']
    sheets={s['properties']['title']:s['properties'] for s in meta['sheets']}
    requests=[]
    selected=[s for s in model['abas'] if s['nome']!='ANALISE']
    for sheet in selected:
        p=sheets[sheet['nome']]
        assert p['gridProperties']['rowCount']>=sheet['rows'] and p['gridProperties']['columnCount']>=sheet['cols']
        requests.append({'repeatCell':{'range':{'sheetId':p['sheetId']},'cell':{},'fields':'userEnteredValue'}})
    # Mesmo procedimento do construtor canônico: escrever fórmulas em en_US,
    # então voltar a pt_BR. A API não converte separadores no envio bruto.
    clear=[{'updateSpreadsheetProperties':{'properties':{'locale':'en_US'},'fields':'locale'}}]+requests
    requests=[]
    for sheet in selected:
        p=sheets[sheet['nome']]; rows={}
        for r,c,v,*_ in sheet['vals']:
            if v is None or v=='':continue
            if isinstance(v,str) and v.startswith("'"):v=v[1:]
            key='boolValue' if isinstance(v,bool) else 'numberValue' if isinstance(v,(int,float)) else 'formulaValue' if v.startswith('=') else 'stringValue'
            rows.setdefault(r,{})[c]={'userEnteredValue':{key:v}}
        for r,cells in sorted(rows.items()):
            lo,hi=min(cells),max(cells)
            requests.append({'updateCells':{'start':{'sheetId':p['sheetId'],'rowIndex':r-1,'columnIndex':lo-1},
                             'rows':[{'values':[cells.get(c,{}) for c in range(lo,hi+1)]}], 'fields':'userEnteredValue'}})
    destination=folder/f'historico-{kind}-lotes';destination.mkdir(exist_ok=True)
    batches=[clear];batch=[];size=0
    for request in requests:
        length=len(json.dumps(request,ensure_ascii=False))
        if batch and (size+length>120000 or len(batch)>=180):batches.append(batch);batch=[];size=0
        batch.append(request);size+=length
    if batch:batches.append(batch)
    batches.append([{'updateSpreadsheetProperties':{'properties':{'locale':'pt_BR'},'fields':'locale'}}])
    for i,batch in enumerate(batches):
        (destination/f'{i:03}.json').write_text(json.dumps({'spreadsheet_id':meta['spreadsheetId'],'requests':batch},ensure_ascii=False),encoding='utf-8')
    summary={'origem':'GS versionados; sem importação XLSX','referencia':model['git_sha'],'diferencas_fixture':0,
             'planilha':meta['spreadsheetId'],'lotes':len(batches),'pedidos':sum(map(len,batches)),
             'abas':len(selected),'analise':'não substituída: exportador sem referências entrantes, fora da comparação histórica'}
    (folder/f'historico-{kind}-plano.json').write_text(json.dumps(summary,ensure_ascii=False),encoding='utf-8')
    print(json.dumps(summary,ensure_ascii=False))


def restaurar(kind,folder):
    model=json.loads((folder/f'historico-{kind}.json').read_text(encoding='utf-8'))
    meta=json.loads((folder/f'historico-{kind}-metadata.json').read_text(encoding='utf-8'))
    plan=json.loads((folder/f'plano-{kind}.json').read_text(encoding='utf-8'))
    values={s['nome']:{L(c)+str(r):v for r,c,v,*_ in s['vals']} for s in model['abas']}
    ids={s['properties']['title']:s['properties']['sheetId'] for s in meta['sheets']}
    from openpyxl.utils.cell import coordinate_to_tuple
    requests=[]
    for s,c,_ in plan['cenarios'][0]['entradas']:
        r,col=coordinate_to_tuple(c);v=values[s].get(c)
        key='boolValue' if isinstance(v,bool) else 'numberValue' if isinstance(v,(int,float)) else 'formulaValue' if isinstance(v,str) and v.startswith('=') else 'stringValue'
        requests.append({'updateCells':{'start':{'sheetId':ids[s],'rowIndex':r-1,'columnIndex':col-1},'rows':[{'values':[{} if v is None or v=='' else {'userEnteredValue':{key:v}}]}],'fields':'userEnteredValue'}})
    (folder/f'historico-{kind}-restaurar.json').write_text(json.dumps({'spreadsheet_id':meta['spreadsheetId'],'requests':requests},ensure_ascii=False),encoding='utf-8')
    print(json.dumps({'restaurar':kind,'celulas':len(requests)}))


def delta(kind,folder):
    assert kind=='custom'
    model=json.loads((folder/f'historico-{kind}.json').read_text(encoding='utf-8'))
    before=json.loads((folder/'historico-campanha.json').read_text(encoding='utf-8'))
    meta=json.loads((folder/f'historico-{kind}-metadata.json').read_text(encoding='utf-8'))
    assert not model['diferencas'] and not before['diferencas']
    assert meta['properties']['title'].startswith('TESTE DESCARTÁVEL B47 — referência histórica')
    ids={s['properties']['title']:s['properties']['sheetId'] for s in meta['sheets']}
    old={s['nome']:{(r,c):v for r,c,v,*_ in s['vals'] if v not in (None,'')} for s in before['abas']}
    requests=[{'updateSpreadsheetProperties':{'properties':{'locale':'en_US'},'fields':'locale'}}]
    changed=[]
    for s in model['abas']:
        if s['nome']=='ANALISE':continue
        new={(r,c):v for r,c,v,*_ in s['vals'] if v not in (None,'')};prior=old[s['nome']]
        for r,c in sorted(set(new)|set(prior)):
            v=new.get((r,c))
            if equivalent(v,prior.get((r,c))):continue
            changed.append([s['nome'],L(c)+str(r)])
            if isinstance(v,str) and v.startswith("'"):v=v[1:]
            key='boolValue' if isinstance(v,bool) else 'numberValue' if isinstance(v,(int,float)) else 'formulaValue' if isinstance(v,str) and v.startswith('=') else 'stringValue'
            requests.append({'updateCells':{'start':{'sheetId':ids[s['nome']],'rowIndex':r-1,'columnIndex':c-1},'rows':[{'values':[{} if v is None else {'userEnteredValue':{key:v}}]}],'fields':'userEnteredValue'}})
    batches=[];batch=[];size=0
    for req in requests:
        length=len(json.dumps(req,ensure_ascii=False))
        if batch and (size+length>120000 or len(batch)>=180):batches.append(batch);batch=[];size=0
        batch.append(req);size+=length
    if batch:batches.append(batch)
    batches.append([{'updateSpreadsheetProperties':{'properties':{'locale':'pt_BR'},'fields':'locale'}}])
    destination=folder/f'historico-{kind}-lotes';destination.mkdir(exist_ok=True)
    for i,batch in enumerate(batches):
        (destination/f'{i:03}.json').write_text(json.dumps({'spreadsheet_id':meta['spreadsheetId'],'requests':batch},ensure_ascii=False),encoding='utf-8')
    summary={'origem':'delta GS versionado, aplicado em cópia da referência campanha restaurada',
             'referencia':model['git_sha'],'base':before['git_sha'],'planilha':meta['spreadsheetId'],
             'diferencas_fixture':0,'lotes':len(batches),'celulas_alteradas':changed}
    (folder/f'historico-{kind}-plano.json').write_text(json.dumps(summary,ensure_ascii=False),encoding='utf-8')
    print(json.dumps({'tipo':kind,'lotes':len(batches),'celulas_alteradas':len(changed)},ensure_ascii=False))


if __name__=='__main__':
    {'lotes':lotes,'restaurar':restaurar,'delta':delta,'preparar':preparar}[sys.argv[3] if len(sys.argv)>3 else 'preparar'](sys.argv[1],Path(sys.argv[2]))
