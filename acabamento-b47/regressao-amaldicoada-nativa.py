"""Executa os cenários/oráculos da Amaldiçoada com respostas nativas Sheets."""
import ast
import gc
import hashlib
import json
import os
from pathlib import Path
import runpy
import sys
from openpyxl import Workbook
from openpyxl.utils.cell import range_boundaries

ROOT=Path(__file__).resolve().parents[1];os.chdir(ROOT);sys.path.insert(0,str(ROOT))
SOURCE=ROOT/'regressao-amaldicoada.py'
helpers=runpy.run_path(str(ROOT/'acabamento-b47/regressao-pessoal-nativa.py'))
SELECTED=('FICHA','DADOS','FICHA AMALDIÇOADA','DADOS_AM')


def contexto():
    source=SOURCE.read_text(encoding='utf-8');tree=ast.parse(source)
    stop=next(i for i,n in enumerate(tree.body) if isinstance(n,ast.Assign) and
              any(isinstance(t,ast.Name) and t.id=='WB' for t in n.targets) and
              isinstance(n.value,ast.Call) and ast.unparse(n.value.func)=='recalcula_tudo')
    func=next(n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name=='prepara')
    body=[]
    for n in func.body:
        if isinstance(n,ast.Assign) and any(isinstance(t,ast.Name) and t.id=='wb' for t in n.targets):
            n.value=ast.parse('Entradas()',mode='eval').body
        if isinstance(n,ast.For) and ast.unparse(n.iter)=='wb':continue
        if isinstance(n,ast.Expr) and isinstance(n.value,ast.Call):
            call=ast.unparse(n.value.func)
            if call=='shutil.copy':continue
            if call=='wb.save':
                n=ast.parse('capturar(nome, wb.cells)').body[0]
        body.append(n)
    func.body=body
    cases=[]
    e={'__file__':str(SOURCE),'__name__':'amaldicoada_nativa','Entradas':helpers['Entradas'],
       'capturar':lambda name,cells:cases.append({'nome':name,'entradas':[[s,c,v] for (s,c),v in cells.items()]})}
    exec(compile(ast.fix_missing_locations(ast.Module(body=tree.body[:stop],type_ignores=[])),str(SOURCE),'exec'),e)
    return source,tree,stop,e,cases


def preparar(folder):
    source,tree,stop,e,cases=contexto()
    wb=e['WB0'];inputs={(s,c) for case in cases for s,c,v in case['entradas']}
    cols={}
    for name in SELECTED:
        for c in wb[name]._cells.values():
            if c.data_type!='f':continue
            lo,hi=cols.get((name,c.column_letter),(c.row,c.row))
            cols[name,c.column_letter]=min(lo,c.row),max(hi,c.row)
    ranges=[f"'{s}'!{c}{lo}:{c}{hi}" for (s,c),(lo,hi) in sorted(cols.items())]
    groups=[];group=[];size=0
    for rg in ranges:
        x=range_boundaries(rg.split('!')[1]);count=(x[2]-x[0]+1)*(x[3]-x[1]+1)
        if group and (size+count>14000 or len(group)>=80):groups.append(group);group=[];size=0
        group.append(rg);size+=count
    if group:groups.append(group)
    plan={'fonte':SOURCE.name,'sha256_teste':hashlib.sha256(source.encode()).hexdigest(),
          'cenarios':cases,'entradas':[[s,c,wb[s][c].value] for s,c in sorted(inputs)],
          'faixas_entrada':helpers['ranges'](inputs),'grupos_saida':groups,'faixas_saida':ranges}
    folder.mkdir(parents=True,exist_ok=True)
    (folder/'plano-amaldicoada.json').write_text(json.dumps(plan,ensure_ascii=False),encoding='utf-8')
    print(json.dumps({'cenarios':len(cases),'entradas':len(inputs),'grupos_leitura':len(groups)},ensure_ascii=False))


def conferir(folder):
    source,tree,stop,e,cases=contexto()
    plan=json.loads((folder/'plano-amaldicoada.json').read_text(encoding='utf-8'))
    assert hashlib.sha256(source.encode()).hexdigest()==plan['sha256_teste']
    assert cases==plan['cenarios'],'cenários mudaram'
    books={}
    for i,case in enumerate(cases):
        wb=Workbook();wb.remove(wb.active)
        needed=set(SELECTED)|{s for s,_,_ in case['entradas']}
        for old in e['WB0']:
            if old.title not in needed:continue
            ws=wb.create_sheet(old.title)
            for c in old._cells.values():
                if c.value is not None:ws[c.coordinate]='__NATIVO_NAO_LIDO__' if c.data_type=='f' else c.value
        # Cada fórmula solicitada foi lida mesmo quando o efetivo é vazio.
        for rg in plan['faixas_saida']:
            name,coords=rg.split('!');name=name.strip("'")
            for row in wb[name][coords]:
                for cell in row:cell.value=None
        for s,c,v in case['entradas']:wb[s][c]=v
        for j in range(len(plan['grupos_saida'])):
            response=json.loads((folder/f'amaldicoada-{i:02}-{j:02}.json').read_text(encoding='utf-8'))
            assert not response.get('isError')
            for s in response['structuredContent']['sheets']:
                for b in s.get('data',[]):
                    for r,row in enumerate(b.get('rowData',[]),b.get('startRow',0)+1):
                        for c,cell in enumerate(row.get('values',[]),b.get('startColumn',0)+1):
                            v=cell.get('effectiveValue',{});val=v.get('numberValue',v.get('boolValue',v.get('stringValue')))
                            if 'errorValue' in v:val='#'+v['errorValue']['type']
                            wb[s['properties']['title']].cell(r,c).value=val
        books[case['nome']]=wb
    e['WB']=books
    # O restante é o código original, inclusive o modelo independente do livro
    # e os ensaios VM. Registrar essas duas categorias separadamente no relatório.
    split=next(i for i,n in enumerate(tree.body) if isinstance(n,ast.Assign) and
               any(isinstance(t,ast.Name) and t.id=='PROG_JS' for t in n.targets))
    exec(compile(ast.Module(body=tree.body[stop+1:split],type_ignores=[]),str(SOURCE),'exec'),e)
    # Não manter 20 workbooks em memória enquanto o simulador Node constrói.
    books.clear();e['WB']={};gc.collect()
    for node in tree.body[split:]:
        exec(compile(ast.Module(body=[node],type_ignores=[]),str(SOURCE),'exec'),e)
        if isinstance(node,ast.FunctionDef) and node.name=='monta_no_sheets':
            original=e['monta_no_sheets']
            inputs=[ROOT/'medidas/sheets-de-mentira.js',ROOT/'medidas/carregar-ficha.js',*sorted((ROOT/'apps-script').glob('*.gs'))]
            sha=hashlib.sha256(e['PROG_JS'].encode()+b''.join(p.read_bytes() for p in inputs)).hexdigest()
            def cache(modo='',original=original,sha=sha):
                path=folder/('amaldicoada-vm-'+(modo or 'padrao')+'.json')
                if path.exists():
                    data=json.loads(path.read_text(encoding='utf-8'))
                    if data.get('sha256_entradas')==sha:return data['resultado']
                result=original(modo)
                if result is not None:path.write_text(json.dumps({'sha256_entradas':sha,'resultado':result},ensure_ascii=False),encoding='utf-8')
                return result
            e['monta_no_sheets']=cache


if __name__=='__main__':
    {'preparar':preparar,'conferir':conferir}[sys.argv[1]](Path(sys.argv[2]))
