"""Planos nativos de Campanha/Custom, preservando os oráculos originais."""
import ast
import hashlib
import json
import math
import os
from pathlib import Path
import runpy
import sys
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1];os.chdir(ROOT);sys.path.insert(0,str(ROOT))
helpers=runpy.run_path(str(ROOT/'acabamento-b47/regressao-pessoal-nativa.py'))


def preparar(kind,folder):
    file=ROOT/f'regressao-{kind}.py';source=file.read_text(encoding='utf-8');tree=ast.parse(source)
    end=next(i for i,n in enumerate(tree.body) if isinstance(n,ast.With))
    nodes=[n for n in tree.body[:end] if not (isinstance(n,ast.If) and 'shutil.which' in ast.unparse(n.test))]
    e={'__file__':str(file),'__name__':kind+'_nativa'}
    exec(compile(ast.Module(body=nodes,type_ignores=[]),str(file),'exec'),e)
    block=tree.body[end]
    stop=next(i for i,n in enumerate(block.body) if isinstance(n,ast.Assign) and
              any(isinstance(t,ast.Name) and t.id==('paths' if kind=='campanha' else 'cached') for t in n.targets))
    block.body=block.body[:stop]
    exec(compile(ast.Module(body=[block],type_ignores=[]),str(file),'exec'),e)
    if kind=='campanha':
        common={('FICHA',e['idx']['caminho']):'Vanguarda'}
        common.update({('FICHA',e['idx'][k]):v for k,v in
                       [('atr_base_Força',3),('atr_base_Destreza',2),('atr_base_Constituição',3),
                        ('atr_base_Inteligência',1),('atr_base_Essência',2),('buff de vida',5),('buff de deslocamento',.5)]})
        cases=[{'nome':name,'historico':src==e['base'],'entradas':[[s,c,v] for (s,c),v in
                {**common,('FICHA',e['idx']['nivel']):level,**edits}.items()]} for name,src,edits,level in e['casos']]
    else:
        cases=[{'nome':name,'historico':src==e['base'],'entradas':[[s,c,v] for (s,c),v in
                {**e['common'],**edits}.items()]} for name,src,edits in e['cases']]
    outputs=set()
    for n in ast.walk(ast.parse(source)):
        if not isinstance(n,ast.Call) or not isinstance(n.func,ast.Name):continue
        if n.func.id=='get' and len(n.args)==3 and all(isinstance(a,ast.Constant) and isinstance(a.value,str) for a in n.args[1:]):
            sheet,coord=[a.value for a in n.args[1:]]
            if __import__('re').fullmatch('[A-Z]+[0-9]+',coord):outputs.add((sheet,coord))
        if n.func.id in ('f','v') and len(n.args)==2 and isinstance(n.args[1],ast.Constant):
            key=n.args[1].value
            if key in e['idx']:outputs.add(('FICHA',e['idx'][key]))
    # Os laços e avisos dinâmicos também fazem parte do oráculo.
    if kind=='campanha':
        outputs.update(('FICHA',e['idx'][k]) for k in ['vida_max','energia_max','defesa','proteção','corpo a corpo','à distância'])
    else:
        outputs.update((e['custom'].NOME,c+str(r)) for c in 'STV' for r in range(2,12))
        outputs.update((e['custom'].NOME,'CP'+str(r)) for r in range(2,26))
        outputs.update((e['custom'].NOME,c) for c in ('B122','B123'))
    inputs={(s,c) for case in cases for s,c,v in case['entradas']}
    plan={'fonte':file.name,'sha256_teste':hashlib.sha256(source.encode()).hexdigest(),
          'cenarios':cases,'entradas':[[s,c] for s,c in sorted(inputs)],
          'faixas_entrada':helpers['ranges'](inputs),'faixas_saida':helpers['ranges'](outputs),
          'idx':e['idx'],'falhas_estruturais':e['FALHAS' if kind=='campanha' else 'falhas']}
    if kind=='custom':plan['contrato_dif']=e['contrato_dif']
    folder.mkdir(parents=True,exist_ok=True)
    (folder/f'plano-{kind}.json').write_text(json.dumps(plan,ensure_ascii=False),encoding='utf-8')
    print(json.dumps({'tipo':kind,'cenarios':len(cases),'estrutura':plan['falhas_estruturais']},ensure_ascii=False))


def conferir(kind,folder):
    file=ROOT/f'regressao-{kind}.py';source=file.read_text(encoding='utf-8');tree=ast.parse(source)
    plan=json.loads((folder/f'plano-{kind}.json').read_text(encoding='utf-8'))
    assert hashlib.sha256(source.encode()).hexdigest()==plan['sha256_teste']
    results={}
    for i,case in enumerate(plan['cenarios']):
        if case['historico']:continue
        data=json.loads((folder/f'{kind}-{i:02}.json').read_text(encoding='utf-8'))
        assert not data.get('isError');values={}
        for s in data['structuredContent']['sheets']:
            cells=values.setdefault(s['properties']['title'],{})
            for b in s.get('data',[]):
                for r,row in enumerate(b.get('rowData',[]),b.get('startRow',0)+1):
                    for c,cell in enumerate(row.get('values',[]),b.get('startColumn',0)+1):
                        v=cell.get('effectiveValue',{});val=v.get('numberValue',v.get('boolValue',v.get('stringValue')))
                        if 'errorValue' in v:val={'N_A':'#N/A','VALUE':'#VALUE!'}.get(v['errorValue']['type'],'#'+v['errorValue']['type'])
                        cells[helpers['get_column_letter'](c)+str(r)]=val
        results[case['nome']]=values
    e={'results':results,'dados':results,'idx':plan['idx'],'FALHAS':[],'falhas':[],
       '__file__':str(file),'ROOT':ROOT,'math':math,'load_workbook':load_workbook,
       'base':ROOT/'nuvem-b47/fixtures/base-parte1.xlsx'}
    if kind=='custom':
        sys.path.insert(0,str(ROOT/'ficha-v01'));import custom;e['custom']=custom
    block=next(n for n in tree.body if isinstance(n,ast.With))
    funcs=[n for n in tree.body+block.body if isinstance(n,ast.FunctionDef) and n.name in ('check','get','f')]
    exec(compile(ast.Module(body=funcs,type_ignores=[]),str(file),'exec'),e)
    # A comparação completa das bases históricas é feita separadamente, sem
    # assumir que o cenário vazio da ficha atual é a antiga referência.
    start=next(i for i,n in enumerate(block.body) if isinstance(n,ast.Expr) and isinstance(n.value,ast.Call)
               and isinstance(n.value.func,ast.Name) and n.value.func.id=='check' and
               isinstance(n.value.args[0],ast.Constant) and (str(n.value.args[0].value).startswith('CAMPANHA vazia:') or str(n.value.args[0].value).startswith('custom vazio:')))+1
    end=next((i for i,n in enumerate(block.body[start:],start) if isinstance(n,ast.Expr) and isinstance(n.value,ast.Call)
              and ast.unparse(n.value.func).endswith('write_text')),len(block.body))
    exec(compile(ast.Module(body=block.body[start:end],type_ignores=[]),str(file),'exec'),e)
    fails=e['FALHAS'] if kind=='campanha' else e['falhas']
    print(kind.upper()+' NATIVA: '+str(len(fails))+' falhas mecânicas; comparação histórica separada')
    sys.exit(bool(fails))


if __name__=='__main__':
    {'preparar':preparar,'conferir':conferir}[sys.argv[1]](sys.argv[2],Path(sys.argv[3]))
