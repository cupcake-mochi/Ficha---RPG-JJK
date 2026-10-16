#!/usr/bin/env python3
"""Plano offline de migração v9→v10. Nunca acessa Google ou executa construir.

Snapshot normalizado: {spreadsheetId, captured_at, backup_id, ranges:{aba:[A1]},
cells:{aba:{A1:{userEnteredValue:{...},note:...}}}}. Ausência de célula só significa
vazio dentro de uma faixa comprovadamente lida. O plano não é licença de envio:
exige releitura/igualdade do snapshot e prova descartável antes da revisão.
"""
import argparse
from copy import deepcopy
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import sys
from openpyxl.utils.cell import range_boundaries, get_column_letter

REVISAO = '1EdEXzFRL8n21dLb60lxE2wbP4PFxFUiBdw7BiTTkNg0'
OFICIAL = '1u0zWQFVaRTTc_VlrAJsabgAiJ1D4rTYhzG7QQWof1ro'
LIMITES = {'familia':10, 'melhoria':30, 'restricao':30}


def sha(value):
    return hashlib.sha256(json.dumps(value,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()).hexdigest()


def raw(model):
    return {a['nome']:{c[0]:c[1] for c in a['celulas']} for a in model['layout']['abas']}


def contract(cells):
    return {v:[cells.get(col+c[1:],'') for col in 'BCD'] for c,v in cells.items()
            if re.fullmatch(r'A\d+',c) and c!='A1'}


def source(value):
    # ANALISE C is emitted as ="'CONFIGURAÇÃO'!B943" to preserve apostrophes.
    if value.startswith('="') and value.endswith('"'): value=value[2:-1].replace('""','"')
    if '!' not in value: return None
    sheet,coord=value.rsplit('!',1)
    return sheet.strip("'"),coord.replace('$','')


def uv(value):
    if value in (None,''):return {}
    if isinstance(value,bool):return {'boolValue':value}
    if isinstance(value,(int,float)):return {'numberValue':value}
    return {'formulaValue' if value.startswith('=') else 'stringValue':value}


def value(cell):
    v=cell.get('userEnteredValue',{})
    return next(iter(v.values()),None)


def canonical(v):
    """Normaliza pontuação e aspas opcionais de aba, preservando textos literais."""
    if v in (None,''): return None
    if not isinstance(v,str) or not v.startswith('='):return v
    out=[];quoted=False
    # An English baseline has argument commas. A native pt_BR formula with
    # multiple arguments has semicolons; only that form has decimal commas.
    native=';' in re.sub(r'"(?:[^"]|"")*"','',v)
    i=0
    while i<len(v):
        ch=v[i]
        if not quoted and ch=="'":
            # Sheets removes quotes around simple ASCII sheet identifiers.
            # Consume the WHOLE single-quoted name, including escaped quotes,
            # so a substring inside an unusual name can never be rewritten.
            match=re.match(r"'((?:[^']|'')*)'!",v[i:])
            if match:
                name=match[1]
                out.append(name+'!' if re.fullmatch(r'[A-Za-z_][A-Za-z0-9_]*',name) else match[0])
                i+=len(match[0]);continue
        if ch=='"':quoted=not quoted
        if not quoted:
            if ch==';':ch=','
            elif native and ch==',' and v[i-1:i].isdigit() and v[i+1:i+2].isdigit():ch='.'
        out.append(ch)
        i+=1
    return ''.join(out)


def equivalent(a,b):
    a,b=canonical(a),canonical(b)
    if isinstance(a,bool)!=isinstance(b,bool):return False
    return a==b


def covered(snapshot,sheet,coord):
    c,r,_,_=range_boundaries(coord)
    for sheet_meta in snapshot.get('metadata',{}).get('sheets',[]):
        props=sheet_meta['properties']
        if props['title']==sheet:
            grid=props['gridProperties']
            if r>grid['rowCount'] or c>grid['columnCount']:return True  # provably outside native grid
    return any(c1<=c<=c2 and r1<=r<=r2 for c1,r1,c2,r2 in
               (range_boundaries(x) for x in snapshot.get('ranges',{}).get(sheet,[])))


def required_ranges(base,target,metadata=None):
    """Full edited catalog area plus every altered value in the other tabs."""
    old,new=raw(base),raw(target)
    ranges={'CONFIGURAÇÃO':[f'A733:P{next(a["linhas"] for a in base["layout"]["abas"] if a["nome"]=="CONFIGURAÇÃO")}']}
    for sheet,cells in new.items():
        if sheet=='CONFIGURAÇÃO':continue
        coords=[c for c in old[sheet].keys()|cells.keys() if not equivalent(old[sheet].get(c),cells.get(c))]
        if coords:
            # Bounding rectangle is explicit and includes blank neighbors; the
            # snapshot reader may split it into bounded connector calls.
            boxes=[range_boundaries(c) for c in coords]
            ranges[sheet]=[f'{get_column_letter(min(b[0] for b in boxes))}{min(b[1] for b in boxes)}:{get_column_letter(max(b[2] for b in boxes))}{max(b[3] for b in boxes)}']
    # Native grids may have blank-looking padding beyond the generator's last
    # row. Every cell about to be removed must be covered, including that tail.
    if metadata:
        shrinking={'CONFIGURAÇÃO','DADOS_PROGRESSAO','DADOS_CRIACAO','DADOS_MELHORIAS','DADOS_RESTRICOES','ANALISE'}
        for sheet in metadata['sheets']:
            props=sheet['properties'];name=props['title']
            if name not in shrinking or name not in ranges:continue
            count=props['gridProperties']['rowCount'];cols=props['gridProperties']['columnCount']
            c,r,cc,rr=range_boundaries(ranges[name][0])
            ranges[name]=[f'{get_column_letter(c)}{r}:{get_column_letter(max(cc,cols))}{max(rr,count)}']
    return ranges


def planejar(base,target,snapshot,*,now=None,permitir_revisao=False,prova=None):
    old,new=raw(base),raw(target);bc,nc=contract(old['ANALISE']),contract(new['ANALISE'])
    if bc['formato.versao'][0]!='9' or nc['formato.versao'][0]!='10':raise ValueError('Exige baseline 9 e alvo 10')
    dest=snapshot.get('spreadsheetId')
    if not dest or dest==OFICIAL:raise ValueError('Destino ausente ou ficha oficial')
    if not snapshot.get('backup_id') or snapshot['backup_id']==dest:raise ValueError('Exige backup nativo integral distinto')
    now=now or datetime.now(timezone.utc)
    captured=datetime.fromisoformat(snapshot['captured_at'].replace('Z','+00:00'))
    if captured.tzinfo is None:raise ValueError('Snapshot exige timestamp com fuso UTC explícito')
    age=(now-captured).total_seconds()
    if not 0<=age<=900:raise ValueError('Snapshot deve ter no máximo 15 minutos; obter nova leitura')
    if dest==REVISAO:
        if not permitir_revisao or not prova or prova.get('resultado')!='OK' or prova.get('alvo_sha256')!=sha(target) or prova.get('destino') in (None,REVISAO,OFICIAL):
            raise ValueError('Revisão exige prova descartável do alvo exato')
    # Validate complete coverage BEFORE interpreting any missing cell as empty.
    required=required_ranges(base,target,snapshot.get('metadata'))
    for sheet,areas in required.items():
        for area in areas:
            c,r,cc,rr=range_boundaries(area)
            for y in range(r,rr+1):
                for x in range(c,cc+1):
                    if not covered(snapshot,sheet,f'{get_column_letter(x)}{y}'):
                        raise ValueError(f'Snapshot incompleto: {sheet}!{get_column_letter(x)}{y}')
    live=snapshot['cells'];inputs={}
    for k,row in bc.items():
        src=source(row[1])
        if src and src[0]=='CONFIGURAÇÃO' and int(re.search(r'\d+',src[1]).group())>=733 and k.startswith('custom.'):
            inputs[k]=src
    preserved={};mapping=[];changes={};destinations={}
    # Retained slots keep their index. Removed occupied slots block migration.
    for key,(sheet,coord) in inputs.items():
        if key.startswith('custom.treino.'):continue
        current=live.get(sheet,{}).get(coord,{})
        v=value(current);baseline=old[sheet].get(coord)
        if key not in nc:
            if not equivalent(v,baseline) or current.get('note','') not in ('', next((a.get('notas',{}).get(coord,'' ) for a in base['layout']['abas'] if a['nome']==sheet),'')):
                raise ValueError('Vaga retirada preenchida: '+key+' em '+sheet+'!'+coord)
            continue
        dst=source(nc[key][1])
        if not dst:raise ValueError('Fonte nova ausente: '+key)
        preserved.setdefault(dst[0],{})[dst[1]]=deepcopy(current)
        destinations[(sheet,coord)]=dst
        if dst!=(sheet,coord):mapping.append({'de':key,'para':key,'origem':[sheet,coord],'destino':list(dst)})
    # Training records have NEW semantic namespaces. Preserve all five fields
    # together, in old slot order, with an explicit per-record manifest.
    counts={'Perícia':0,'Ofício':0}
    for n in range(1,132):
        fields={field:inputs[f'custom.treino.{n}.{field}'] for field in ('tipo','nome','mestre_viu','atributo','efeito')}
        values={field:live.get(sh,{}).get(co,{}) for field,(sh,co) in fields.items()}
        occupied=any(not equivalent(value(values[f]),old[sh].get(co)) or values[f].get('note','') not in ('',next(a.get('notas',{}).get(co,'') for a in base['layout']['abas'] if a['nome']==sh)) for f,(sh,co) in fields.items())
        if not occupied:continue
        typ=value(values['tipo'])
        if typ not in counts:raise ValueError(f'Treino {n} preenchido sem tipo válido')
        counts[typ]+=1
        if counts[typ]>10:raise ValueError('Capacidade excedida: '+typ)
        prefix=f'custom.{"pericia" if typ=="Perícia" else "oficio"}.{counts[typ]}'
        for field,cell in values.items():
            dst=source(nc[prefix+'.'+field][1]);origin=fields[field]
            if 'formulaValue' in cell.get('userEnteredValue',{}) and dst!=origin:
                raise ValueError('Fórmula de entrada movida exige revisão explícita: '+str(origin))
            preserved.setdefault(dst[0],{})[dst[1]]=deepcopy(cell);destinations[origin]=dst
        mapping.append({'de':f'custom.treino.{n}','para':prefix})
    # Any changed unexpected content (including cells outside known inputs)
    # blocks instead of being overwritten or dropped.
    input_coords=set(inputs.values())
    for sheet,areas in required.items():
        for area in areas:
            c,r,cc,rr=range_boundaries(area)
            for y in range(r,rr+1):
                for x in range(c,cc+1):
                    coord=f'{get_column_letter(x)}{y}';cell=live.get(sheet,{}).get(coord,{})
                    before,after=old[sheet].get(coord),new[sheet].get(coord)
                    cut_tail=sheet in {'DADOS_PROGRESSAO','DADOS_CRIACAO','DADOS_MELHORIAS','DADOS_RESTRICOES','ANALISE'} and y>next(a['linhas'] for a in target['layout']['abas'] if a['nome']==sheet)
                    touched=(sheet=='CONFIGURAÇÃO' and y>=733) or cut_tail or not equivalent(before,after)
                    if not touched:continue
                    if (sheet,coord) not in input_coords and not equivalent(value(cell),before):
                        raise ValueError('Conteúdo divergente da base: '+sheet+'!'+coord)
                    if (sheet,coord) not in input_coords:
                        note=next(a.get('notas',{}).get(coord,'') for a in base['layout']['abas'] if a['nome']==sheet)
                        if cell.get('note','') not in ('',note):raise ValueError('Nota manual fora de entrada: '+sheet+'!'+coord)
                    if not equivalent(before,after):changes.setdefault(sheet,{})[coord]=after
    # Preserve notes with every surviving input. Refuse formulas that would
    # acquire a different relative meaning after moving their input cell.
    for origin,dst in destinations.items():
        cell=preserved[dst[0]][dst[1]]
        if origin!=dst and 'formulaValue' in cell.get('userEnteredValue',{}):
            raise ValueError('Fórmula de entrada movida exige revisão explícita: '+str(origin))
    return {'formato_de':9,'formato_para':10,'destino':dest,'backup_id':snapshot['backup_id'],
            'snapshot_sha256':sha(snapshot),'base_sha256':sha(base),'alvo_sha256':sha(target),
            'requer_releitura_identica':True,'mapeamento':mapping,'entradas':preserved,
            'valores_alterados':changes,'faixas':required,
            'pendente':'Aplicar lotes em descartável, validar e salvar/reler Apps Script antes da revisão.'}


def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--base',type=Path,required=True);ap.add_argument('--alvo',type=Path,required=True)
    ap.add_argument('--snapshot',type=Path);ap.add_argument('--metadata',type=Path,help='Metadados nativos para incluir caudas da grade nas faixas');ap.add_argument('--saida',type=Path,required=True)
    a=ap.parse_args();base=json.loads(a.base.read_text());target=json.loads(a.alvo.read_text())
    metadata=json.loads(a.metadata.read_text()) if a.metadata else None
    result=planejar(base,target,json.loads(a.snapshot.read_text())) if a.snapshot else {'faixas_necessarias':required_ranges(base,target,metadata),'base_sha256':sha(base),'alvo_sha256':sha(target)}
    a.saida.write_text(json.dumps(result,ensure_ascii=False,indent=2));print(a.saida)

if __name__=='__main__':main()
