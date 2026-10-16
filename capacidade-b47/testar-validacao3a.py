#!/usr/bin/env python3
"""As validações básicas da 3a continuam iguais após dividir fórmulas longas."""
import argparse,json,re,shutil,subprocess,sys,tempfile
from pathlib import Path
from openpyxl import Workbook,load_workbook
ROOT=Path(__file__).resolve().parent.parent;sys.path[:0]=[str(ROOT),str(ROOT/'ficha-v01')]
import custom
ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('--modelo',type=Path,default=ROOT/'nuvem-b47/saidas/pecas-modelo.json');args=ap.parse_args()
model=json.loads(args.modelo.read_text());raw={a['nome']:{v[0]:v[1] for v in a['celulas']} for a in model['layout']['abas']}
calc=shutil.which('libreoffice') or shutil.which('soffice')
if not calc:raise SystemExit('Calc ausente: teste impedido')
specs=[('Arma','Arma Teste',['Porrete',1,'Concussão','Versátil','—','—','—'],['1d6',4,1.5,'—','—',0]),('Escudo','Escudo Teste',[],[4,1,4,2,0]),('Uniforme','Uniforme Teste',['Grau 4'],[3,2,4,2,0]),('Item comum','Item Teste',[],[.7,0]),('Ferramenta amaldiçoada','Ferramenta Teste',['Arma Teste','Grau 2'],['Estigma',2,0,0]),('Origem','Origem Teste',custom.listas()['pericias'][:4]+['Fundamento','Nenhuma','—'],['CE4']),('Marcante','Marcante Teste',['Origem Teste'],['Uma vez por cena']),('Menor','Menor Teste',['Origem Teste','Marcante Teste'],['Uma vez por missão'])]
results={};checks=[]
def check(name,ok):checks.append(bool(ok));print(('OK ' if ok else 'FALHA ')+name,flush=True)
with tempfile.TemporaryDirectory(prefix='b47-val3a-') as temp:
 temp=Path(temp);out=temp/'out';out.mkdir()
 for case in ('vazio','validos','invalidos'):
  wb=Workbook();wb.remove(wb.active)
  for sheet in ('CONFIGURAÇÃO','DADOS_CUSTOM','DADOS_PROGRESSAO'):
   ws=wb.create_sheet(sheet)
   for c,v in raw[sheet].items():
    if v is None:continue
    if sheet=='CONFIGURAÇÃO' and isinstance(v,str) and v.startswith('='):continue
    if isinstance(v,str) and v.startswith('='):v=re.sub(r'(?<![A-Z_.])(IFS|SWITCH|TEXTJOIN)\(',r'_xlfn.\1(',v)
    ws[c]=v
  cfg=wb['CONFIGURAÇÃO']
  if case!='vazio':
   for n,(typ,name,closed,free) in enumerate(specs,1):
    g=custom.cartao(n);cfg[g['tipo']]=typ;cfg[g['nome']]=name
    vals=closed+[None]*(7-len(closed))+free+[None]*(6-len(free))
    for c,v in zip(g['campos'],vals):cfg[c]=v
  if case=='invalidos':cfg[custom.cartao(1)['nome']]='Bastão';cfg[custom.cartao(3)['campos'][7]]='três'
  file=temp/(case+'.xlsx');wb.save(file);wb.close()
  run=subprocess.run([calc,'-env:UserInstallation='+(temp/'perfil').as_uri(),'--headless','--convert-to','xlsx','--outdir',str(out),str(file)],capture_output=True,text=True,timeout=120)
  if run.returncode or not(out/file.name).exists():raise RuntimeError(run.stderr)
  wb=load_workbook(out/file.name,data_only=True);results[case]={n:wb['DADOS_CUSTOM']['U'+str(n+1)].value for n in range(1,26)};wb.close()
check('25 registros vazios não geram erro básico',all(v in (None,'') for v in results['vazio'].values()))
check('oito moldes válidos sem perda de campos ou novo requisito',all(results['validos'][n] in (None,'') for n in range(1,9)))
check('colisão com livro continua rejeitada','Nome repetido' in str(results['invalidos'][1]))
check('número inválido continua rejeitado','Informe número' in str(results['invalidos'][3]))
check('nenhum aviso básico é erro de fórmula',not any(str(v).startswith('#') for row in results.values() for v in row.values()))
print(json.dumps({'resultado':'OK' if all(checks) else 'FALHA','checagens':len(checks),'resultados':results},ensure_ascii=False));raise SystemExit(not all(checks))
