#!/usr/bin/env python3
"""Primeira/última vaga e literalidade dos dois catálogos em Calc local."""
import argparse,json,re,shutil,subprocess,tempfile,sys
from pathlib import Path
from openpyxl import Workbook,load_workbook
ROOT=Path(__file__).resolve().parent.parent
sys.path[:0]=[str(ROOT),str(ROOT/'ficha-v01')]
import custom_progressao as p
ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('--modelo',type=Path,required=True);a=ap.parse_args()
model=json.loads(a.modelo.read_text());raw={a['nome']:{c[0]:c[1] for c in a['celulas']} for a in model['layout']['abas']}
calc=shutil.which('libreoffice') or shutil.which('soffice')
if not calc:raise SystemExit('Calc ausente: teste impedido')
checks=[]
def check(name,ok):checks.append(bool(ok));print(('OK ' if ok else 'FALHA ')+name,flush=True)
with tempfile.TemporaryDirectory(prefix='b47-treinos-') as temp:
    temp=Path(temp);out=temp/'out';out.mkdir();results={}
    for case in ('validos','tipo_incorreto','duplicado','tipo_apagado_vazio'):
        wb=Workbook();cfg=wb.active;cfg.title='CONFIGURAÇÃO';d=wb.create_sheet(p.NOME)
        for c,v in raw['CONFIGURAÇÃO'].items():
            if v is not None and not(isinstance(v,str) and v.startswith('=')):cfg[c]=v
        for c,v in raw[p.NOME].items():
            if v is not None:
                if isinstance(v,str) and v.startswith('='):v=re.sub(r'(?<![A-Z_.])(IFS|SWITCH|TEXTJOIN)\(',r'_xlfn.\1(',v)
                d[c]=v
        for n,typ,name,attr in [(1,'Perícia','Rito*','Essência'),(2,'Perícia','Ritox','Força'),(10,'Perícia','Última perícia ? ~','Destreza'),(11,'Ofício','Primeiro ofício','Inteligência'),(20,'Ofício','Último ofício *','Constituição')]:
            g=p.curto(n)
            for k,v in [('tipo',typ),('nome',name),('atributo',attr)]:cfg[g[k]]=v
        if case=='tipo_incorreto':cfg[p.curto(10)['tipo']]='Ofício'
        if case=='duplicado':cfg[p.curto(2)['nome']]='Rito*'
        if case=='tipo_apagado_vazio':
            cfg[p.curto(3)['tipo']]=None;cfg[p.curto(12)['tipo']]=None
        file=temp/(case+'.xlsx');wb.save(file);wb.close()
        run=subprocess.run([calc,'-env:UserInstallation='+(temp/'perfil').as_uri(),'--headless','--convert-to','xlsx','--outdir',str(out),str(file)],capture_output=True,text=True,timeout=120)
        if run.returncode or not(out/file.name).exists():raise RuntimeError('Calc não converteu '+case+' '+run.stderr)
        wb=load_workbook(out/file.name,data_only=True);d=wb[p.NOME]
        results[case]={c:d[c].value for c in ['G4','G13','G2','G3','G11','G12','G21','AG25','AG34','AH34','AJ13','AJ22','AK22','AV11','AW11']};wb.close()
    r=results['validos'];check('curinga em nome é literal; Ritox não duplica Rito*',not r['G2'] and not r['G3'])
    check('primeira Perícia/Ofício entram depois do livro',r['AG25']=='Rito*' and r['AJ13']=='Primeiro ofício')
    check('última Perícia/Ofício entram com atributos corretos',r['AG34']=='Última perícia ? ~' and r['AH34']=='Destreza' and r['AJ22']=='Último ofício *' and r['AK22']=='Constituição')
    check('pools dos consumidores alcançam a décima vaga',r['AV11']=='Última perícia ? ~' and r['AW11']=='Último ofício *')
    r=results['tipo_incorreto'];check('colar Ofício em Perícia invalida somente esse cadastro',bool(r['G11']) and not r['AG34'] and r['AJ22']=='Último ofício *')
    check('apagar TIPO num cartão vazio não cria avisos',not results['tipo_apagado_vazio']['G4'] and not results['tipo_apagado_vazio']['G13'])
    r=results['duplicado'];check('duplicata literal exata invalida ambos os registros',bool(r['G2']) and bool(r['G3']) and not r['AG25'])
print(json.dumps({'resultado':'OK' if all(checks) else 'FALHA','cenarios':results},ensure_ascii=False))
raise SystemExit(not all(checks))
