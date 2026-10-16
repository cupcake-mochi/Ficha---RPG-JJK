#!/usr/bin/env python3
"""3c: preservação do contrato, cadastros e herança nas contas reais das cartas.

LibreOffice é uma segunda calculadora local; nenhum arquivo é importado no
Google. Os testes nativos de menus, fórmulas e visual continuam separados.
"""
import argparse
import json
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import range_boundaries

ROOT=Path(__file__).resolve().parent
sys.path[:0]=[str(ROOT/'ficha-v01')]
import custom_pecas as cp
import custom_criacao as cc
import ficha_amaldicoada as fa
import indice_ficha as ix


def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--modelo',type=Path,required=True)
    ap.add_argument('--anterior',type=Path,required=True)
    ap.add_argument('--estrutura',action='store_true')
    ap.add_argument('--evidencia',type=Path)
    ap.add_argument('--preparar-nativo',type=Path,help='Exporta entradas e resultados esperados; não acessa o Google')
    a=ap.parse_args()
    model=json.loads(a.modelo.read_text()); old=json.loads(a.anterior.read_text())
    tabs={s['nome']:s for s in model['layout']['abas']}
    raw={s:{c[0]:c[1] for c in t['celulas']} for s,t in tabs.items()}
    before={s['nome']:{c[0]:c[1] for c in s['celulas']} for s in old['layout']['abas']}
    checks=[]
    def check(name,ok,detail=None):
        print(('OK ' if ok else 'FALHA ')+name+(' '+str(detail) if not ok else ''),flush=True)
        checks.append(dict(teste=name,ok=bool(ok),detalhe=detail if not ok else None))
    check('catálogos compartilhados aprovados: 30 Melhorias + 30 Restrições',cp.MELHORIAS==30 and cp.RESTRICOES==30 and len(model['registros'])==60)
    def contract(cells):
        return {v:[cells.get(col+c[1:],'') for col in 'BCD']
                for c,v in cells.items() if re.fullmatch(r'A\d+',c) and c!='A1'}
    prior, current=contract(before['ANALISE']), contract(raw['ANALISE'])
    dif=[]
    for key, record in prior.items():
        if key.startswith(('custom.treino.','aviso.custom.treino.')): continue
        retired=re.fullmatch(r'(?:custom|aviso.custom).familia.(\d+)(?:\..*)?',key)
        if retired and int(retired[1])>10: continue
        expected=list(record)
        if key=='formato.versao': expected[0]='10'
        # These advisories keep their meaning and source; only the custom
        # selection pool shrinks from 131 to 10.
        if re.fullmatch(r'aviso.personagem.(pericia|oficio)_custom.\d+',key):
            expected[0]=expected[0].replace('$AV$132','$AV$11').replace('$AW$132','$AW$11')
            expected[0]=re.sub(r"COUNTIF\(([^,]+),('FICHA'!\$[A-Z]+\$\d+)\)",lambda m:'COUNTIF('+m[1]+','+(cc.criterio(m[2]) if m[2]==re.search(r"'FICHA'!\$[A-Z]+\$\d+",expected[0]).group() else m[2])+')',expected[0])
        # Family slot identities survive; their cards move after the smaller
        # training section. Derive the exact relocation, never ignore a field.
        fm=re.fullmatch(r'custom.familia.(\d+).(nome|estado|mestre_viu|descricao)',key)
        if fm:
            g=cc.familia(int(fm[1]));new_ref=cc.ref(g[fm[2]],'CONFIGURAÇÃO')
            old_ref=re.search(r"'CONFIGURAÇÃO'!\$[A-Z]+\$\d+",expected[0]).group()
            expected[0]=expected[0].replace(old_ref,new_ref)
            expected[1]=expected[1].replace(old_ref.replace('$',''),new_ref.replace('$',''))
        if current.get(key)!=expected: dif.append(key)
    check('contrato anterior retido preserva sentido, fontes e partes por chave',not dif,dif[:5])
    keys=list(current)
    check('7356 chaves únicas; 131 treinos substituídos por tipos explícitos',len(keys)==7356 and not any('custom.treino.' in k for k in keys))
    cfg=tabs['CONFIGURAÇÃO']; occupied=set(); overlaps=[]
    for region in cfg['mescladas']:
        c1,r1,c2,r2=range_boundaries(region)
        for r in range(r1,r2+1):
            for c in range(c1,c2+1):
                if (r,c) in occupied:overlaps.append(region)
                occupied.add((r,c))
    check('mesclagens de todos os blocos sem sobreposição',not overlaps,overlaps[:4])
    check('cadastros ocultos e fora da paleta',all(tabs[s]['estado']=='hidden' and tabs[s]['paleta_fixa'] for s in [cp.MEL,cp.RES]))
    literal='="COUNTIF($A2:$D2,""Rápido"")"'
    check('texto literal e contagem de vagas não são canonizados',
          cp.canoniza_contagens(literal,'$A2:$D2','$F2:$I2',{'Rápido'})==literal and
          cp.canoniza_contagens('=COUNTIF($A2:$D2,"?*")','$A2:$D2','$F2:$I2',{'Rápido'})=='=COUNTIF($A2:$D2,"?*")')
    check('identidade da base conserva a regra nomeada',cp.canoniza_contagens('=COUNTIF($A2:$D2,"Rápido")','$A2:$D2','$F2:$I2',{'Rápido'})=='=COUNTIF($F2:$I2,"Rápido")')
    lookup='=VLOOKUP(A2,\'DADOS_MELHORIAS\'!$O$2:$W$830,3,FALSE)'
    transformed=cp.buscas_literais(lookup,'DADOS_AM',{(cp.MEL,'O2:W830')})
    check('consultas exatas escapam curingas do nome',transformed=='=VLOOKUP('+cc.criterio('A2')+',\'DADOS_MELHORIAS\'!$O$2:$W$830,3,FALSE)')
    regs={(g['tipo'],g['n']):g for g in model['registros']}
    check('últimos cartões são editáveis e exportados',all(f'custom.{kind}.{n}.nome' in keys for kind,n in [('melhoria',30),('restricao',30)]))
    if a.estrutura:
        sys.exit(not all(x['ok'] for x in checks))

    # Um conjunto de cadastros cobre bases publicadas, criação do zero, erros,
    # curingas e os últimos espaços. As cartas pares usam o livro diretamente.
    wb=Workbook();wb.remove(wb.active)
    for s in tabs: wb.create_sheet(s)
    keep={cp.MEL,cp.RES,cc.NOME}
    for s in keep:
        for c,v in raw[s].items():
            if v is not None:wb[s][c]=v
    for c,v in raw['CONFIGURAÇÃO'].items():
        if v is not None and not (isinstance(v,str) and v.startswith('=')):
            wb['CONFIGURAÇÃO'][c]=v
    for s in [fa.DADOS_AM,'DADOS_INVOC']:
        meta=model['consumidores'][s]
        names=['fam','classe','formas','alcance','pecas','res','feit','carta_zero'] if s==fa.DADOS_AM else ['classe','formas','alcance','pecas','res','fichas','cartas']
        selected=set(raw[s])
        for name in names:
            c,r,rr,w=meta['tabelas'][name]
            for row in range(1,rr+1):
                for col in range(c,c+w): selected.add(cp.coord(row,col))
        prior_cols=next(t['colunas'] for t in old['layout']['abas'] if t['nome']==s)
        selected.update(c for c in raw[s] if ix._lc(c)[1]>prior_cols)
        for c in selected:
            v=raw[s].get(c)
            if v is not None:wb[s][c]=v
        # Fixed context keeps this test focused on the edited assembly rules.
        for key,rr in meta['globais'].items():
            c=rr.rsplit('!',1)[-1].replace('$','')
            wb[s][c]=dict(nível=30,maestria=4,refino=10,essência=5,rota=1,
                **{'maior classe':7,'liberações':3,'classe 0 quantos':5,'classe 0 dados':3,'classe do índice':1}).get(key,0)
    # Rota has a separate one-cell table.
    meta=model['consumidores'][fa.DADOS_AM]
    wb[fa.DADOS_AM][cp.coord(2,meta['tabelas']['rota'][0])]=1
    families=list(fa.regras()['familias'])
    for i,(col,_) in enumerate(fa.geometria()['cols_fam']):
        wb[fa.NOME][col+str(fa.geometria()['familias']+1)]='Livre' if i<2 else 'Neutra'
    fg=cc.familia(1)
    for key,value in [('nome','Família própria'),('estado','Livre')]:wb['CONFIGURAÇÃO'][fg[key]]=value
    def reg(kind,n,**values):
        g=regs[kind,n]
        for key,value in values.items():wb['CONFIGURAÇÃO'][g[key]]=value
    aliases=[('Pressa própria','Rápido'),('Distância própria','Longe'),('Distância ampla','Muito Longe'),
             ('Resposta própria','Reação'),('Alma própria','Toca a Alma'),('Acerto próprio','Inescapável'),
             ('Prisão própria','Prende')]
    for i,(name,base) in enumerate(aliases,1):reg('melhoria',i,nome=name,base=base)
    reg('melhoria',8,nome='Lua X',peso='Leve')
    reg('melhoria',9,nome='Lua*',peso='Pesada',familia='Família própria',exige_tr='Sim')
    reg('melhoria',10,nome='Duplicada',peso='Leve')
    reg('melhoria',11,nome='duplicada',peso='Leve')
    reg('melhoria',12,nome='Peso inválido',peso='muito')
    reg('melhoria',13,nome='Base inválida',base='Inexistente',peso='Leve')
    reg('melhoria',14,efeito='Sem nome')
    reg('melhoria',15,nome='Longe',peso='Leve')
    reg('melhoria',16,nome='Checkbox inválido',peso='Leve',mestre_viu=1)
    reg('melhoria',17,nome='Controle herdado',base='Prende',familia='Nenhuma')
    reg('melhoria',18,nome='Sem peso')
    reg('melhoria',19,nome='Rajada própria',base='Rajada')
    reg('melhoria',30,nome='Última Melhoria',base='Reação')
    reg('restricao',1,nome='Demora própria',base='Atrasar')
    reg('restricao',2,nome='Frequência própria',devolucao='Leve',frequencia='Sim')
    reg('restricao',3,nome='Segunda frequência',devolucao='Média',frequencia='Sim')
    reg('restricao',4,nome='Frequência inválida',devolucao='Leve',frequencia='Talvez')
    reg('restricao',30,nome='Última Restrição',base='Tudo ou Nada')
    # Compare the actual card calculations, with and without renamed bases.
    pairs=[]
    for i,(alias,base) in enumerate(aliases):
        for d,name in enumerate([alias,base]):
            row=2+2*i+d;cel=meta['lugares'][row-2][3]
            for key,value in [('nome','Teste'),('classe',4),('forma','Projétil')]:wb[fa.NOME][cel[key]]=value
            wb[fa.NOME][cel['mel'][0]]=name
        pairs.append((2+2*i,3+2*i,alias))
    # Alias for a frequency restriction must combine with the same incompatibility.
    for row,restrict in [(16,'Demora própria'),(17,'Atrasar')]:
        cel=meta['lugares'][row-2][3]
        for key,value in [('nome','Teste'),('classe',4),('forma','Projétil')]:wb[fa.NOME][cel[key]]=value
        wb[fa.NOME][cel['mel'][0]]='Rápido';wb[fa.NOME][cel['res'][0]]=restrict
    pairs.append((16,17,'Demora própria'))
    for row,pieces,restrictions in [(18,['Lua*'],['Frequência própria','Segunda frequência']),
                                     (19,['Peso inválido'],[])]:
        cel=meta['lugares'][row-2][3]
        for key,value in [('nome','Teste'),('classe',4),('forma','Explosão')]:wb[fa.NOME][cel[key]]=value
        for key,values in [('mel',pieces),('res',restrictions)]:
            for c,v in zip(cel[key],values):wb[fa.NOME][c]=v
    # Entity cards and menus have their own family openings and discounts.
    im=model['consumidores']['DADOS_INVOC']
    for i,g in enumerate(im['fichas']):
        for j in range(9):
            c=cp.coord(i+2,im['colf']['f'+str(j)])
            wb['DADOS_INVOC'][c]=families[j] if i<2 else ''
        for field,value in [('n',30),('cl',7),('db',3),('atq',6),('cd',16)]:
            wb['DADOS_INVOC'][cp.coord(i+2,im['colf'][field])]=value
    # Ninth slot opens the new Family, without making it a free position.
    wb['DADOS_INVOC'][cp.coord(2,im['colf']['f8'])]='Família própria'
    wb['DADOS_INVOC'][cp.coord(3,im['colf']['f8'])]='Família própria'
    iv_pairs=[]
    for i,(alias,base) in enumerate(aliases):
        pair=[]
        for entity,name in enumerate([alias,base]):
            row,card=next((j+2,c) for j,c in enumerate(im['cartas']) if c[0]==entity and c[1]=='esp' and c[2]==i)
            cel=card[3];pair.append(row)
            for key,value in [('nome','Teste'),('classe',4),('forma','Projétil')]:wb['INVOCAÇÕES'][cel[key]]=value
            wb['INVOCAÇÕES'][cel['mel'][0]]=name
        iv_pairs.append((*pair,alias))
    # Class 0 also uses the base for its distance ladder.
    for i,name in enumerate(['Distância própria','Longe']):
        cel=meta['zero'][i]
        wb[fa.NOME][cel['nome']]='Teste';wb[fa.NOME][cel['forma']]='Projétil';wb[fa.NOME][cel['mel']]=name

    if a.preparar_nativo:
        # Só as entradas temporárias diferentes do modelo. A restauração nativa
        # exige uma leitura dos valores originais antes de aplicar este cenário.
        entradas={s:{} for s in tabs}
        for ws in wb:
            for cell in ws._cells.values():
                if cell.value!=raw[ws.title].get(cell.coordinate):
                    entradas[ws.title][cell.coordinate]=cell.value
        a.preparar_nativo.mkdir(parents=True,exist_ok=True)
        (a.preparar_nativo/'entradas.json').write_text(json.dumps({s:v for s,v in entradas.items() if v},ensure_ascii=False,indent=2))

    # XLSX compatibility spelling only, not a change to production formulas.
    for ws in wb:
        for cell in list(ws._cells.values()):
            if isinstance(cell.value,str) and cell.value.startswith('='):
                cell.value=re.sub(r'(?<![A-Z_.])(TEXTJOIN|IFS|SWITCH)\(',r'_xlfn.\1(',cell.value)
    with tempfile.TemporaryDirectory(prefix='b47-pecas-') as td:
        tmp=Path(td);out=tmp/'out';out.mkdir();path=tmp/'pecas.xlsx';wb.save(path);wb.close()
        run=subprocess.run([shutil.which('libreoffice') or 'soffice','-env:UserInstallation='+(tmp/'profile').as_uri(),
            '--headless','--convert-to','xlsx','--outdir',str(out),str(path)],text=True,capture_output=True,timeout=600)
        if run.returncode or not (out/path.name).exists():raise RuntimeError(run.stdout+run.stderr)
        result=load_workbook(out/path.name,data_only=True,read_only=True)
        values={s:{c.coordinate:c.value for row in result[s] for c in row if c.value is not None}
                for s in [fa.DADOS_AM,'DADOS_INVOC',cp.MEL,cp.RES]}
        result.close()
    consultas={}
    def value(s,c):
        v=values[s].get(c,'')
        consultas.setdefault(s,{})[c]=v
        return v
    nbook=model['limites'][cp.MEL]['livro']
    def mel(n,c):return value(cp.MEL,c+str(nbook+n+1))
    for n in [1,2,3,4,5,6,7,8,9,17,30]:check('Melhoria válida '+str(n),value(cp.MEL,'M'+str(n+1))=='',value(cp.MEL,'M'+str(n+1)))
    for n in [10,11,12,13,14,15,16,18]:check('cadastro inválido excluído '+str(n),value(cp.MEL,'M'+str(n+1))!='' and mel(n,'O')=='')
    check('criação do zero tem peso, desconto, TR e identidade própria',
        [mel(9,c) for c in ['O','P','Q','R','W','X']]==['Lua*','Família própria',3,6,1,'Lua*'],[mel(9,c) for c in ['O','P','Q','R','W','X']])
    check('base Controle conserva controle e TR',
        [mel(17,c) for c in ['P','T','W']]==['',1,1],[mel(17,c) for c in ['P','T','W']])
    check('Rajada herda classe reduzida da entidade',mel(19,'Y')==1,mel(19,'Y'))
    mr=model['limites'][cp.RES]['livro']
    check('última Restrição herda a identidade',value(cp.RES,'N'+str(mr+31))=='Tudo ou Nada')
    check('frequência inválida fica fora do catálogo',value(cp.RES,'K5')!='' and value(cp.RES,'M'+str(mr+5))=='')
    check('nona Família abre o menu só das entidades que a escolheram',value(cp.MEL,'AN'+str(nbook+10))=='Lua*' and value(cp.MEL,'AR'+str(nbook+10))=='')
    fc=meta['col']
    check('caixas vazias não cobram cadastros excluídos do catálogo',
        all(value(fa.DADOS_AM,cp.coord(18,fc['k'+str(k)]))==0 for k in [2,3,4]))
    check('duas Restrições de frequência são sinalizadas','frequência' in str(value(fa.DADOS_AM,cp.coord(18,fc['erros']))))
    check('devolução paga só Melhorias e conserva o custo da Forma',
        value(fa.DADOS_AM,cp.coord(18,fc['usa']))==4 and value(fa.DADOS_AM,cp.coord(18,fc['dv']))==6,
        {k:value(fa.DADOS_AM,cp.coord(18,fc[k])) for k in ['classe','k0','k1','k2','k3','k4','usa','dv','g','s']})
    check('carta sinaliza peça ausente ou cadastro inválido','não cadastrada ou com erro' in str(value(fa.DADOS_AM,cp.coord(19,fc['erros']))))
    norm=lambda x: '' if x is None else x
    def normalized(x):
        if not isinstance(x,str):return norm(x)
        for alias,base in aliases+[('Demora própria','Atrasar')]:x=x.replace(alias,base)
        return x
    fields=['k1','ct','ptr','g','dv','usa','s','d','dep','erros','acao','resolve','lg','mr','q1','q2','alcance']
    for s,pairlist in [(fa.DADOS_AM,pairs),('DADOS_INVOC',iv_pairs)]:
        m=model['consumidores'][s]
        for r1,r2,label in pairlist:
            mismatches=[(k,value(s,cp.coord(r1,m['col'][k])),value(s,cp.coord(r2,m['col'][k]))) for k in fields
                        if normalized(value(s,cp.coord(r1,m['col'][k])))!=normalized(value(s,cp.coord(r2,m['col'][k])))]
            errors=[(k,value(s,cp.coord(r1,m['col'][k]))) for k in fields if str(value(s,cp.coord(r1,m['col'][k]))).startswith(('#REF!','#VALUE!','#NAME?','#N/A'))]
            check(s+' herda '+label,not mismatches and not errors,(mismatches+errors)[:4])
    z=meta['tabelas']['carta_zero'][0]+7
    zero_values=[value(fa.DADOS_AM,cp.coord(r,z)) for r in [2,3]]
    check('Classe 0 herda alcance da Melhoria',zero_values[0]==zero_values[1] and 'm' in str(zero_values[0]),zero_values)
    summary=dict(resultado='OK' if all(c['ok'] for c in checks) else 'FALHOU',checagens=checks)
    if a.preparar_nativo:
        (a.preparar_nativo/'esperados.json').write_text(json.dumps(consultas,ensure_ascii=False,indent=2))
    if a.evidencia:a.evidencia.write_text(json.dumps(summary,ensure_ascii=False,indent=2))
    sys.exit(not all(c['ok'] for c in checks))


if __name__=='__main__':main()
