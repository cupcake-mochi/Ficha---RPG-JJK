#!/usr/bin/env python3
"""Inventário somente de leitura das fórmulas geradas para a parte 2 da B47.

Lê o artefato interno produzido por monta.py; não importa/exporta no Sheets.
Varre TODAS as fórmulas com o tokenizer, que distingue números de endereços,
texto e nomes de função. A associação semântica às regras é explícita e
revisável: números não classificados não são automaticamente regras.
Inclui valores fixos que não aparecem em fórmulas (nível e tabela de XP).
Não executa o gerador, não recalcula e não altera nenhum arquivo de entrada.
"""
import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer

RAIZ = Path(__file__).resolve().parent


def numeros(formula):
    return [t.value for t in Tokenizer(formula).items
            if t.type == 'OPERAND' and t.subtype == 'NUMBER']


def regra(chave, titulo, livro, celulas, trecho, fonte, observacao=''):
    return dict(chave=chave, regra=titulo, livro=livro, celulas=celulas,
                trecho_livro=trecho, gerador=fonte, observacao=observacao)


REGRAS = [
    regra('nivel_inicial', 'Nível inicial', '2', ['FICHA!AH11'],
          'O padrão deste livro é nível 2', 'ficha-v01/layout.json; apps-script/Codigo.gs: nivelDoXp_, nivelQueSobe_, nivelPelaXP_',
          'Entrada fixa, não fórmula. Decidir o tratamento do XP de quem começa acima de 2; o livro admite nível 1, mas a tabela de custos começa em 2.'),
    regra('atributos_pontos', 'Pontos de atributo iniciais', '9', ['DADOS!AK6', 'FICHA!Z15'],
          'Distribua 9 pontos', 'ficha-v01/ficha_automatica.py: pontos_criacao / pontos de atributo',
          'Preservar os pontos adicionais dos marcos.'),
    regra('atributos_teto_criacao', 'Teto de atributo na criação', '3', ['FICHA!Z15'],
          'Nenhum pode começar abaixo de 0 ou acima de 3', 'ficha-v01/ficha_automatica.py: teto_criacao',
          'Hoje a verificação de criação vale enquanto nenhum marco foi alcançado; não depende de um nível inicial configurável.'),
    regra('atributos_teto', 'Teto de atributo na progressão', '6', ['FICHA!Z15'],
          'O máximo de cada atributo é 6', 'ficha-v01/ficha_automatica.py: teto_atributo',
          'Proposta de separar do teto da criação; depende de aprovação.'),
    regra('xp_curva', 'Curva de XP', 'Tabela por nível (2–30)',
          ['DADOS!AR'+str(r) for r in range(5,34)],
          'Nível atual | XP para o próximo', 'ficha-v01/ficha_automatica.py: tab_prog; apps-script/Codigo.gs: nivelPelaMissao_, nivelPelaXP_',
          'AR guarda XP acumulado, não custo para o próximo. Fonte do custo: catalogo-projeto-m.json progressao.tabela_impressa. Não confundir com recompensa/desconto de missão.'),
    regra('pericias_com_oficio', 'Perícias com ofícios', '9', ['DADOS!AK17','DADOS!AK19','DADOS!AK25'],
          'a criação reúne nove perícias e dois ofícios', 'ficha-v01/ficha_automatica.py: pericias_com',
          'Total inclui treinos da Origem e Caminho. Precisamos decidir se a alteração concede apenas escolhas livres adicionais.'),
    regra('oficios_criacao', 'Ofícios na criação', '2', ['DADOS!AK17','DADOS!AK20','DADOS!AK26'],
          'a criação reúne nove perícias e dois ofícios', 'ficha-v01/ficha_automatica.py: oficios_com'),
    regra('pericias_sem_oficio', 'Perícias sem ofícios', '10', ['DADOS!AK17','DADOS!AK19','DADOS!AK27','DADOS!AK28'],
          'ou dez perícias sem ofícios', 'ficha-v01/ficha_automatica.py: pericias_troca',
          'A alternativa tem zero ofícios; preservar a inferência da rota e as trocas por arma existentes.'),
    regra('deslocamento_base', 'Deslocamento base', '9 m', ['FICHA!O67'],
          'O deslocamento padrão é 9 m', 'ficha-v01/layout.json; ficha-v01/ficha_pessoal.py: metros; ficha-v01/configuracao.py',
          'Só o 9 é o valor base. 4,5 é limite da Exaustão; 2/1 são fatores; 9999 é sentinela. Preservar carga, Integridade, Buff e efeitos.'),
    regra('morrendo_aguentar', 'Morrendo · Aguentar', '1/4, para cima, sem teto', ['FICHA!D41','FICHA!D42'],
          'Aguentar | Fica Inconsciente', 'ficha-v01/ficha_pessoal.py: morrendo / frac(4)',
          'Atualizar a legenda junto da conta.'),
    regra('morrendo_insistir_1', 'Morrendo · Insistir na primeira queda', '1/3, para cima', ['FICHA!O41','FICHA!O42'],
          'Insistir pela primeira vez na cena |', 'ficha-v01/ficha_pessoal.py: morrendo / frac(3)'),
    regra('morrendo_insistir_2', 'Morrendo · Insistir nas seguintes', '1/2, para cima', ['FICHA!Z41','FICHA!Z42'],
          'Insistir pela segunda vez ou mais na cena |', 'ficha-v01/ficha_pessoal.py: morrendo / frac(2)'),
    regra('morrendo_teto', 'Morrendo · teto de Insistir', '15 × maior Classe', ['FICHA!O41','FICHA!Z41','FICHA!O42','FICHA!Z42'],
          'Se o resultado superar 15 × sua maior Classe', 'ficha-v01/ficha_pessoal.py: teto'),
    regra('morrendo_prazo', 'Morrendo · prazo', '3 rodadas, menos 1 por Sequela', ['FICHA!AK41','FICHA!AK42'],
          'Sequelas antes de cair | Janela de Aguentar ou Insistir', 'ficha-v01/ficha_pessoal.py: morrendo prazo',
          'Há três caixas de Sequelas. Mudar o prazo para mais de três exige decidir como registrar as adicionais; dano também reduz a janela na mesa.'),
    regra('vida_rolada', 'Vida rolada', 'Opcional; padrão usa valor fixo do Caminho', ['FICHA!J26'] + ['DADOS!'+c+str(r) for c in ['P','Q'] for r in range(5,11)],
          'Como variante, com permissão do mestre, role o dado do Caminho', 'ficha-v01/dados_catalogo.py: e_dos_caminhos; ficha-v01/ficha_automatica.py: vida; ficha-v01/configuracao.py',
          'A fórmula atual não recebe rolagens. Criar registro persistente por nível (proposta); não usar ALEATÓRIO. O dado de cada Caminho existe no manual, mas não na tabela atual P:Q.'),
    regra('morte_permanente', 'Morte permanente', 'Não como regra padrão; campanha pode adotar', [],
          'Na regra de campanha padrão, derrota não causa morte automática', 'Novo registro em CONFIGURAÇÃO / ANALISE, sem conta existente',
          'CONFLITO: PLANO-CONFIGURACAO.html diz sim. Pedir decisão antes de preencher LIVRO/CAMPANHA. Pode afetar a redação do prazo se isso for aprovado.'),
]


def levantar(arquivo):
    manual = (RAIZ/'manual.txt').read_text().splitlines()
    fonte = json.loads((RAIZ/'manual-fonte.json').read_text())
    cat = json.loads((RAIZ/'catalogo-projeto-m.json').read_text())
    wb = load_workbook(arquivo, read_only=True, data_only=False)
    celulas, formulas, erros = {}, [], []
    try:
        for sheet in wb:
            for row in sheet.iter_rows():
                for c in row:
                    if c.value is None:
                        continue
                    endereco = f'{sheet.title}!{c.coordinate}'
                    valor = c.value
                    if c.data_type == 'f' and not isinstance(valor, str):
                        valor = getattr(valor, 'text', None)
                        if not isinstance(valor, str):
                            raise ValueError(f'Fórmula não suportada: {endereco}')
                    celulas[endereco] = valor
                    if c.data_type == 'f':
                        try:
                            ns = numeros(valor)
                        except Exception as e:
                            erros.append(dict(celula=endereco, erro=str(e)))
                            continue
                        formulas.append(dict(celula=endereco, formula=valor, numeros=ns))
    finally:
        wb.close()
    if erros:
        raise ValueError(f'Não foi possível analisar todas as fórmulas: {erros[:5]}')
    regras = []
    associacoes = {}
    for original in REGRAS:
        r = dict(original)
        linhas = [i for i,s in enumerate(manual,1) if r['trecho_livro'] in s]
        if not linhas:
            raise ValueError(f'Fonte do livro mudou: {r["chave"]}')
        r['linhas_livro'] = linhas
        r['ocorrencias'] = []
        for endereco in r['celulas']:
            if endereco not in celulas:
                raise ValueError(f'Célula ausente; atualizar mapa: {endereco}')
            valor = celulas[endereco]
            r['ocorrencias'].append(dict(celula=endereco, valor=valor))
            associacoes.setdefault(endereco,[]).append(r['chave'])
        regras.append(r)
    for f in formulas:
        f['regras'] = associacoes.get(f['celula'], [])
    return dict(
        fonte_livro=fonte, manual_sha256=hashlib.sha256((RAIZ/'manual.txt').read_bytes()).hexdigest(),
        artefato=str(arquivo), artefato_sha256=hashlib.sha256(arquivo.read_bytes()).hexdigest(),
        resumo=dict(formulas=len(formulas), formulas_com_numeros=sum(bool(f['numeros']) for f in formulas),
                    formulas_associadas=sum(bool(f['regras']) for f in formulas), regras=len(regras),
                    celulas_associadas=len(associacoes), erros_de_tokenizacao=len(erros)),
        limites='Varredura lexical completa das fórmulas do artefato atual; associação semântica explícita somente para a lista B47. Números de argumentos/índices não são automaticamente regras. Condicionais, validações e textos não associados ainda exigem revisão na implementação.',
        regras=regras,
        xp_custos={n:row['xp'] for n,row in cat['progressao']['tabela_impressa'].items()},
        formulas=formulas)


def markdown(d):
    s=['# Levantamento das regras de campanha — B47, parte 2','',
       'Proposta para aprovação. A ficha e suas fórmulas não foram alteradas.', '',
       'Fonte: manual.txt do gerador, identificado por manual-fonte.json como R51. O texto antigo de outra pasta não foi usado para decidir números.', '',
       'São 16 parâmetros propostos para detalhar os grupos aprovados no plano. A separação dos tetos e dos valores de perícia/Morrendo precisa de aprovação; ainda não determina o número de linhas do protótipo.', '',
       '| Regra | Livro R51 | Células atuais |', '|---|---|---|']
    for r in d['regras']:
        coords = ', '.join(r['celulas']) if len(r['celulas'])<8 else ', '.join(r['celulas'][:2])+f' … ({len(r["celulas"])} células; lista completa no JSON)'
        s.append(f'| {r["regra"]} | {r["livro"]} | {coords or "Sem campo atual"} |')
    s += ['', '## Fontes e dependências', '']
    for r in d['regras']:
        s += [f'### {r["regra"]}', f'- Livro: manual.txt, linhas {", ".join(map(str,r["linhas_livro"]))}.',f'- Gerador: {r["gerador"]}.']
        if r['observacao']:s.append(f'- {r["observacao"]}')
    s += ['', '## Decisões antes do protótipo', '',
          '1. **Morte permanente:** proposta: LIVRO = Não; CAMPANHA em branco usa Não e pode receber Sim. O plano diz Sim, mas o R51 diz que derrota não mata automaticamente. Confirmar qual comportamento deseja.',
          '2. **Tetos e perícias:** proposta: dois tetos independentes (criação e progressão) e três totais de treino (perícias com ofícios, ofícios, perícias sem ofícios). As perícias fixas da Origem/Caminho continuam fixas; mudanças no total alteram as escolhas livres. Confirmar.',
          '3. **XP:** proposta para desenho: tabela de custo por nível, com Livro / Campanha / Em uso, recolhível abaixo das regras. Um multiplicador sozinho não permite qualquer curva. Decidir também se XP zero corresponde ao nível inicial configurado; o custo do nível 1 para 2 não aparece na tabela do R51 e precisa ser informado caso essa opção seja usada.',
          '4. **Vida rolada:** proposta: valor inicial do Caminho preservado e um resultado anotado por nível posterior, somando Constituição normalmente. Confirmar a interpretação da variante antes de criar os campos. O gerador deve avisar quando faltar rolagem, em vez de inventar um resultado.',
          '5. **Prazo:** proposta: manter desconto de uma rodada por Sequela. Se o prazo puder superar três, precisamos ampliar o registro de Sequelas; hoje há só três caixas. Definir essa extensão no protótipo.',
          '', '## Cuidados da futura implementação', '',
          '- CAMPANHA vazia deve conservar os resultados atuais; incluir avisos, textos, notas e comparação com a parte 1.',
          '- Na vida e no deslocamento, a configuração da regra entra na conta base; a prioridade já aprovada de trocar por continua fixando o resultado final.',
          '- A subida automática por XP hoje depende do onEdit em Codigo.gs. Conferir na cópia nativa o requisito de funcionar só ao copiar, preservando o nível editável; não substituir por fórmula nem mudar esse comportamento sem combinar.',
          '- A ficha usa a vida máxima atual para o Morrendo. O R51 fala em máximo de referência da queda; conferir como registrar a referência em mesa se o máximo mudar depois da queda. Não foi alterado nesta etapa.',
          '- Maestria, TRs, Refino, carga, recompensas de missão e valores de técnica também contêm números. A varredura os registra, mas não amplia a lista aprovada de regras de campanha.',
          '', '## Alcance da varredura', '', d['limites'], '',
          json.dumps(d['resumo'],ensure_ascii=False), '',
          'O JSON guarda as fórmulas completas e seus literais numéricos, inclusive as não classificadas. A etapa não executa monta.py nem muda o contrato ANALISE 3.', '']
    return '\n'.join(s)


def teste():
    assert numeros('=IF(A12>3,"15 × Classe",ROUNDUP($J$26/4,0))') == ['3','4','0']
    assert numeros("=VLOOKUP(A1,'ABA 2'!$B$4:$E$30,3,FALSE)+1.5") == ['3','1.5']
    assert numeros('=LOG10(100)+1E-3') == ['100','1E-3']
    assert numeros('="texto ""12"""&A99') == []
    print('OK: números separados de referências, texto, funções e notação científica.')


if __name__ == '__main__':
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--arquivo', type=Path, default=RAIZ/'ficha-v01/ficha-projeto-m-0.1.xlsx')
    p.add_argument('--saida', type=Path)
    p.add_argument('--teste', action='store_true')
    a=p.parse_args()
    if a.teste: teste()
    else:
        if not a.saida: p.error('--saida é obrigatório para guardar a evidência')
        d=levantar(a.arquivo.resolve())
        a.saida.mkdir(parents=True,exist_ok=True)
        (a.saida/'regras-campanha.json').write_text(json.dumps(d,ensure_ascii=False,indent=2)+'\n')
        (a.saida/'REGRAS-CAMPANHA.md').write_text(markdown(d))
        print(json.dumps(d['resumo'],ensure_ascii=False))
