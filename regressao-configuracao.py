# -*- coding: utf-8 -*-
"""Recalcula a ficha real para testar desligado, somas, trocas, erros e conflitos.

As expectativas das mudanças são aritmética independente; não executam o
gerador nem avaliam suas fórmulas em Python. LibreOffice roda uma vez, serial.
"""
import hashlib
import json
import math
import os
import re
import runpy
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent
os.chdir(RAIZ)
ARQ = RAIZ / "ficha-v01/ficha-projeto-m-0.1.xlsx"
FALHAS = []


def checa(nome, certo, detalhe=""):
    print(f"[{'OK' if certo else 'FALHA'}] {nome}" + ("" if certo else f": {detalhe}"), flush=True)
    if not certo:
        FALHAS.append(nome)


def numero(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return v
    if isinstance(v, str) and v.startswith("d20 + "):
        return float(v[6:].replace(",", "."))
    if isinstance(v, str) and v.endswith(" m"):
        return float(v[:-2].replace(",", "."))      # o Calc em pt-BR escreve 9,5 m
    return v


leitor = runpy.run_path(str(RAIZ / "ler-analise.py"))
checa("o leitor normaliza decimais de Excel e Sheets e preserva a versão textual da ficha",
      leitor["normaliza"]("4.5") == 4.5 and leitor["normaliza"]("-1,5") == -1.5
      and leitor["de_linhas"]([["ficha.versao", "1.0", "", "cabeçalho"]])["ficha.versao"]["valor"] == "1.0")
wb = load_workbook(ARQ)
an = [[c.value for c in r] for r in wb["ANALISE"].iter_rows(min_row=2, max_col=4)]
antigas = [list(r) for r in an if r[0] and not r[0].startswith(("efeito.", "aviso.efeito.", "aviso.efeito_caixa.", "regra.", "vida.rolagem.", "aviso.regras_campanha", "aviso.vida_rolada", "sequela.extras", "custom.", "aviso.custom."))]
for r in antigas:
    if r[0] == "formato.versao": r[1] = "2"
canon = json.dumps(antigas, ensure_ascii=False, separators=(",", ":"))
snapshot = json.loads((RAIZ / "ficha-v01/analise-v2-preservada.json").read_text())
# 17/10/2026: a fotografia do formato 2 deixou de valer quando a fase A retirou chaves de propósito (capacidade
# reduzida) e C levou ao formato 11. A preservação entre versões está provada nas migrações de C
# (conteudos-b47/comum/MIGRACAO-C-DESCARTAVEL.json e MIGRACAO-C-REVISAO.json); aqui fica o contrato atual.
checa("a ANALISE está no formato 11, com as 8.190 chaves do contrato B47",
      an[0][0:2] == ["formato.versao", "11"] and len(an) == 8190, f"{an[0][0:2]} · {len(an)} linhas")
checa("nenhuma chave foi repetida", len(an) == len({r[0] for r in an}))
checa("os seis campos de cada efeito são exportados separadamente", sum(str(r[0]).startswith("efeito.") for r in an) == 180)
checa("o formato atual declara regras de campanha para o bot",
      an[0][0:2] == ["formato.versao", "11"] and any(str(r[0]).startswith("regra.") for r in an))

alvos = []
for r in wb["DADOS_CONFIG"].iter_rows(min_row=2, max_col=8):
    nome, origem = r[0].value, r[7].value
    if not nome: continue
    m = re.fullmatch(r"='([^']+)'!\$([A-Z]+)\$(\d+)", origem or "")
    assert m, (nome, origem)
    alvos.append((nome, m[1], m[2] + m[3], r[0].row))
checa(f"cada uma das {len(alvos)} caixas tem destino próprio (eram 50 antes de C)",
      len(alvos) >= 50 and len({a[1:3] for a in alvos}) == len(alvos), str(len(alvos)))

# A geometria da FICHA é lida do índice, como nos outros testes.
sys.path.insert(0, str(RAIZ / "ficha-v01"))
import indice_ficha as ix
idx = {wb["DADOS"].cell(r, 53).value: ix.endereco(wb["DADOS"].cell(r, 54).value)
       for r in range(5, 200) if wb["DADOS"].cell(r, 53).value}
f = wb["FICHA"]
f[idx["caminho"]] = "Vanguarda"
f[idx["nivel"]] = 10
for nome, valor in zip(("Força", "Destreza", "Constituição", "Inteligência", "Essência"), (3, 2, 4, 1, 3)):
    f[idx["atr_base_" + nome]] = valor
for campo, valor in (("buff de defesa", 2), ("buff de vida", 5), ("buff de energia", 1), ("buff de deslocamento", 0.5)):
    f[idx[campo]] = valor
for ws in wb:
    ws._images = []
    for c in tuple(ws._cells.values()):
        if isinstance(c.value, str) and c.value.startswith("="):
            c.value = re.sub(r"(?<![A-Z_.])IFS\(", "_xlfn.IFS(", c.value)

def limpar():
    for r in range(11, 41):
        for col in "BCDEFG": wb["CONFIGURAÇÃO"][f"{col}{r}"] = False if col == "B" else None


def efeito(n, caixa, como, valor, ligado=True, tipo="Decisão do mestre", fonte="Teste"):
    r = 10 + n
    for col, v in zip("BCDEFG", (ligado, caixa, como, valor, tipo, fonte)):
        wb["CONFIGURAÇÃO"][f"{col}{r}"] = v


if not shutil.which("libreoffice"):
    sys.exit("LibreOffice indisponível: a regressão NÃO RODOU")
with tempfile.TemporaryDirectory(prefix="b47-config-") as pasta:
    pasta = Path(pasta)
    arquivos, casos = [], {}
    def salvar(nome, mudancas):
        p = pasta / (nome + ".xlsx")
        wb.save(p)
        arquivos.append(str(p))
        casos[nome] = mudancas
    limpar()
    salvar("desligado", {})
    for lote in range(2):
        limpar()
        mudancas = {}
        for n, alvo in enumerate(alvos[lote * 25:(lote + 1) * 25], 1):
            valor = (2, -1, 0.5)[(n - 1) % 3]
            efeito(n, alvo[0], "somar", valor)
            mudancas[alvo[0]] = ("somar", valor)
        salvar(f"somas{lote + 1}", mudancas)
    limpar()
    nomes = ("Vida máxima", "PE máximo", "Integridade máxima", "Defesa", "Proteção", "Iniciativa",
             "CD de feitiço", "Conjuração", "Corpo a corpo", "À distância", "Deslocamento", "Limite de carga",
             "Perícia · Furtividade", "Ofício · Condução", "TR · Físico")
    mudancas = {}
    for n, nome in enumerate(nomes, 1):
        valor = (0, -2, 4.5)[(n - 1) % 3]
        efeito(2 * n - 1, nome, "somar", 99)
        efeito(2 * n, nome, "trocar por", valor)
        mudancas[nome] = ("trocar por", valor)
    salvar("trocas", mudancas)
    limpar()
    efeito(1, "Defesa", "trocar por", 0)
    efeito(2, "Defesa", "trocar por", 10)
    efeito(3, "Perícia · Furtividade", "somar", "dois")
    efeito(4, "Caixa inexistente", "somar", 2)
    efeito(5, "PE máximo", "somar", 2, fonte="")
    efeito(6, "Vida máxima", "multiplicar", 2)
    efeito(7, "Proteção", "somar", 2, ligado="sim")
    efeito(8, "Caixa inexistente", "errado", "inválido", ligado=False, tipo="errado", fonte="")
    efeito(9, "Caixa inexistente", "errado", "inválido", ligado=None, tipo="errado", fonte="")
    salvar("invalidos", {})
    saida = pasta / "saida"
    perfil = pasta / "perfil"
    p = subprocess.run(["libreoffice", f"-env:UserInstallation={perfil.as_uri()}", "--headless", "--convert-to", "xlsx",
                        "--outdir", str(saida), *arquivos], capture_output=True, text=True, timeout=1800)
    checa("LibreOffice recalculou todos os casos", p.returncode == 0 and all((saida / Path(a).name).exists() for a in arquivos), p.stderr[-500:])
    if p.returncode or not all((saida / Path(a).name).exists() for a in arquivos):
        sys.exit(1)
    for nome, mudancas in casos.items():
        lido = load_workbook(saida / (nome + ".xlsx"), data_only=True, read_only=True)
        # Percorre cada aba uma vez: acesso avulso em read_only relê o XML inteiro.
        valores = {}
        for aba in {a[1] for a in alvos} | {"DADOS_CONFIG"}:
            valores[aba] = {c.coordinate: c.value for linha in lido[aba].iter_rows()
                            for c in linha if c.value is not None}
        le = lambda aba, coord: valores[aba].get(coord)
        if nome == "invalidos":
            for caixa in ("Defesa", "Perícia · Furtividade", "PE máximo", "Vida máxima", "Proteção"):
                alvo = next(a for a in alvos if a[0] == caixa)
                checa(f"{caixa}: erro ligado fica visível, sem certificar zero", le(alvo[1], alvo[2]) == "#N/A")
            checa("duas trocas indicam conflito", "Conflito" in str(le("DADOS_CONFIG", "E5")))
            for n in range(3, 8):
                checa(f"entrada inválida {n} gera aviso", bool(le("DADOS_CONFIG", f"M{n + 1}")))
            checa("um efeito inválido desligado não altera nem acusa", le("DADOS_CONFIG", "M9") in (None, ""))
            checa("apagar a caixa de seleção também deixa o efeito desligado", le("DADOS_CONFIG", "M10") in (None, ""))
            lido.close()
            continue
        for caixa, aba, coord, r in alvos:
            atual = numero(le(aba, coord))
            base = numero(le("DADOS_CONFIG", f"G{r}"))
            modo, valor = mudancas.get(caixa, ("somar", 0))
            esperado = valor if modo == "trocar por" else base + valor if isinstance(base, (int, float)) else base
            bate = math.isclose(atual, esperado, abs_tol=1e-9) if isinstance(atual, (int, float)) and isinstance(esperado, (int, float)) else atual == esperado
            checa(f"{nome}: {caixa}", bate and atual not in ("#REF!", "#VALUE!", "#NAME?", "#N/A", "#DIV/0!"), f"{atual!r} != {esperado!r}")
        lido.close()
print(f"\n{'CONFIGURAÇÃO CONFERE' if not FALHAS else str(len(FALHAS)) + ' FALHAS'}")
sys.exit(bool(FALHAS))
