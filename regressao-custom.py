#!/usr/bin/env python3
"""B47/3a: contrato preservado, geometria e cenários recalculados independentes.

Executar sob o teto de recursos documentado. Não toca Sheets nem Apps Script.
LibreOffice é apenas uma segunda calculadora para as contas escalares; os
menus com FILTER/UNIQUE são conferidos no Sheets nativo, nunca por importação.
"""
import json, math, re, shutil, subprocess, sys, tempfile
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parent
sys.path[:0]=[str(ROOT/'ficha-v01'),str(ROOT)]
import custom,indice_ficha as ix
(ROOT/'nuvem-b47/saidas').mkdir(parents=True,exist_ok=True)
falhas=[]
def check(nome,ok,det=''):
 print(('OK ' if ok else 'FALHA ')+nome+((' '+str(det)) if not ok else ''),flush=True)
 if not ok:falhas.append(nome)
novo=ROOT/'ficha-v01/ficha-projeto-m-0.1.xlsx'
base=ROOT/'nuvem-b47/fixtures/base-parte2-local.xlsx'
w=load_workbook(novo)
rows=[list(x) for x in w['ANALISE'].iter_rows(min_row=2,max_col=4,values_only=True) if x[0]]
old=json.loads((ROOT/'ficha-v01/analise-v4-preservada.json').read_text())
mapa={r[0]:r for r in rows}
check('5274 chaves anteriores preservadas integralmente',all((mapa[r[0]]==r if r[0]!='formato.versao' else mapa[r[0]][1]=='6') for r in old) and len(old)==5274)
check('contrato sem chaves repetidas',len(rows)==len(mapa))
check('25 registros completos inclusive vazios',all(f'custom.{n}.campo.base_do_livro' in mapa and f'aviso.custom.{n}' in mapa for n in range(1,26)))
cfg=w['CONFIGURAÇÃO'];merges=list(cfg.merged_cells.ranges)
check('painel esquerdo preserva bônus, campanha, XP e vida',all(cfg[c].value==v for c,v in [('B9','BÔNUS E MUDANÇAS'),('B43','REGRAS DE CAMPANHA'),('B62','XP · CUSTO POR NÍVEL'),('B97','VIDA ROLADA · REGISTRO')]))
check('cartões sem sobreposição de mesclagens',not any(a.min_col<=b.max_col and b.min_col<=a.max_col and a.min_row<=b.max_row and b.min_row<=a.max_row for i,a in enumerate(merges) for b in merges[i+1:]))
check('contorno da descrição em cada cartão',all(str(cfg[custom.cartao(n)['efeito']].border.left.style)=='thin' and str(cfg[custom.cartao(n)['efeito']].border.right.style)=='thin' for n in range(1,26)))
check('sem faixa cinza de congelamento',cfg.freeze_panes is None)
check('aviso global preserva bônus, campanha e vida rolada',all(x in cfg['J6'].value for x in ['DADOS_CONFIG!$P$2','DADOS_CONFIG!$T$17','DADOS_CONFIG!$AF$32']))
dvs=list(cfg.data_validations.dataValidation)
check('175 seletores de campos fechados estritos',sum(v.showErrorMessage is True and v.type=='list' and str(v.sqref).startswith('L') for v in dvs)==175)
check('25 seletores de tipo estritos da parte 3a',sum(v.showErrorMessage is True and v.type=='list' and str(v.sqref).startswith('J') and int(re.search(r'\d+',str(v.sqref)).group())<=732 for v in dvs)==25)
check('itens R51 têm volume real',next(i for i in custom.itens_livro() if i['nome']=='Corda')['volume']==.5)
check('rotas próprias incluem Marcial com e sem energia',len(custom.ROTAS)==4 and custom.SEM_BASE=='Nenhuma')
idx={w['DADOS'].cell(r,53).value:ix.endereco(w['DADOS'].cell(r,54).value) for r in range(5,w['DADOS'].max_row+1) if w['DADOS'].cell(r,53).value}
w.close()
if '--estrutura' in sys.argv:sys.exit(bool(falhas))
if not shutil.which('libreoffice'):sys.exit('LibreOffice indisponível: cálculo não executado')
with tempfile.TemporaryDirectory(prefix='b47-custom-') as tmp:
 tmp=Path(tmp);inp=tmp/'in';out=tmp/'out';inp.mkdir();out.mkdir()
 def record(edits,n,t,name,closed,free,effect='Efeito de teste'):
  g=custom.cartao(n)
  edits[('CONFIGURAÇÃO',g['tipo'])]=t;edits[('CONFIGURAÇÃO',g['nome'])]=name
  edits[('CONFIGURAÇÃO',g['efeito'])]=effect
  vals=closed+[None]*(7-len(closed))+free+[None]*(6-len(free))
  for c,v in zip(g['campos'],vals):edits[('CONFIGURAÇÃO',c)]=v
 full={}
 record(full,1,'Arma','Arma de teste',['Porrete',1,'Concussão','Versátil','—','—','—'],['1d6',4,1.5,'—','—',0])
 record(full,2,'Escudo','Escudo de teste',[],[4,1,4,2,0])
 record(full,3,'Uniforme','Uniforme de teste',['Grau 4'],[3,2,4,2,0])
 record(full,4,'Item comum','Item de teste',[],[.7,0])
 record(full,5,'Ferramenta amaldiçoada','Ferramenta de teste',['Arma de teste','Grau 2'],['Estigma de teste',2,0,0])
 per=custom.listas()['pericias'][:4]
 record(full,6,'Origem','Origem de teste',per+['Fundamento','Nenhuma','—'],['Talento registrado'])
 record(full,7,'Marcante','Marcante de teste',['Origem de teste'],['Uma vez por cena'],'Texto marcante')
 record(full,8,'Menor','Menor de teste',['Origem de teste','Marcante de teste'],['Uma vez por missão'],'Texto menor')
 record(full,9,'Ferramenta amaldiçoada','Uniforme ferramenta',['Uniforme de teste','Grau 1'],['Marca',3,None,0])
 record(full,10,'Ferramenta amaldiçoada','Acessório ferramenta',['Item de teste','Grau 3'],['Marca',1,.4,0])
 full.update({('FICHA PESSOAL','D92'):'Arma de teste',('FICHA PESSOAL','K92'):2,('FICHA PESSOAL','D93'):'Escudo de teste',
  ('FICHA PESSOAL','D64'):'Arma de teste',('FICHA PESSOAL','S64'):'Escudo de teste',('FICHA PESSOAL','AH64'):'Uniforme de teste',
  ('FICHA PESSOAL','D104'):'Item de teste',('FICHA PESSOAL','T104'):3,('FICHA',idx['origem']):'Origem de teste',
  ('FICHA PESSOAL','D44'):'Marcante de teste',('FICHA PESSOAL','R44'):'Menor de teste'})
 common={('FICHA',idx['caminho']):'Vanguarda',('FICHA',idx['nivel']):10,
  ('FICHA',idx['atr_base_Força']):3,('FICHA',idx['atr_base_Destreza']):3,
  ('FICHA',idx['atr_base_Constituição']):3,('FICHA',idx['atr_base_Inteligência']):1,('FICHA',idx['atr_base_Essência']):2,
  ('FICHA',idx['buff de defesa']):2}
 cases=[('base',base,{}),('vazia',novo,{}),('completo',novo,full),
  ('troca_zero',novo,{**full,('CONFIGURAÇÃO','B11'):True,('CONFIGURAÇÃO','C11'):'Defesa',('CONFIGURAÇÃO','D11'):'trocar por',('CONFIGURAÇÃO','E11'):0,('CONFIGURAÇÃO','F11'):'Decisão do mestre',('CONFIGURAÇÃO','G11'):'Teste'}),
  ('ferramenta',novo,{**full,('FICHA PESSOAL','D92'):'Ferramenta de teste',('FICHA PESSOAL','D64'):'Ferramenta de teste',('FICHA PESSOAL','AH64'):'Uniforme ferramenta',('FICHA PESSOAL','D104'):'Acessório ferramenta',('FICHA PESSOAL','T104'):2}),
  ('marcial_sem',novo,{**full,('CONFIGURAÇÃO',custom.cartao(6)['campos'][4]):'Técnica Marcial sem energia'}),
  ('duplicado',novo,{**full,('CONFIGURAÇÃO',custom.cartao(1)['nome']):'Bastão'}),
  ('marcante_livro',novo,{**full,('CONFIGURAÇÃO',custom.cartao(7)['campos'][0]):'Latente'}),
  ('marcante_restricao',novo,{**full,('CONFIGURAÇÃO',custom.cartao(7)['campos'][0]):custom.RESTRICAO}),
  ('base_latente',novo,{**full,('CONFIGURAÇÃO',custom.cartao(6)['campos'][5]):'Latente',('FICHA PESSOAL','D44'):'Tentativa e Erro'}),
  ('invalido',novo,{**full,('CONFIGURAÇÃO',custom.cartao(3)['campos'][7]):'três'})]
 cached={}
 for name,source,edits in cases:
  if source not in cached:
   cached[source]=load_workbook(source)
  wb=cached[source]
  for ss in wb:
   ss._images=[]
   for cc in tuple(ss._cells.values()):
    if isinstance(cc.value,str) and cc.value.startswith('='):
     cc.value=re.sub(r'(?<![A-Z_.])(IFS|SWITCH|TEXTJOIN)\(',r'_xlfn.\1(',cc.value)
  changed={**common,**edits}
  restore=[(s,c,wb[s][c].value) for s,c in changed]
  for (s,c),v in changed.items():wb[s][c]=v
  path=inp/(name+'.xlsx');wb.save(path)
  for s,c,v in restore:wb[s][c]=v
  r=subprocess.run(['libreoffice',f'-env:UserInstallation={(tmp/"perfil").as_uri()}','--headless','--convert-to','xlsx','--outdir',str(out),str(path)],capture_output=True,text=True,timeout=600)
  check('recalcular '+name,r.returncode==0 and (out/path.name).exists(),r.stderr[-150:])
 for wb in cached.values():wb.close()
 cached.clear()
 results={}
 for name,*_ in cases:
  wb=load_workbook(out/(name+'.xlsx'),read_only=True,data_only=True)
  results[name]={s:{c.coordinate:c.value for row in wb[s].iter_rows() for c in row if c.value is not None} for s in wb.sheetnames if s not in ['CONFIGURAÇÃO','ANALISE']}
  wb.close()
 def get(name,s,c):return results[name].get(s,{}).get(c)
 def f(name,k):return get(name,'FICHA',idx[k])
 dif=[]
 for s,cs in results['base'].items():
  for c,v in cs.items():
   nv=get('vazia',s,c)
   equal=math.isclose(v,nv,rel_tol=1e-10,abs_tol=1e-9) if isinstance(v,(int,float)) and isinstance(nv,(int,float)) else v==nv
   if not equal:dif.append((s,c,v,nv))
 check('custom vazio: IGUAIS a todos os valores preexistentes',not dif,dif[:12])
 for n in range(1,11):check(f'cadastro {n} válido',get('completo',custom.NOME,f'S{n+1}') in ['',None],get('completo',custom.NOME,f'S{n+1}'))
 check('arma calcula volume × quantidade',get('completo','FICHA PESSOAL','AG92')==3)
 check('escudo conserva proteção e teto',get('completo','FICHA PESSOAL','W93')=='Proteção +4 · teto de Destreza 1')
 check('uniforme + escudo somam Proteção',f('completo','proteção')==7,f('completo','proteção'))
 check('Defesa usa menor teto e Buff',f('completo','defesa')==20,f('completo','defesa'))
 check('sem Força corta deslocamento uma vez',f('completo','deslocamento')=='4.5 m',f('completo','deslocamento'))
 check('item multiplica volume real',math.isclose(get('completo','FICHA PESSOAL','V104'),2.1))
 check('trocar por zero domina equipamento custom e Buff',f('troca_zero','defesa')==0)
 check('ferramenta herda arma e volume zero explícito',get('ferramenta','FICHA PESSOAL','U92')=='1d6' and get('ferramenta','FICHA PESSOAL','AG92')==0)
 check('ferramenta mantém grau, estigma e desgaste',get('ferramenta','FICHA PESSOAL','AI92')=='Grau 2' and get('ferramenta','FICHA PESSOAL','AK92')=='Estigma de teste' and get('ferramenta','FICHA PESSOAL','AQ92')==2)
 check('ferramenta de uniforme conserva Defesa',f('ferramenta','defesa')==20,f('ferramenta','defesa'))
 check('acessório ferramenta usa seu volume × quantidade',get('ferramenta','FICHA PESSOAL','V104')==.8)
 check('Origem do zero publica rota própria',get('completo',custom.NOME,'B122')=='Fundamento' and get('completo',custom.NOME,'B123')=='Nenhuma')
 check('Marcial sem energia troca os três indicadores',get('marcial_sem','DADOS','AK30')==1 and get('marcial_sem','DADOS','AK31')==0 and get('marcial_sem','DADOS','AK32')==1)
 check('Sem Energia conserva PE como Pontos de Esforço',f('marcial_sem','energia_max')==50,f('marcial_sem','energia_max'))
 check('marcante própria em Origem do livro é rejeitada',get('marcante_livro',custom.NOME,'T8') is False and 'Marcante própria' in get('marcante_livro',custom.NOME,'V8'))
 check('Restrição corporal permite marcante própria',get('marcante_restricao',custom.NOME,'T8') is True)
 check('base Latente rejeita marcante própria mas continua Origem válida',get('base_latente',custom.NOME,'T7') is True and get('base_latente',custom.NOME,'T8') is False)
 check('base Latente permite marcante publicada',not str(get('base_latente','FICHA PESSOAL','D51') or '').startswith('⚠'))
 check('capacidades custom trazem texto',get('completo','FICHA PESSOAL','D45')=='Texto marcante' and get('completo','FICHA PESSOAL','R45')=='Texto menor')
 check('colisão com livro tem diagnóstico',bool(get('duplicado',custom.NOME,'S2')) and get('duplicado',custom.NOME,'T2') is False)
 check('número inválido não vira cadastro válido',bool(get('invalido',custom.NOME,'S4')) and get('invalido',custom.NOME,'T4') is False)
 (ROOT/'nuvem-b47/saidas/resultados-mecanicos-locais.json').write_text(json.dumps({'falhas':falhas,'diferencas_vazio':dif[:50], 'resultados':results},ensure_ascii=False,default=str)+'\n')
print('CUSTOM CONFERE' if not falhas else f'{len(falhas)} FALHAS',flush=True)
sys.exit(bool(falhas))
