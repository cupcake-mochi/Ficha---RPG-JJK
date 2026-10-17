#!/usr/bin/env python3
"""B47/3b: contrato anterior, estrutura e cenários escalares independentes.

O artefato local é só uma calculadora de apoio. Nunca é importado no Google.
FILTER/UNIQUE, menus, aparência e montagem são verificados no Sheets nativo.
"""
import gzip, json, re, shutil, subprocess, sys, tempfile
from pathlib import Path
from openpyxl import Workbook, load_workbook

ROOT=Path(__file__).resolve().parent
sys.path[:0]=[str(ROOT/'ficha-v01'),str(ROOT/'ficha')]
import custom_progressao as p, indice_ficha as ix, emitir_gs
EVID=ROOT/'nuvem-b47/saidas'
EVID.mkdir(parents=True,exist_ok=True)
# A fresh package prevents a stale evidence directory from validating old code.
# The 03a459e snapshot is a versioned fixture, so a clean checkout can run this.
subprocess.run([sys.executable,str(ROOT/'ficha-v01/monta.py'),'--exportar-progressao',
                str(EVID/'progressao-em-integracao.json')],check=True,
               stdout=subprocess.DEVNULL,timeout=300)
pack=json.loads((EVID/'progressao-em-integracao.json').read_text())
layout=pack['layout'];abas={a['nome']:a for a in layout['abas']}
old=json.loads(gzip.decompress((ROOT/'ficha-v01/progressao-base3a.json.gz').read_bytes()))
fails=[]
def check(name,ok,details=''):
    print(('OK ' if ok else 'FALHA ')+name+(' '+str(details) if not ok else ''),flush=True)
    if not ok:fails.append(name)
raw={n:{v[0]:v[1] for v in a['celulas']} for n,a in abas.items()}
prior=next(a for a in old if a['nome']=='ANALISE')
prior_rows={v[0]:{} for v in prior['vals']}
for r,c,v,*_ in prior['vals']:prior_rows[r][c]=v
keys={v[0]:v[1] for v in [[raw['ANALISE'].get(f'A{r}'),raw['ANALISE'].get(f'B{r}')] for r in range(2,abas['ANALISE']['linhas']+1)]}
check('6174 chaves anteriores presentes',sum(bool(v.get(1)) for r,v in prior_rows.items() if r>1)==6174 and all(v[1] in keys for r,v in prior_rows.items() if r>1 and v.get(1)))
check('sem chaves repetidas',len(keys)==abas['ANALISE']['linhas']-1)
check('formato 11 explícito (C)',str(keys.get('formato.versao'))=='11')
check('um Caminho/uma Trilha e 10 Perícias + 10 Ofícios',len([k for k in keys if re.fullmatch(r'custom.pericia.\d+.tipo',str(k))])==10 and len([k for k in keys if re.fullmatch(r'custom.oficio.\d+.tipo',str(k))])==10 and 'custom.caminho.nivel.30.texto' in keys and 'custom.trilha.campo.base_do_livro' in keys)
check('nenhum arquivo atinge 900000 bytes',all(x.stat().st_size<900000 for x in (ROOT/'apps-script').glob('*.gs')))
check('fragmentos recompõem todas as abas do modelo',set(a['nome'] for a in emitir_gs.abas_do_script(str(ROOT/'apps-script/Ficha.gs')))==set(abas)-{'INVOCAÇÃO','CATÁLOGO','DADOS_INV'})
for name in ['CONFIGURAÇÃO','FICHA']:
    ms=[(*ix._lc(x.split(':')[0]),*ix._lc(x.split(':')[-1])) for x in abas[name]['mescladas']]
    overlap=[]
    for i,(r,c,rr,cc) in enumerate(ms):
        for s,d,ss,dd in ms[i+1:]:
            if max(r,s)<=min(rr,ss) and max(c,d)<=min(cc,dd):overlap.append((r,c,s,d))
    check(name+': mesclagens sem sobreposição',not overlap,overlap[:5])
    prior_sheet=next(a for a in old if a['nome']==name)
    expected={tuple(m) for m in prior_sheet['merges']}
    current={tuple(m) for m in ms}
    check(name+': todas as mesclagens anteriores preservadas',expected<=current)
    # A tabela de limites de entidade ocupa uma folga aprovada B133:G149.
    # Não aceitar qualquer outra nova mesclagem dentro da parte anterior.
    allowed={(153,2,154,7),(156,2,156,4),(162,2,164,7)}|{(r,2,r,4) for r in range(157,161)}|{(133,2,134,7),(135,2,136,7),(138,2,138,4),(147,2,149,7)}|{(r,2,r,4) for r in range(139,146)} if name=='CONFIGURAÇÃO' else set()
    check(name+': nenhuma mesclagem nova invade a área anterior',all(m in expected or m in allowed or m[0]>prior_sheet['rows'] for m in current))
# Correção pedida pelo usuário: mesmas cartas/níveis do livro, sem espelho paralelo.
check('Caminho tem cinco níveis e Trilha quatro, todos do livro',
      [v['nivel'] for v in p.longo('Caminho')['niveis'] if v['nome']]==[2,7,15,23,30] and
      [v['nivel'] for v in p.longo('Trilha')['niveis'] if v['nome']]==[2,11,19,27])
check('FICHA não contém bloco duplicado de habilidades custom',
      not any(isinstance(v,str) and 'HABILIDADES CUSTOM' in v for v in raw['FICHA'].values()))
for typ,total in [('Perícia',23),('Ofício',11)]:
    rows=[v for v in pack['linhas_ficha'] if v['tipo']==typ]
    check(typ+': capacidade preservada em duas colunas',len(rows)==total and len({ix._lc(v['nome'])[1] for v in rows})==2)
    check(typ+': duplicata confere as duas colunas',all(v['valido'].count('COUNTIF(\'FICHA\'!')==2 for v in rows))
styles={v[0]:v[2] for v in abas['FICHA']['celulas']}
check('divisória lateral continua até o fim',all(layout['estilos'][styles['B'+str(r)]][2].get('right')==['medium','FF8A7EC4']
      for r in range(408,abas['FICHA']['linhas']+1)))
check('CONFIGURAÇÃO fixa e sem congelamento',abas['CONFIGURAÇÃO']['paleta_fixa'] and not abas['CONFIGURAÇÃO'].get('congeladas'))
changed=[]
generated={a['nome']:{ix._letras(c)+str(r):v for r,c,v,*_ in a['vals']}
           for a in emitir_gs.abas_do_script(str(ROOT/'apps-script/Ficha.gs'))}
def contrato_esperado(r,c,v):
    # Ampliação solicitada: agregados leem quatro extras; DEFINIÇÃO desce uma
    # linha para deixar o respiro. Todas as demais fontes anteriores ficam.
    match=re.fullmatch(r'invocacao\.(\d+)\.familias',str(prior_rows[r].get(1,'')))
    if match and c in (2,3):
        row=int(match[1])+1
        if c==2:return v[:-1]+''.join(f',DADOS_INVOC!${col}${row}' for col in ('QV','QW','QX','QY'))+')'
        return v+''.join(f' DADOS_INVOC!{col}{row}' for col in ('QV','QW','QX','QY'))
    match=re.fullmatch(r'invocacao\.(\d+)\.definicao',str(prior_rows[r].get(1,'')))
    if match and c in (2,3):
        n=int(match[1])-1;row=52+198*(n%6);col=ix._letras(9+19*(n//6))
        return v.replace(f'{col}${row}',f'{col}${row+1}').replace(f'{col}{row}',f'{col}{row+1}')
    return v
contract_diff=[(r,c,contrato_esperado(r,c,v),generated['ANALISE'].get(ix._letras(c)+str(r)))
               for r,cs in prior_rows.items() if r>1 and cs.get(1)
               for c,v in cs.items() if c<=4 and not (r==2 and c==2)
               and contrato_esperado(r,c,v)!=generated['ANALISE'].get(ix._letras(c)+str(r))]
check('contrato anterior preservado, com quatro extras nos agregados e Definição uma linha abaixo',not contract_diff,contract_diff[:3])
def fatoracao_basica_exata(original, atual, row, valores, antes, depois):
    """Recognize only the approved U2:U26 -> JY:KA factorization.

    Rebuild the complete frozen formula from PRE-progression helpers, requiring
    identical guards, diagnostics and order. Modified helpers must also match
    the recorded post-progression formula. No arbitrary old-cell changes pass.
    """
    if not 2<=row<=26 or atual!=f'=TEXTJOIN(" · ",TRUE,JY{row}:KA{row})':return False
    pat=r'=IF\((.+),"",TEXTJOIN\(" · ",TRUE,(.*)\)\)'
    old_formula=re.fullmatch(pat,original) if isinstance(original,str) else None
    if not old_formula:return False
    parts=[]
    for col in ('JY','JZ','KA'):
        coord=col+str(row);value=valores.get(coord)
        if coord in antes and (coord not in depois or depois[coord]!=value):return False
        pre=antes.get(coord,value)
        m=re.fullmatch(pat,pre) if isinstance(pre,str) else None
        if not m or m[1]!=old_formula[1]:return False
        parts.append(m[2])
    rebuilt='=IF('+old_formula[1]+',"",TEXTJOIN(" · ",TRUE,'+','.join(parts)+'))'
    return rebuilt==original

basic_old={r:v for a in old if a['nome']=='DADOS_CUSTOM' for r,c,v,*_ in a['vals'] if c==21 and 2<=r<=26}
basic_now=generated['DADOS_CUSTOM'];basic_before=pack['antes'].get('DADOS_CUSTOM',{});basic_after=pack['depois'].get('DADOS_CUSTOM',{})
def factor(row=2,original=None,current=None,values=None,before=None,after=None):
    return fatoracao_basica_exata(basic_old[row] if original is None else original,
           basic_now['U'+str(row)] if current is None else current,row,
           basic_now if values is None else values,basic_before if before is None else before,basic_after if after is None else after)
check('25 validações 3a recompõem a fórmula histórica inteira, sem mudar ordem ou mensagens',len(basic_old)==25 and all(factor(row=r) for r in basic_old))
check('contraprova: mudar condição, guarda ou mensagem histórica é recusado',all(not factor(original=basic_old[2].replace(a,b,1)) for a,b in [('COUNTIF','COUNTIFS'),('LEN(TEXTJOIN','N(TEXTJOIN'),('Preencha NOME','Nome opcional')]))
check('contraprova: auxiliar omitida ou linha trocada é recusada',not factor(current='=TEXTJOIN(" · ",TRUE,JY2:JZ2)') and not factor(current='=TEXTJOIN(" · ",TRUE,JY3:KA3)'))
broken=dict(basic_after);broken['JY2']=basic_now['JY2']+' '
check('contraprova: auxiliar divergente do depois registrado é recusada',not factor(after=broken))

for a in old:
    if a['nome']=='ANALISE':continue
    for r,c,v,*_ in a['vals']:
        coord=ix._letras(c)+str(r);nv=generated[a['nome']].get(coord)
        norm=lambda value:None if value is None or value=='' else value
        if norm(v)==norm(nv):continue
        if a['nome']=='DADOS_CUSTOM' and c==21 and fatoracao_basica_exata(v,nv,r,basic_now,basic_before,basic_after):continue
        expected=pack['antes'].get(a['nome'],{}).get(coord,'__não declarado__')
        after=pack['depois'].get(a['nome'],{}).get(coord,'__não declarado__')
        # O emissor GS serializa None como string vazia. Normalizar somente essa
        # representação; fórmulas e textos continuam exigindo igualdade inteira.
        if norm(expected)!=norm(v) or norm(after)!=norm(nv):changed.append((a['nome'],coord,v,nv))
check('toda alteração preexistente tem antes/depois exatos',not changed,[(sheet,cell,len(str(before)),len(str(after))) for sheet,cell,before,after in changed[:5]])
if '--estrutura' in sys.argv:sys.exit(bool(fails))

def short(n,typ,name,attr):
    g=p.curto(n)
    return {('CONFIGURAÇÃO',g[k]):v for k,v in [('tipo',typ),('nome',name),('atributo',attr)]}
path=p.longo('Caminho');trail=p.longo('Trilha')
def path_fields(values):return {('CONFIGURAÇÃO',path['campos'][k]):v for k,v in values.items()}
fixture={**short(1,'Perícia','Rituais Urbanos','Essência'),**short(11,'Ofício','Cartógrafo do Véu','Inteligência'),
         ('CONFIGURAÇÃO',path['nome']):'Caminho Teste',('CONFIGURAÇÃO',trail['nome']):'Trilha Teste',
         ('CONFIGURAÇÃO',trail['campos']['CAMINHO']):'Caminho Teste',('CONFIGURAÇÃO',trail['campos']['BASE DO LIVRO']):'Múltiplas Invocações',
         **path_fields({'VIDA INICIAL':10,'VIDA POR NÍVEL':6,'DADO DE VIDA':'d8','PE POR NÍVEL':7,'ATRIBUTO NATURAL 1':'Força','ATRIBUTO NATURAL 2':'Constituição',
                        'PERÍCIA FIXA 1':'Rituais Urbanos','PERÍCIA FIXA 2':'Acrobacia','TROCA PERÍCIAS POR ARMA':'Sim','TREINO DE ARMA':'Categorias selecionadas','CATEGORIA DE ARMA 1':'Yumi'}),
         ('CONFIGURAÇÃO',path['niveis'][1]['nome']):'Passo do Véu',('CONFIGURAÇÃO',path['niveis'][1]['texto']):'Descrição integral do nível 2.',
         ('CONFIGURAÇÃO',path['niveis'][6]['nome']):'Entrega futura',('CONFIGURAÇÃO',path['niveis'][6]['texto']):'Descrição do nível 7.',
         ('FICHA','D11'):'Caminho Teste',('FICHA','N11'):'Trilha Teste',('FICHA','AK81'):'1 arma (-2 pericias)'}
per=next(g for g in pack['linhas_ficha'] if g['tipo']=='Perícia');off=next(g for g in pack['linhas_ficha'] if g['tipo']=='Ofício')
fixture.update({('FICHA',per['nome']):'Rituais Urbanos',('FICHA',per['treino']):True,('FICHA',off['nome']):'Cartógrafo do Véu',('FICHA',off['especializacao']):True})
common={('FICHA','AH11'):5,('FICHA','D21'):3,('FICHA','L21'):2,('FICHA','T21'):3,('FICHA','AB21'):2,('FICHA','AJ21'):4}
cases=[('livro_base',None,{}),('vazio',raw,{}),('custom',raw,fixture),
       ('troca_zero',raw,{**fixture,('CONFIGURAÇÃO','B11'):True,('CONFIGURAÇÃO','C11'):'Perícia · Rituais Urbanos',('CONFIGURAÇÃO','D11'):'trocar por',('CONFIGURAÇÃO','E11'):0,('CONFIGURAÇÃO','F11'):'Decisão do mestre',('CONFIGURAÇÃO','G11'):'Teste'}),
       ('duplicado',raw,{**fixture,**short(3,'Perícia','Rituais Urbanos','Força')}),
       ('numero_invalido',raw,{**fixture,**path_fields({'VIDA POR NÍVEL':'seis'})}),
       ('atributo_invalido',raw,{**fixture,**short(1,'Perícia','Rituais Urbanos','Carisma')}),
       ('rolada',raw,{**fixture,('CONFIGURAÇÃO','E59'):'Sim',**{('CONFIGURAÇÃO','E'+str(r)):v for r,v in zip(range(101,105),[8,4,6,2])}}),
       ('dado_invalido',raw,{**fixture,('CONFIGURAÇÃO','E59'):'Sim',**{('CONFIGURAÇÃO','E'+str(r)):v for r,v in zip(range(101,105),[9,4,6,2])}})]
cases.append(('trilha_do_zero',raw,{**fixture,('CONFIGURAÇÃO',trail['campos']['BASE DO LIVRO']):'Nenhuma',
   ('CONFIGURAÇÃO','B11'):True,('CONFIGURAÇÃO','C11'):'Invocações · teto de ativas',('CONFIGURAÇÃO','D11'):'trocar por',('CONFIGURAÇÃO','E11'):5,('CONFIGURAÇÃO','F11'):'Decisão do mestre',('CONFIGURAÇÃO','G11'):'Trilha própria',
   ('CONFIGURAÇÃO','B12'):True,('CONFIGURAÇÃO','C12'):'Invocações · famílias da trilha',('CONFIGURAÇÃO','D12'):'somar',('CONFIGURAÇÃO','E12'):2,('CONFIGURAÇÃO','F12'):'Decisão do mestre',('CONFIGURAÇÃO','G12'):'Trilha própria'}))
oldraw={a['nome']:{ix._letras(c)+str(r):v for r,c,v,*_ in a['vals']} for a in old}
keep=['FICHA','DADOS','CONFIGURAÇÃO','DADOS_CONFIG','DADOS_CUSTOM','FICHA PESSOAL',p.NOME,'DADOS_INVOC']
results={}
with tempfile.TemporaryDirectory(prefix='b47-3b-') as tmp:
    tmp=Path(tmp);out=tmp/'out';out.mkdir()
    for name,source,edits in cases:
        data={s:dict(v) for s,v in (source or oldraw).items() if s in keep}
        for (s,c),v in {**common,**edits}.items():data[s][c]=v
        wb=Workbook();wb.remove(wb.active)
        for s,vals in data.items():
            ws=wb.create_sheet(s)
            for c,v in vals.items():
                if v is None or v=='':continue
                # Calculations only: omit wide text cards and unused entity UI.
                if s=='FICHA' and 105<=ix._lc(c)[0]<413:continue
                if isinstance(v,str) and v.startswith('='):v=re.sub(r'(?<![A-Z_.])(IFS|SWITCH|TEXTJOIN)\(',r'_xlfn.\1(',v)
                ws[c]=v
        file=tmp/(name+'.xlsx');wb.save(file);wb.close()
        run=subprocess.run([shutil.which('libreoffice') or 'soffice',f'-env:UserInstallation={(tmp/"profile").as_uri()}','--headless','--convert-to','xlsx','--outdir',str(out),str(file)],capture_output=True,text=True,timeout=300)
        check('recalcular '+name,run.returncode==0 and (out/file.name).exists(),run.stderr[-200:])
        if not (out/file.name).exists():continue
        wb=load_workbook(out/file.name,data_only=True,read_only=True)
        wanted={'FICHA':['J26','J30','D67',per['resultado'],off['resultado']],'DADOS':['AK13','AK14','AK15','AK18'],p.NOME:['G2','G3','G4','K2','K3','K4','K7','K8','K9','M8','T8','BG3','BH3','BG7','BG8','BH8'], 'DADOS_INVOC':['DB8','DB11']}
        results[name]={s:{c:wb[s][c].value for c in cs} for s,cs in wanted.items() if s in wb.sheetnames}
        wb.close()
def get(n,s,c):return results.get(n,{}).get(s,{}).get(c)
check('vazio conserva vida, PE, maestria e orçamentos do livro',all(get('vazio',s,c)==v for s,vs in results['livro_base'].items() for c,v in vs.items() if s in ['FICHA','DADOS'] and c not in [per['resultado'],off['resultado']]))
check('Caminho custom válido: vida 49 e PE 35',get('custom','FICHA','J26')==49 and get('custom','FICHA','J30')==35,results.get('custom'))
check('dado d8 vira limite numérico 8',get('custom',p.NOME,'K9')==8)
check('vida rolada preserva inicial fixa e soma quatro dados + Constituição',get('rolada','FICHA','J26')==45,get('rolada','FICHA','J26'))
check('rolagem acima do dado é rejeitada',str(get('dado_invalido','FICHA','J26')).startswith('#'),get('dado_invalido','FICHA','J26'))
mae=get('custom','FICHA','D67')
check('perícia e ofício usam atributos e especialização existentes',get('custom','FICHA',per['resultado'])==4+mae and get('custom','FICHA',off['resultado'])==2+mae+mae//2)
check('custom acrescenta um treino de perícia e uma especialização',get('custom','DADOS','AK13')==get('vazio','DADOS','AK13')+1 and get('custom','DADOS','AK15')==get('vazio','DADOS','AK15')+1)
check('troca de perícias por arma custa 2',get('custom','DADOS','AK18')==2)
check('Trilha herda Múltiplas Invocações: 4 ativas e 4 famílias',get('custom','DADOS_INVOC','DB8')==4 and get('custom','DADOS_INVOC','DB11')==4)
check('Trilha do zero ajusta limites por BÔNUS E MUDANÇAS',get('trilha_do_zero','DADOS_INVOC','DB8')==5 and get('trilha_do_zero','DADOS_INVOC','DB11')==5)
check('nível 2 inteiro, nível 6 inexistente e nível 7 disponível para a carta riscada',get('custom',p.NOME,'BG3')=='Passo do Véu' and get('custom',p.NOME,'BH3')=='Descrição integral do nível 2.' and not get('custom',p.NOME,'BG7') and get('custom',p.NOME,'BG8')=='Entrega futura' and get('custom',p.NOME,'BH8')=='Descrição do nível 7.')
check('trocar por 0 fixa resultado da perícia custom',get('troca_zero','FICHA',per['resultado'])==0)
check('nomes duplicados invalidam os dois registros',all('repetido' in str(get('duplicado',p.NOME,c)) for c in ['G2','G4']))
check('número inválido impede entrada do Caminho no catálogo',bool(get('numero_invalido',p.NOME,'K2')) and not get('numero_invalido',p.NOME,'M8'))
check('atributo inexistente é rejeitado',bool(get('atributo_invalido',p.NOME,'G2')))
(EVID/'resultados-locais-3b.json').write_text(json.dumps(dict(falhas=fails,resultados=results),ensure_ascii=False,indent=2))
print('PROGRESSÃO CONFERE' if not fails else f'{len(fails)} FALHAS',flush=True)
sys.exit(bool(fails))
