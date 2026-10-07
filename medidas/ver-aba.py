# -*- coding: utf-8 -*-
"""Desenha uma aba da ficha em HTML, a partir do que o script manda para o Sheets.

O Apps Script não roda fora do Google, e o .xlsx aberto no LibreOffice não mostra o que o Sheets mostra (fonte,
caixa de seleção, barra, grupo fechado). Este desenhador lê o `ABAS` do apps-script/Ficha.gs, que é a lista de
instruções que o construir() executa, e pinta cada célula como o montarAba_ pinta: largura e altura em pixel, fundo,
fonte, mesclagem, borda por lado, menu, caixa de seleção, nota. O valor de cada fórmula vem de uma cópia da ficha
gerada, preenchida com um exemplo e recalculada pelo LibreOffice.

Serve para olhar a aba antes de testar no Sheets e para comparar com o estudo que o Mizuki aprovou. Não é o Sheets:
a barra (SPARKLINE) é desenhada aqui pela conta dela, e a fonte é a do Google Fonts que o navegador baixar.

    python3 medidas/ver-aba.py                      # a FICHA PESSOAL com a Kaori do estudo -> /tmp/ver-aba.html
    python3 medidas/ver-aba.py --saida x.html --fechado   # com o painel de XP fechado, como a aba nasce
"""
import argparse, html, json, os, re, shutil, subprocess, sys, tempfile
from openpyxl import load_workbook

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(RAIZ, "ficha-v01"))
sys.path.insert(0, os.path.join(RAIZ, "ficha"))
import indice_ficha as ix, ficha_pessoal as fp, ficha_amaldicoada as fa, ficha_invocacoes as fi, emitir_gs

ARQ = os.path.join(RAIZ, "ficha-v01", "ficha-projeto-m-0.1.xlsx")
GS = os.path.join(RAIZ, "apps-script", "Ficha.gs")


def abas_do_script():
    """o ABAS como o script o usa: com as fileiras copiadas da FICHA AMALDIÇOADA escritas por extenso"""
    return emitir_gs.abas_do_script(GS)


def arte_do_script():
    """os PNG de uma cor só que o construir() põe na ficha, em base64, pelo nome"""
    src = open(GS, encoding="utf-8").read()
    bloco = re.search(r"var ARTE = \{([\s\S]*?)\n\};", src).group(1)
    arte = {}
    for nome, resto in re.findall(r'"([^"]+\.png)":((?:\s*\+?\s*"[^"]*")+)', bloco):
        arte[nome] = "".join(re.findall(r'"([^"]*)"', resto))
    return arte


def exemplo(wb, R, G):
    """a Kaori do estudo: nível 2, Bastião, Força 3, com o kit que o estudo mostra"""
    f, p, c = wb["FICHA"], wb[fp.NOME], wb["CARTEIRA"]
    idx = {}
    dd = wb["DADOS"]
    for r in range(5, 200):
        k, v = dd.cell(row=r, column=53).value, ix.endereco(dd.cell(row=r, column=54).value)
        if k and v:
            idx[k] = v
    for linha in c.iter_rows():                      # o nome do portador, digitado na CARTEIRA
        for cel in linha:
            if isinstance(cel.value, str) and cel.value.startswith("Coloque o nome"):
                cel.value = "Kaori"
    f[idx["atr_base_Força"]], f[idx["atr_base_Destreza"]], f[idx["caminho"]], f[idx["nivel"]] = 3, 2, "Bastião", 2
    p[G["principal"]], p[G["secundaria"]], p[G["situacao"]] = "Faca", "Broquel", "Escuro"
    p[G["ienes"]] = 61000
    for i, (nome, qtd) in enumerate([("Kanabō", 1), ("Faca", 1), ("Kunai", 2), ("Broquel", 1)]):
        p[f"D{G['equip_ini'] + i}"], p[f"K{G['equip_ini'] + i}"] = nome, qtd
    for i, nome in enumerate(["Corda", "Lanterna", "Kit de escalada"]):
        p[f"D{G['itens_ini'] + i}"], p[f"T{G['itens_ini'] + i}"] = nome, 1
    for col, lin, _ in list(G["armas"].values()) + list(G["grupos"].values()):   # Bastião treina todas
        p.cell(row=lin, column=col).value = True
    p[G["historia"]] = ("O clã da Kaori perdeu o nome faz três gerações, e ela cresceu ouvindo a história de quem perdeu. A avó era "
                "a única que ainda sabia alguma coisa de valor, ervas principalmente, e fez questão de ensinar.")
    for i, (nome, tipo, mult, desc) in enumerate([("Missão curta (exemplo)", "Curta", "1,25x", None),
                                                  ("3ª da semana (exemplo)", "Padrão", None, "3ª · 50%"),
                                                  ("5ª da semana (exemplo)", "Padrão", None, "5ª · 12,5%")]):
        lin = G["missao_ini"] + i
        for desloc, v in ((0, nome), (fp.C_XP, tipo), (fp.C_ADIC, mult), (fp.C_DESC, desc)):
            p.cell(row=lin, column=fp.B1 + desloc).value = v


ALTURAS_DO_SCRIPT = {}       # as linhas que o Codigo.gs estica (a caixa das cartas de Habilidades): {linha: px}


def _escrever_habilidades(f, caminho, trilha):
    """escreve nas cartas da seção 7 da FICHA o que o Codigo.gs escreve quando o Caminho e a Trilha são escolhidos: o nome
    e o texto do livro, e a caixa esticada para o texto. As cartas são achadas pelo endereço que a DADOS_AM publica"""
    import habilidades as _hb
    d = f.parent["DADOS_AM"]
    cab = {d.cell(row=1, column=c).value: c for c in range(1, d.max_column + 1)}
    livro = {(l[1], l[0], l[2]): l for l in _hb.linhas_do_livro()}
    M = _hb.medida()
    for r in range(2, 40):
        carta = d.cell(row=r, column=cab["habilidade: carta"]).value
        if not carta:
            break
        fonte, nivel = carta.split(" ")[0], int(d.cell(row=r, column=cab["nível da carta"]).value)
        l = livro.get((fonte, caminho if fonte == "Caminho" else trilha, nivel))
        if not l:
            continue
        cn, ct = (ix.endereco(d.cell(row=r, column=cab[k]).value) for k in ("célula do nome", "célula do texto"))
        f[cn], f[ct] = l[3], l[4]
        px = max(M["minima"], -(-(l[5] * M["linha"] + M["respiro"]) // M["caixa"]))
        for k in range(M["caixa"]):
            f.row_dimensions[f[ct].row + k].height = px * 0.75
            ALTURAS_DO_SCRIPT[f[ct].row + k] = px


def exemplo_amaldicoada(wb):
    """a Kaori do estudo da Ficha Amaldiçoada: nível 10, com a técnica dela e seis feitiços prontos do livro"""
    f, a, c = wb["FICHA"], wb[fa.NOME], wb["CARTEIRA"]
    G = fa.geometria()
    idx = {}
    dd = wb["DADOS"]
    for r in range(5, 200):
        k, v = dd.cell(row=r, column=53).value, ix.endereco(dd.cell(row=r, column=54).value)
        if k and v:
            idx[k] = v
    for linha in c.iter_rows():
        for cel in linha:
            if isinstance(cel.value, str) and cel.value.startswith("Coloque o nome"):
                cel.value = "Kaori"
    f[idx["atr_base_Força"]], f[idx["atr_base_Essência"]], f[idx["caminho"]], f[idx["nivel"]] = 3, 2, "Bastião", 10
    f[idx["refino escolhido"]] = 2
    a[G["nome_tecnica"]], a[G["tipo_dano"]] = "Peso Emprestado", "Impacto"
    a[f"D{G['regra'] + 1}"] = "Tudo que eu prendo entre as minhas mãos fica mais pesado."
    a[f"L{G['descricao'] + 1}"] = "As duas mãos precisam se tocar antes."
    a[f"L{G['descricao'] + 4}"] = "Ela sabe o peso exato de qualquer coisa que encoste nela."
    for fam, estado in {"Controle": "Livre", "Castigo": "Livre", "Amparo": "Fechada", "Área": "Fechada", "Auxiliares": "Fechada"}.items():
        i = list(fa.regras()["familias"]).index(fam)
        a[f"{G['cols_fam'][i][0]}{G['familias'] + 1}"] = estado
    feiticos = [("Estalo", 1, "Projétil", [], [], "Ela bate as mãos e o ar entre elas ganha peso. O que sai é um soco sem braço."),
                ("Perfurar", 1, "Projétil", ["Precisão"], ["Parado"], "Parada, ela aperta o ar até virar uma ponta e solta num alvo só."),
                ("Golpe Cru", 1, "Toque", [], [], "A mão encosta e o peso entra direto no corpo do outro."),
                ("Lança Negra", 2, "Projétil", ["Fura"], ["Atrasar"], "Uma rodada inteira apertando o ar entre as palmas. Sai uma haste escura que atravessa proteção."),
                ("Marca do Carrasco", 3, "Projétil", ["Marca", "Queima"], ["Uma Vez"], "O peso fica grudado no alvo depois do golpe e continua esmagando. Uma vez por cena."),
                ("Palma Trovejante", 2, "Cone", ["Derrubado"], ["Atrasar"], "Ela abre as mãos de uma vez e o peso sai em leque, derrubando o que estiver na frente.")]
    for (nome, classe, forma, mel, res, como), pos in zip(feiticos, G["feiticos"]):
        cel = fa.celulas_do_feitico(*pos)
        a[cel["nome"]], a[cel["classe"]], a[cel["forma"]], a[cel["como"]] = nome, classe, forma, como
        for k, m in zip(cel["mel"], mel):
            a[k] = m
        for k, r in zip(cel["res"], res):
            a[k] = r
    cel = fa.celulas_do_feitico(*G["libs"][0])
    a[cel["nome"]], a[cel["como"]] = "Golpe do Voto", "Tudo que ela segurou na luta inteira, devolvido num golpe só."
    for p, pos in zip(["Raiz", "Fluxo"], G["passivas"]):
        a[fa.celulas_da_passiva(*pos)["nome"]] = p
    a[fa.celulas_da_passiva(*G["passivas"][1])["texto"]] = "O peso que ela solta volta para as mãos dela, devagar."
    a[G["tm_nome"]], a[f"D{G['tm_como'] + 1}"] = "Sentença de Chumbo", "Tudo o que ela tocou na luta pesa ao mesmo tempo."
    a[G["dom_nome"]], a[G["degrau"]] = "Balança Quebrada", "Incompleta"
    # a seção 7 da FICHA (as Habilidades): o livro do Bastião e do Muro, como o script escreve; as cartas acima do nível
    # 10 da Kaori ainda não abriram e saem riscadas
    _escrever_habilidades(f, "Bastião", "Muro")
    for p, pos in zip(["Projetar energia", "Barreira Simples"], G["aptidoes"]):
        a[fa.celulas_da_aptidao(*pos)["nome"]] = p
    pc = fa.celulas_do_pacto(G["pactos"][0])
    a[pc["nome"]], a[pc["forma"]] = "Mostrar a mão", "temporário"
    a[pc["dou"]], a[pc["recebo"]] = "Explicar a própria técnica ao adversário.", "A técnica fica mais forte."
    cz = G["cols_zero"]
    a[f"{cz['nome'][0]}{G['zero_ini']}"], a[f"{cz['forma'][0]}{G['zero_ini']}"] = "Tapa de Peso", "Toque"
    a[f"{cz['nome'][0]}{G['zero_ini'] + 1}"], a[f"{cz['mel'][0]}{G['zero_ini'] + 1}"] = "Pedrada", "Empurrão"


def exemplo_invocacoes(wb):
    """o Cão de sombra do capítulo 17, no nível 5 do Kaito (Essência 2): a ficha que o livro monta passo a passo"""
    f, a, c = wb["FICHA"], wb[fi.NOME], wb["CARTEIRA"]
    idx = {}
    dd = wb["DADOS"]
    for r in range(5, 200):
        k, v = dd.cell(row=r, column=53).value, ix.endereco(dd.cell(row=r, column=54).value)
        if k and v:
            idx[k] = v
    for linha in c.iter_rows():
        for cel in linha:
            if isinstance(cel.value, str) and cel.value.startswith("Coloque o nome"):
                cel.value = "Kaito"
    f[idx["atr_base_Essência"]], f[idx["caminho"]], f[idx["nivel"]] = 2, "Evocador", 5
    g = fi.celulas_da_ficha(fi.L0, 0)
    a[g["nome"]], a[g["acerto"]], a[g["fis"]], a[g["trT"]] = "Cão de sombra", "Força", "Força", "Físico"
    a[g["def"]] = "Um cão feito de sombra que reconhece vestígios de energia amaldiçoada e persegue o que seu invocador aponta."
    a[g["corpo"]] = ("Médio, quatro patas, sem mãos; usa a boca para segurar. Visão, audição e olfato comuns. Entende ordens "
                     "faladas e responde por latidos e gestos.")
    a[g["vida"]] = 19
    for cel, v in zip(g["pts"], (3, 2, 2, 1, 1)):
        a[cel] = v
    for cel, v in zip(g["fam"], ("Mira", None, "Alcance", "Controle", None)):
        a[cel] = v
    for cel, v in zip(g["per"], ("Atletismo", "Furtividade", "Percepção", "Sobrevivência")):
        a[cel] = v
    a[g["tal"][0]] = "Farejador"
    for chave, nome, forma, mel, como in ((("bas", 0), "Mordida", "Toque", None, "Consome a atuação básica do cão."),
                                          (("esp", 0), "Mordida precisa", "Toque", "Precisão", "Não aplica condição nem deixa efeito contínuo.")):
        cel = fi.celulas_da_carta(*g["cartas"][chave])
        a[cel["nome"]], a[cel["forma"]], a[cel["tdano"]], a[cel["como"]], a[cel["mel"][0]] = nome, forma, "Perfurante", como, mel
    # 07/10/2026, a grade de 2 × 6: o Vigia de papel do capítulo 17 no primeiro lugar da segunda coluna de fichas (a 7)
    if fi.N_COLUNAS > 1:
        g = fi.celulas_da_ficha(fi.L0, 1)
        a[g["nome"]], a[g["acerto"]], a[g["fis"]], a[g["trT"]] = "Vigia de papel", "Inteligência", "Destreza", "Intelecto"
        a[g["def"]] = "Uma figura de tiras de papel que vigia um lugar, avisa o invocador e ampara quem cai."
        for cel, v in zip(g["pts"], (0, 2, 2, 3, 2)):
            a[cel] = v
        for cel, v in zip(g["fam"], ("Amparo", None, "Auxiliares", "Alcance", None)):
            a[cel] = v
        for cel, v in zip(g["per"], ("Acrobacia", "Furtividade", "Investigação", "Percepção", "Sobrevivência")):
            a[cel] = v
        for chave, nome, classe, forma, mels in ((("bas", 0), "Orientação", None, "Apoio", ["Impulso"]),
                                                 (("esp", 0), "Tiras de resgate", 2, "Apoio", ["Guarda", "Empurrão"]),
                                                 (("esp", 1), "Remendo de papel", 1, "Cura", [])):
            cel = fi.celulas_da_carta(*g["cartas"][chave])
            a[cel["nome"]], a[cel["forma"]] = nome, forma
            if classe:
                a[cel["classe"]] = classe
            for c_, m_ in zip(cel["mel"], mels):
                a[c_] = m_


def recalculada(preenche):
    d = tempfile.mkdtemp(prefix="ver-aba-")
    copia = os.path.join(d, "ficha.xlsx")
    shutil.copy(ARQ, copia)
    wb = load_workbook(copia)
    preenche(wb)
    for ws in wb:
        for linha in ws.iter_rows():
            for c in linha:
                if isinstance(c.value, str) and c.value.startswith("="):
                    c.value = re.sub(r"(?<![A-Z_.])(IFS|TEXTJOIN)\(", r"_xlfn.\1(", c.value)
    wb.save(copia)
    subprocess.run(["libreoffice", "--headless", "--convert-to", "xlsx", "--outdir", os.path.join(d, "s"), copia],
                   capture_output=True, timeout=300)
    lido = load_workbook(os.path.join(d, "s", "ficha.xlsx"), data_only=True)
    cru = load_workbook(copia)
    shutil.rmtree(d, ignore_errors=True)
    return lido, cru


def numero(v):
    if isinstance(v, bool) or v is None:
        return v
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).replace(".", ",") if isinstance(v, float) else v


def desenha(spec, valores, crus, aberto=True, barras=None, titulo="", pintura=None, linhas=None, arte=None, linhas_abertas=None):
    """`pintura` é a aba depois de uma troca de paleta (o que o medidas/pintar-paletas.js grava): fundo e fonte de cada
    célula, a régua e a cor de cada imagem. `linhas` corta a aba nas primeiras N. `arte` são os PNG do Ficha.gs.
    `linhas_abertas` diz o que fazer com os grupos de LINHAS: None abre todos, "nasce" deixa como a aba nasce."""
    return pagina(grade(spec, valores, crus, aberto, barras, pintura, linhas, arte, linhas_abertas), titulo)


def grade(spec, valores, crus, aberto=True, barras=None, pintura=None, linhas=None, arte=None, linhas_abertas=None):
    nl, nc = min(spec["rows"], linhas or spec["rows"]), spec["cols"]
    grupos = spec.get("grupos") or {"lin": [], "col": []}
    col_fechada = set()
    # `aberto` é até que profundidade os grupos de colunas estão abertos: 0 = como a aba nasce, 1 = o painel aberto
    # com a extensão fechada, 2 = tudo aberto
    for g in grupos["col"]:
        if int(aberto) < (g[3] if len(g) > 3 else 1):
            col_fechada |= set(range(g[0], g[1] + 1))
    largs = [spec["larg"]] * (nc + 1)
    for a, b, px in spec.get("largs", []):
        for c in range(a, min(b, nc) + 1):
            largs[c] = px
    alts = [int(spec["alturas"].get(str(r), 21)) for r in range(nl + 1)]
    # os grupos de linhas que nascem fechados somem, como no Sheets: a linha fica com altura zero
    lin_fechada = set()
    if linhas_abertas == "nasce":
        for g in grupos.get("lin", []):
            if g[2]:
                lin_fechada |= set(range(g[0], g[1] + 1))
    for r in lin_fechada:
        if r <= nl:
            alts[r] = 0
    bg = [[spec.get("fundo_base", "#120F1D")] * (nc + 1) for _ in range(nl + 1)]
    for r, c1, c2, cor in spec["fundos"]:
        if r <= nl:
            for c in range(c1, c2 + 1):
                bg[r][c] = cor
    if pintura:
        for r in range(1, nl + 1):
            for c in range(1, nc + 1):
                bg[r][c] = pintura["bg"][r - 1][c - 1]
    regua = (pintura or {}).get("regua")
    pad = spec.get("padrao", ["Roboto", 11, "#F4F1F7"])
    est, val = {}, {}
    for t in spec["vals"]:
        val[(t[0], t[1])] = t[2]
        if len(t) > 3:
            est[(t[0], t[1])] = spec["estilos"][t[3]]
    dono, cantos = {}, {}
    for r1, c1, r2, c2 in spec["merges"]:
        if r1 > nl:
            continue
        r2 = min(r2, nl)
        cantos[(r1, c1)] = (r2, c2)
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                dono[(r, c)] = (r1, c1)
    # a borda, lado a lado: cada faixa do ABAS pinta o lado de fora dela
    lados = {}
    for lado, traco, cor, faixas in spec.get("bordas", []):
        px = {"medium": 2, "thick": 3}.get(traco, 1)
        for f in faixas:
            a, _, b = f.partition(":")
            (r1, c1), (r2, c2) = ix._lc(a), ix._lc(b or a)
            if regua and cor.upper() == "#8A7EC4":
                cor = regua
            for r in range(r1, min(r2, nl) + 1):
                for c in range(c1, c2 + 1):
                    if (lado == "top" and r == r1) or (lado == "bottom" and r == r2) or (lado == "left" and c == c1) or (lado == "right" and c == c2):
                        lados.setdefault((r, c), {})[lado] = (px, cor)
    menus = set()
    for onde, _ in spec.get("dv", []):
        for parte in onde.split():
            a, _, b = parte.partition(":")
            (r1, c1), (r2, c2) = ix._lc(a), ix._lc(b or a)
            for r in range(r1, r2 + 1):
                menus.add((r, c1))
    caixas = {(lin + i, col) for col, lin, n in spec.get("caixas", []) for i in range(n)}
    notas = {ix._lc(a): t for a, t in spec.get("notas", [])}
    avisos = []
    for regra in spec.get("condicional", []):
        for a1 in regra["faixas"]:
            ini, _, fim = a1.partition(":")
            avisos.append((ix._lc(ini), ix._lc(fim or ini), regra))
    fmt = dict((ix._lc(a), f_) for a, f_ in spec.get("formatos", []))
    cols = [c for c in range(1, nc + 1) if c not in col_fechada]
    grade_c = " ".join(f"{largs[c]}px" for c in cols)
    grade_r = " ".join(f"{alts[r]}px" for r in range(1, nl + 1))
    pos = {c: i + 1 for i, c in enumerate(cols)}
    out = []
    for r in range(1, nl + 1):
        for c in cols:
            if (r, c) in dono and dono[(r, c)] != (r, c):
                continue
            r2, c2 = cantos.get((r, c), (r, c))
            vis = [k for k in range(c, c2 + 1) if k in pos]
            if not vis:
                continue
            e = est.get((r, c))
            fonte, tam, cor = (e[0], e[1], e[2] or pad[2]) if e else pad
            if pintura:
                cor = pintura["fc"][r - 1][c - 1]
            alinh, valinh = (e[4], e[5]) if e else ("left", "middle")
            v = valores.get((r, c))
            bruto = crus.get((r, c))
            if (r, c) in caixas:
                txt = '<span class="cx%s"></span>' % (" on" if v in (True, 1) and v is not False else "")
            elif barras and (r, c) in barras:
                pct, corb = barras[(r, c)]
                txt = f'<span class="barra" style="width:{pct}%;background:{corb}"></span>'
                alinh = "left"
            else:
                txt = html.escape("" if v is None else str(numero(v)))
            if (r, c) in fmt and isinstance(v, (int, float)) and not isinstance(v, bool):
                # o texto do formato vem antes do número ("Classe "0), depois dele (0" pt") ou antes, com o sinal ("B/D "+0)
                f_ = fmt[(r, c)].split(";")[0]
                antes, depois = re.match(r'(?:"([^"]*)")?[+]?[#,0]+(?:"([^"]*)")?$', f_).groups()
                num = f"{int(v):,}".replace(",", ".") if "#" in f_ else ("+" if "+" in f_ and int(v) > 0 else "") + str(int(v))
                txt = (antes or "") + num + (depois or "")
            fundo, corf, riscado = bg[r][c], cor, False
            for (r1_, c1_), (r2_, c2_), regra in avisos:
                texto = str(v or "")
                if not (r1_ <= r <= r2_ and c1_ <= c <= c2_):
                    continue
                if regra.get("formula"):
                    # 02/10/2026: a única regra por fórmula é a das Habilidades, '=LEFT($D$92,4)="Abre"' (ver habilidades.py)
                    m = re.fullmatch(r'=LEFT\(\$([A-Z]+)\$(\d+),(\d+)\)="(.*)"', regra["formula"])
                    bate = bool(m) and str(valores.get(ix._lc(m.group(1) + m.group(2)), "") or "")[:int(m.group(3))] == m.group(4)
                else:
                    bate = texto.startswith(regra["comeca"]) if regra.get("comeca") else regra["contem"] in texto
                if bate:
                    fundo, corf, riscado = regra.get("fundo", fundo), regra.get("fonte", corf), bool(regra.get("riscado"))
                    break
            if all(rr in lin_fechada for rr in range(r, r2 + 1)):
                continue
            b = {}
            for rr in range(r, r2 + 1):
                for cc in range(c, c2 + 1):
                    for lado, x in lados.get((rr, cc), {}).items():
                        if (lado == "top" and rr == r) or (lado == "bottom" and rr == r2) or (lado == "left" and cc == c) or (lado == "right" and cc == c2):
                            b[lado] = x
            css = [f"grid-area:{r}/{pos[vis[0]]}/{r2 + 1}/{pos[vis[-1]] + 1}", f"background:{fundo}", f"color:{corf}",
                   f"font-family:'{fonte}',sans-serif", f"font-size:{tam}pt",
                   "justify-content:" + {"left": "flex-start", "center": "center", "right": "flex-end"}.get(alinh, "flex-start"),
                   "align-items:" + {"top": "flex-start", "middle": "center", "bottom": "flex-end"}.get(valinh, "center")]
            if e and e[8]:
                css.append("white-space:pre-wrap")      # quebra como o Sheets: na largura, e em cada quebra de linha do texto
            if riscado:
                css.append("text-decoration:line-through")
            if e and e[6]:
                txt = f'<span style="writing-mode:vertical-rl;transform:rotate(180deg)">{txt}</span>'
            for lado, (px, corb) in b.items():
                css.append(f"box-shadow:none;border-{lado}:{px}px solid {corb}")
            marcas = ('<i class="menu">▾</i>' if (r, c) in menus else "") + ('<i class="nota"></i>' if (r, c) in notas else "")
            formula = isinstance(bruto, str) and bruto.startswith("=")
            out.append(f'<div class="c{" f" if formula else ""}" style="{";".join(css)}" title="{html.escape(notas.get((r, c), ""))}">{txt}{marcas}</div>')
    # a arte (a pincelada do cabeçalho): um traço no lugar dela, da cor da régua
    for l1, c1, l2, c2, nome in spec.get("imgs", []):
        vis = [k for k in range(c1, c2 + 1) if k in pos]
        if vis and l1 <= nl and arte and nome in arte:
            cor = ((pintura or {}).get("imgs") or {}).get(f"{l1},{c1}") or "#" + FABRICA_DA_ARTE.get(re.sub(r"-\d+x\d+\.png$", "", nome), "8A7EC4")
            w, h = (int(x) for x in re.search(r"-(\d+)x(\d+)\.png$", nome).groups())
            out.append(f'<div class="c" style="grid-area:{l1}/{pos[vis[0]]}/{min(l2, nl) + 1}/{pos[vis[-1]] + 1};padding:0;align-items:center;justify-content:center">'
                       f'<span title="{html.escape(nome)} {cor}" style="display:block;width:{w}px;height:{h}px;max-width:100%;max-height:100%;background:{cor};'
                       f'-webkit-mask:url(data:image/png;base64,{arte[nome]}) center/contain no-repeat;mask:url(data:image/png;base64,{arte[nome]}) center/contain no-repeat"></span></div>')
        elif vis and l1 <= nl:
            out.append(f'<div class="c" style="grid-area:{l1}/{pos[vis[0]]}/{l2 + 1}/{pos[vis[-1]] + 1};align-items:center">'
                       f'<span class="arte" title="{html.escape(nome)}"></span></div>')
    larg_total = sum(largs[c] for c in cols)
    return (f'<div class="aba" style="grid-template-columns:{grade_c};grid-template-rows:{grade_r};width:{larg_total}px">'
            + "".join(out) + "</div>")


# a cor com que cada imagem nasce, antes de qualquer troca de paleta (ficha-v01/arte)
FABRICA_DA_ARTE = {"carteira-1": "756588", "carteira-2": "756588", "carteira-3": "998BA9", "carteira-4": "8A7EC4",
                   "ficha-1": "756588", "ficha-2": "8A7EC4", "carteira-canto": "8A7EC4"}
CSS = """
.aba{display:grid}
.c{position:relative;box-sizing:border-box;display:flex;overflow:hidden;white-space:nowrap;padding:0 3px;min-width:0;line-height:1.15}
.cx{display:block;width:11px;height:11px;border:1.5px solid currentColor;border-radius:2px}
.cx.on{background:currentColor;position:relative}
.cx.on::after{content:"✓";position:absolute;left:1px;top:-3px;font:bold 10px sans-serif;color:#0A0810;mix-blend-mode:difference}
.barra{display:block;height:100%;min-width:1px}
.arte{display:block;width:100%;height:5px;border-radius:3px;background:linear-gradient(90deg,#8A7EC4,#3D2E78 70%,transparent)}
.menu{position:absolute;right:3px;top:50%;transform:translateY(-50%);font:normal 9px sans-serif;opacity:.6}
.nota{position:absolute;top:0;right:0;border-style:solid;border-width:0 7px 7px 0;border-color:transparent currentColor transparent transparent}
"""
FONTES = ('<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Castoro&family=Oswald:wght@400;500'
          '&family=Roboto:wght@400;500&family=Courier+Prime&family=Yuji+Syuku&display=swap">')


def pagina(corpo, titulo="", fundo="#0A0810"):
    return (f'<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><title>{html.escape(titulo)}</title>{FONTES}'
            f'<style>body{{margin:0;background:{fundo}}}{CSS}</style></head><body>{corpo}</body></html>')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--aba", default=fp.NOME)
    ap.add_argument("--saida", default="/tmp/ver-aba.html")
    ap.add_argument("--fechado", action="store_true", help="o painel de XP fechado, como a aba nasce")
    ap.add_argument("--vazia", action="store_true", help="a ficha de fábrica, sem o exemplo")
    ap.add_argument("--tudo", action="store_true", help="a extensão do painel aberta também")
    ap.add_argument("--nasce", action="store_true", help="os grupos de linhas como a aba nasce (os fechados, fechados)")
    a = ap.parse_args()
    # 02/10/2026: a FICHA vai com a Kaori da Ficha Amaldiçoada, para o menu rápido da seção 8 ter o que mostrar
    if a.aba in (fa.NOME, "FICHA"):
        lido, cru = recalculada((lambda wb: None) if a.vazia else exemplo_amaldicoada)
        ws, wc = lido[a.aba], cru[a.aba]
        valores = {(c.row, c.column): c.value for linha in ws.iter_rows() for c in linha if c.value is not None}
        crus = {(c.row, c.column): c.value for linha in wc.iter_rows() for c in linha if c.value is not None}
        spec = next(s for s in abas_do_script() if s["nome"] == a.aba)
        if a.aba == "FICHA":       # a caixa das Habilidades como o script a estica, e não como a aba nasce
            spec["alturas"].update({str(r): px for r, px in ALTURAS_DO_SCRIPT.items()})
        open(a.saida, "w", encoding="utf-8").write(desenha(spec, valores, crus, 2, None, a.aba, arte=arte_do_script(),
                                                           linhas_abertas="nasce" if a.nasce else None))
        print(f"{a.aba}: {spec['rows']} linhas, {spec['cols']} colunas -> {a.saida}")
        return
    # 06/10/2026: a INVOCAÇÕES vai com o Cão de sombra do livro; a barra de vida é desenhada pela conta dela
    if a.aba == fi.NOME:
        lido, cru = recalculada((lambda wb: None) if a.vazia else exemplo_invocacoes)
        ws, wc, dd = lido[a.aba], cru[a.aba], lido[fi.DADOS_IV]
        valores = {(c.row, c.column): c.value for linha in ws.iter_rows() for c in linha if c.value is not None}
        crus = {(c.row, c.column): c.value for linha in wc.iter_rows() for c in linha if c.value is not None}
        # a primeira coluna com cada nome: a tabela das fichas vem antes da das cartas, que repete "tem"
        cab = {}
        for c in range(1, dd.max_column + 1):
            cab.setdefault(dd.cell(row=1, column=c).value, c)
        barras = {}
        for i, (_, j, k) in enumerate(fi.lugares()):
            atual, vida = (dd.cell(row=2 + i, column=cab[x]).value for x in ("atual", "vida"))
            tem = dd.cell(row=2 + i, column=cab["tem"]).value
            if tem and vida:
                barras[ix._lc(fi.celulas_da_ficha(fi.linha_da_fileira(j), k)["barra"])] = (round(100 * atual / vida), "#E8DCD4")
        spec = next(s for s in abas_do_script() if s["nome"] == a.aba)
        open(a.saida, "w", encoding="utf-8").write(desenha(spec, valores, crus, 2, barras, a.aba, arte=arte_do_script(),
                                                           linhas_abertas="nasce" if a.nasce else None))
        print(f"{a.aba}: {spec['rows']} linhas, {spec['cols']} colunas -> {a.saida}")
        return
    CAT = json.load(open(os.path.join(RAIZ, "catalogo-projeto-m.json"), encoding="utf-8"))
    R = fp.regras(CAT)
    G = fp.geometria(R)
    lido, cru = recalculada((lambda wb: None) if a.vazia else (lambda wb: exemplo(wb, R, G)))
    ws, wc = lido[a.aba], cru[a.aba]
    valores = {(c.row, c.column): c.value for linha in ws.iter_rows() for c in linha if c.value is not None}
    crus = {(c.row, c.column): c.value for linha in wc.iter_rows() for c in linha if c.value is not None}
    barras = {}
    if a.aba == fp.NOME:
        # as duas barras, pela conta que a SPARKLINE faz: a carga contra o limite, e o XP dentro do nível
        m = re.match(r"([\d.,]+) de (\d+)", str(valores.get(ix._lc(G["carga"]), "")))
        if m:
            carga, limite = float(m.group(1).replace(",", ".")), float(m.group(2))
            barras[ix._lc(G["carga_barra"])] = (min(100, round(100 * carga / limite)), "#C2334D" if carga > limite else "#E8DCD4")
        total, falta = valores.get(ix._lc(G["xp_total"])), valores.get(ix._lc(G["falta"]))
        if isinstance(total, (int, float)) and isinstance(falta, (int, float)) and total + falta:
            barras[ix._lc(G["falta_barra"])] = (round(100 * total / (total + falta)), "#E8DCD4")
    spec = next(s for s in abas_do_script() if s["nome"] == a.aba)
    open(a.saida, "w", encoding="utf-8").write(desenha(spec, valores, crus, 0 if a.fechado else 2 if a.tudo else 1, barras, a.aba,
                                                       arte=arte_do_script()))
    print(f"{a.aba}: {spec['rows']} linhas, {spec['cols']} colunas -> {a.saida}")


if __name__ == "__main__":
    main()
