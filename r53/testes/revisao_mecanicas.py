"""Planos offline para QA descartável; oráculos aritméticos do livro R53.

Não envia pedidos ao Google. Matrizes editam os cartões reais de uma cópia,
sem substituir as fórmulas calculadas por uma segunda implementação.
"""
import argparse
import json
import math
from pathlib import Path
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter

OUTPUTS=['orcamento','base_funcoes','custo_forma','saldo','dano_d8','cura_d8','vida_temporaria','pe','valida','avisos']


def typed(v):
    if isinstance(v,bool):return {'boolValue':v}
    if isinstance(v,(int,float)):return {'numberValue':v}
    return {'stringValue':v}


def matriz(modelo, limites=False):
    tr=modelo['r53'];tabs={a['nome']:a for a in modelo['layout']['abas']}
    sid={'FICHA':2091222604,'FICHA AMALDIÇOADA':543479115,'MONTAGENS':53018003}
    requests=[];cases=[]
    def put(sh,addr,v):
        r,c=coordinate_to_tuple(addr)
        requests.append({'updateCells':{'start':{'sheetId':sid[sh],'rowIndex':r-1,'columnIndex':c-1},'rows':[{'values':[{'userEnteredValue':typed(v)}]}],'fields':'userEnteredValue'}})
    put('FICHA','AH11',30)
    contexts=[g for g in tr['contextos'] if g['modo']=='Comum' and g['aba']=='FICHA AMALDIÇOADA']
    assert len(contexts)>=35
    specs=[dict(classe=c,tipo=k) for c in range(1,8) for k in ['dano','cura','vt','coletiva','mista']]
    if limites:
        specs=[]
        for c in range(1,8):
            specs += [dict(classe=c,tipo='toque_cura'),dict(classe=c,tipo='toque_precisao'),dict(classe=c,tipo='toque_pessoal'),dict(classe=c,tipo='coletiva_vt')]
        specs += [
            dict(classe=1,tipo='negativo',forma='Explosão',pontos1=0,pontos2=1,m1='Rápido',esperado={'saldo':-1,'valida':False},contem='Orçamento excedido'),
            dict(classe=3,tipo='incompativel',forma='Toque',pontos1=7,m1='Longe · um degrau',esperado={'valida':False},contem='Toque/Aura não aceita Longe'),
            dict(classe=3,tipo='repetida',pontos1=5,m1='Precisão',m2='Precisão',esperado={'valida':False},contem='Compra repetida sem permissão'),
            dict(classe=3,tipo='distribuicao',pontos1=10,esperado={'valida':False},contem='Distribuição excede o saldo'),
            dict(classe=3,tipo='fracionaria',pontos1=1.5,esperado={'valida':False},contem='Distribua pontos inteiros não negativos'),
            dict(classe=3,tipo='negativa',pontos1=-1,esperado={'valida':False},contem='Distribua pontos inteiros não negativos'),
            dict(classe=3,tipo='coletivo_excede',forma='Aura',pontos1=0,pontos2=3,pontos3=2,esperado={'valida':False},contem='Cura e vida temporária excedem juntas o teto coletivo'),
            dict(classe=3,tipo='guarda',forma='Aura',pontos1=1,m1='Guarda',esperado={'base_funcoes':3,'custo_forma':2,'melhorias_pagas':3,'saldo':1,'beneficiarios_fixos':4,'valida':True}),
        ]
    for i,spec in enumerate(specs):
        c,kind=spec['classe'],spec['tipo']
        g=contexts[i];base=c if kind in ['cura','coletiva','mista'] else math.ceil(c/2) if kind=='vt' else 0
        form=math.ceil(c/2) if kind=='coletiva' else 0
        saldo=3*c-base-form
        pd=saldo if kind=='dano' else c if kind=='mista' else 0
        pc=saldo if kind in ['cura','coletiva'] else c if kind=='mista' else 0
        pv=saldo if kind=='vt' else 0
        title=f'QA {kind} C{c}'
        values={'forma':'Aura' if kind=='coletiva' else 'Projétil','pessoal':False,'pontos1':pd,'pontos2':pc,'pontos3':pv,
                'm1':'','m2':'','m3':'','m4':'','r1':'','r2':'','controle':'Sem reforço','fragil_pago':0}
        if kind.startswith('toque_'):
            paid=math.ceil(c/2) if kind!='toque_cura' else 0
            refund=paid if kind=='toque_precisao' else 0
            base=c;form=0;saldo=3*c-base-paid+refund;pd=0;pc=saldo;pv=0
            values.update(forma='Toque',pessoal=kind=='toque_pessoal',pontos1=0,pontos2=pc,pontos3=0,m1='Precisão' if paid else '')
        elif kind=='coletiva_vt':
            base=c;form=math.ceil(c/2);saldo=3*c-base-form;pd=pc=0;pv=saldo
            values.update(forma='Aura',pontos1=0,pontos2=0,pontos3=pv)
        values.update({k:v for k,v in spec.items() if k in values})
        pd,pc,pv=[values['pontos'+str(n)] for n in range(1,4)]
        put(g['aba'],g['cel']['classe'],c);put(g['aba'],g['cel']['nome'],title)
        for field,addr in g['entradas_diretas'].items():put(g['aba'],addr,values[field])
        # Preparação de uma versão, não um ensaio de onEdit/persistência.
        col=get_column_letter(g['classes'].index(c)+2)
        store={'preparada':True,'forma':values['forma'],'funcao1':'Dano' if pd else '',
               'funcao2':'Cura' if pc else '', 'funcao3':'PV temporários' if pv else '',
               'pontos1':pd,'pontos2':pc,'pontos3':pv}
        for v in g['classes']:put('MONTAGENS',get_column_letter(g['classes'].index(v)+2)+str(g['campos']['preparada']),False)
        for field,value in store.items():put('MONTAGENS',col+str(g['campos'][field]),value)
        expected={'orcamento':3*c,'base_funcoes':base,'custo_forma':form,'saldo':saldo,
                  'dano_d8':pd,'cura_d8':pc,'vida_temporaria':3*pv,'pe':3*c,'valida':True}
        if kind.startswith('toque_'):expected.update(melhorias_pagas=paid,devolucao=refund)
        if 'esperado' in spec:expected=spec['esperado']
        keys=list(dict.fromkeys(OUTPUTS+list(expected)))
        cases.append({'nome':title,'id':g['id'],'linha':g['linha'],'esperado':expected,
                      'contem':spec.get('contem',''),
                      'leituras':{k:"'DADOS_R53'!"+tr['colunas'][k]+str(g['linha']) for k in keys}})
    ranges=[]
    for key in dict.fromkeys(k for case in cases for k in case['leituras']):
        col=tr['colunas'][key];ranges.append(f"'DADOS_R53'!{col}2:{col}{len(cases)+1}")
    ranges+=['\'DADOS_CRIACAO\'!H2:I10',"'FICHA'!AH11"]
    return dict(casos=cases,pedidos=requests,faixas=ranges,escopo='Fórmulas nativas dos cartões; não testa eventos de edição')


def cells(snapshot):
    out={}
    for sh in snapshot['sheets']:
        for b in sh.get('data',[]):
            for i,row in enumerate(b.get('rowData',[])):
                for j,cell in enumerate(row.get('values',[])):
                    addr=get_column_letter(b.get('startColumn',0)+j+1)+str(b.get('startRow',0)+i+1)
                    ev=cell.get('effectiveValue',{})
                    out[sh['properties']['title']+'!'+addr]=ev.get('errorValue') or next(iter(ev.values()),'')
    return out


def matriz_entidades(modelo):
    """Ensaios aritméticos dos 12 grupos; não aprovam um repertório de 12 criaturas."""
    tr=modelo['r53'];requests=[];cases=[]
    sid={'FICHA':2091222604,'INVOCAÇÕES':996249592,'MONTAGENS':53018003}
    def put(sh,addr,value):
        row,col=coordinate_to_tuple(addr)
        requests.append({'updateCells':{'start':{'sheetId':sid[sh],'rowIndex':row-1,'columnIndex':col-1},'rows':[{'values':[{'userEnteredValue':typed(value)}]}],'fields':'userEnteredValue'}})
    put('FICHA','AH11',30)
    inv=next(a for a in modelo['layout']['abas'] if a['nome']=='INVOCAÇÕES')
    for family in inv['familias_custom']:
        for addr,value in zip(family['vagas'],['Amparo','Área','Mira','','','','','','']):put('INVOCAÇÕES',addr,value)
        row,col=coordinate_to_tuple(family['vagas'][0]);put('INVOCAÇÕES',get_column_letter(col-6)+str(row-35),'QA entidade '+str(family['entidade']))
    for entity in range(12):
        groups=[g for g in tr['contextos'] if g.get('entidade')==entity and g['modo']=='Entidade' and '.esp.' in g['id']][:7]
        basic=[g for g in tr['contextos'] if g.get('entidade')==entity and g['modo']=='Básica'][:2]
        assert len(groups)==7 and len(basic)==2
        specs=[(g,c,'cura') for c,g in enumerate(groups,1)]+[(basic[0],0,'basica'),(basic[1],0,'media_proibida')]
        for g,c,kind in specs:
            budget=[0,2,4,6,9,11,13,16][c] if c else 3
            heal=budget-c if c else 0
            values=dict(forma='Projétil',pessoal=False,pontos1=0,pontos2=heal,pontos3=0,usar_dano=c==0,usar_cura=False,usar_vt=False,m1='Anteparo' if kind=='media_proibida' else '',m2='',m3='',m4='',r1='',r2='',controle='Sem reforço',fragil_pago=0)
            name=f'QA E{entity+1} {kind} C{c}'
            if len(g['classes'])>1:put(g['aba'],g['cel']['classe'],c)
            put(g['aba'],g['cel']['nome'],name)
            for k,addr in g['entradas_diretas'].items():put(g['aba'],addr,values[k])
            col=get_column_letter(g['classes'].index(c)+2)
            for cl in g['classes']:put('MONTAGENS',get_column_letter(g['classes'].index(cl)+2)+str(g['campos']['preparada']),False)
            for k,value in dict(preparada=True,forma='Projétil',funcao1='Dano' if c==0 else '',pontos1=0,funcao2='Cura' if c else '',pontos2=heal,funcao3='',pontos3=0).items():put('MONTAGENS',col+str(g['campos'][k]),value)
            expected={'orcamento':budget,'base_funcoes':c,'saldo':heal if c else 3,'cura_d8':heal,'vida_temporaria':0,'pe':3*c,'valida':True}
            contains=''
            if kind=='basica':expected.update(dano_d8=3,dado_dano='d6')
            if kind=='media_proibida':expected={'valida':False};contains='Classe 0 aceita somente Melhoria Leve'
            keys=list(dict.fromkeys(OUTPUTS+list(expected)))
            cases.append(dict(nome=name,id=g['id'],linha=g['linha'],esperado=expected,contem=contains,leituras={k:"'DADOS_R53'!"+tr['colunas'][k]+str(g['linha']) for k in keys}))
    keys=dict.fromkeys(k for case in cases for k in case['leituras'])
    ranges=[f"'DADOS_R53'!{tr['colunas'][k]}47:{tr['colunas'][k]}202" for k in keys]
    return dict(casos=cases,pedidos=requests,faixas=ranges,escopo='Fórmulas nativas de 12 grupos isolados; não testa aquisição, limite global ou onEdit')


def conferir(plan,snapshot):
    actual=cells(snapshot);checks=0;fail=[]
    for case in plan['casos']:
        for key,wanted in case['esperado'].items():
            addr=case['leituras'][key].replace("'",'');got=actual.get(addr,'');checks+=1
            if type(got) is not type(wanted) or got!=wanted:
                # API numbers may be represented as integer or float.
                if not (type(got) in (int,float) and type(wanted) in (int,float) and got==wanted):
                    fail.append(dict(caso=case['nome'],campo=key,esperado=wanted,obtido=got,aviso=actual.get(case['leituras']['avisos'].replace("'",''))))
        if case.get('contem'):
            checks+=1;got=actual.get(case['leituras']['avisos'].replace("'",''),'')
            if case['contem'] not in str(got):fail.append(dict(caso=case['nome'],campo='avisos',esperado=case['contem'],obtido=got))
    return dict(cenarios=len(plan['casos']),verificacoes=checks,falhas=fail)


def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--modelo',type=Path);p.add_argument('--limites',action='store_true');p.add_argument('--entidades',action='store_true');p.add_argument('--plano',type=Path);p.add_argument('--resposta',type=Path);p.add_argument('--saida',type=Path,required=True);a=p.parse_args()
    if a.modelo:
        model=json.loads(a.modelo.read_text());out=matriz_entidades(model) if a.entidades else matriz(model,a.limites)
    else:out=conferir(json.loads(a.plano.read_text()),json.loads(a.resposta.read_text()))
    a.saida.write_text(json.dumps(out,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:v for k,v in out.items() if k not in ['casos','pedidos','faixas']},ensure_ascii=False))

if __name__=='__main__':main()
