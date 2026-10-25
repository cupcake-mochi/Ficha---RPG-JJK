import sys, unittest, random
from pathlib import Path
ROOT = Path(__file__).resolve().parents[2]
sys.path[:0] = [str(ROOT/'r53'), str(ROOT/'r53/testes'), str(ROOT/'ficha-v01')]
import modelo as M
from formulas import X, evaluate, native
from compactar import fatorar_sufixos, compartilhar_buscas, native_compact
from verificar import CASOS, BASE
from openpyxl import Workbook
from integridade_celula import atribuir


class Compactacao(unittest.TestCase):
    def test_equivalencia_arvores_e_buscas(self):
        rng = random.Random(531023)
        corpus = [(dict(BASE, **c['entrada']), c) for c in CASOS]
        for i in range(1024):
            corpus.append((dict(BASE, classe=rng.randrange(8),
                modalidade=rng.choice(['Comum','Classe 0','Entidade','Básica','Técnica Máxima','Máxima de entidade','Liberação Máxima','Liberação de entidade']),
                controle=rng.choice(['Sem reforço','Sem dano','Dano limitado à Classe']),
                **{f'm{j}':rng.choice(list(M.OPCOES)+['','* ? ~ desconhecida']) for j in range(1,5)}), {}))
        refs = {k:'A1' for k in M.INPUTS+list(M.ARVORE)}
        trees = [M.arvore(False), M.arvore(True)]
        for inputs, case in corpus:
            env = M.calcular(inputs, case.get('catalogo'), case.get('familias_custom'), case.get('formas'))
            catalog = dict(M.OPCOES, **case.get('catalogo', {}))
            for tree in trees:
                for field in ('efeitos','avisos'):
                    expr = tree[field]
                    before = evaluate(expr, env, catalog)
                    if field == 'efeitos': expr = fatorar_sufixos(expr)
                    bindings, compact = compartilhar_buscas(expr, refs, 'CAT!A2:O200', M.LOOKUP_FIELDS)
                    trial = dict(env)
                    for name, node in bindings:
                        trial[name] = evaluate(node, env, catalog)
                    self.assertEqual(before, evaluate(compact, trial, catalog))

    def test_emissao_menor_e_identificadores(self):
        refs = {k:'A1' for k in M.INPUTS+list(M.ARVORE)}
        for sep in (',',';'):
            for livre in (False, True):
                for field in ('efeitos','avisos'):
                    e = M.arvore(livre)[field]
                    old = native(e,refs,'CAT!A2:O200',M.LOOKUP_FIELDS,sep)
                    new = native_compact(e,refs,'CAT!A2:O200',M.LOOKUP_FIELDS,sep,field=='efeitos')
                    self.assertLess(len(new), len(old))
                    self.assertLess(len(new),32767)
                    self.assertTrue(new.startswith('LET('))

    def test_integridade_textos_e_matrizes(self):
        ws = Workbook().active
        for v, array in [('=1+2',None),('texto',None),('=SUM(A1:A2)','B1:B2')]:
            atribuir(ws['B1'], v, array)
        for v in ('x'*32768, '='+'1+'*17000+'1'):
            for array in (None,'B1:B2'):
                with self.assertRaises(ValueError): atribuir(ws['B1'],v,array)


if __name__ == '__main__': unittest.main()
