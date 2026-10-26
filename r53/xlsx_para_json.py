import openpyxl,json,sys,datetime
src=sys.argv[1]
f=openpyxl.load_workbook(src); v=openpyxl.load_workbook(src,data_only=True)
out={}
def ser(x):
  if isinstance(x,datetime.datetime): return {'d':x.isoformat()}
  if isinstance(x,datetime.date): return {'d':x.isoformat()}
  if isinstance(x,datetime.time): return str(x)
  if hasattr(x,'text'): return '='+x.text if not str(x.text).startswith('=') else x.text
  return x
for ws in f.worksheets:
  wv=v[ws.title]; cells={}
  for row in ws.iter_rows():
    for c in row:
      fv=c.value; vv=wv.cell(c.row,c.column).value
      note=c.comment.text if c.comment else None
      if fv is None and vv is None and not note: continue
      e={}
      if isinstance(fv,str) and fv.startswith('=') or hasattr(fv,'text'): e['f']=ser(fv); e['v']=ser(vv)
      else: e['v']=ser(fv)
      if note: e['n']=note
      cells[f'{c.row},{c.column}']=e
  out[ws.title]={'rows':ws.max_row,'cols':ws.max_column,'hidden':ws.sheet_state!='visible','cells':cells}
json.dump(out,open(sys.argv[2],'w'))
print(len(out),sum(len(s['cells']) for s in out.values()))
