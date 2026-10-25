"""Prepara redução de FD na cópia B; nunca transmite nem modifica a oficial.

Requer snapshots nativos atuais em diretório privado. Confere cada fórmula
contra o modelo completo antes de gerar dois pedidos (semente + copyPaste).
Não passa pelo Workbook, que limita textos longos antes da emissão.
"""
import hashlib
import json
from pathlib import Path

from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import get_column_letter

from ensaiar_fatoracao import M, candidate, refs
from revisar import normalized

ROOT = Path(__file__).resolve().parents[2]
PRIVATE = ROOT / 'nuvem-b47/saidas/diagnostico-mobile-20261023/android-ab'
OFFICIAL = '1u0zWQFVaRTTc_VlrAJsabgAiJ1D4rTYhzG7QQWof1ro'
CONTROL = '1mhrotVJIwblthbTT_lCoau8OmKYDezA0TphTSusUZqc'
TARGET = '1rH1csUCpUWBoYRmyw4FdQahzdRqgcP9wk3Jer0rRUBE'


def cells(document):
    result = {}
    for sheet in document['sheets']:
        for data in sheet.get('data', []):
            for row, item in enumerate(data.get('rowData', []), data.get('startRow', 0)+1):
                for col, cell in enumerate(item.get('values', []), data.get('startColumn', 0)+1):
                    result[get_column_letter(col)+str(row)] = cell
    return result


def digest(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def main():
    meta = json.loads((PRIVATE/'otimizada-metadados.json').read_text())
    assert meta['file']['id'] == TARGET and TARGET not in (OFFICIAL, CONTROL)
    sheet = next(s['properties'] for s in meta['metadata']['sheets'] if s['properties']['title'] == 'DADOS_R53')
    assert sheet['hidden'] and sheet['gridProperties']['rowCount'] == 251
    before = {}
    for start in range(1, 251, 25):
        doc = json.loads((PRIVATE/f'antes-{start}.json').read_text())
        assert doc['spreadsheetId'] == TARGET
        block = cells(doc)
        assert not before.keys() & block.keys()
        before.update(block)
    assert before['FD1']['userEnteredValue'] == {'stringValue': 'efeitos'}
    trial, _, _ = candidate()
    assert refs(M.ARVORE['efeitos']) == refs(trial.ARVORE['efeitos'])
    cols = {key: get_column_letter(i+1) for i, key in enumerate(M.INPUTS+list(M.ARVORE))}
    assert cols['efeitos'] == 'FD' and cols['nome'] == 'B'
    old, new = {}, {}
    for row in range(2, 251):
        address = f'FD{row}'
        cell = before[address]
        assert not cell.get('dataValidation') and 'errorValue' not in cell.get('effectiveValue', {})
        old[address] = cell['userEnteredValue']['formulaValue']
        def formula(model):
            inner = model.formulas(row, cols, "'DADOS_R53_CATALOGO'!$A$2:$O$179", saidas={'efeitos'}, separator=';')['efeitos']
            return f'=IF(B{row}="";"";'+inner[1:]+')'
        assert normalized(old[address]) == normalized(formula(M)), address
        new[address] = formula(trial)
        assert len(new[address]) < 32767
        assert normalized(Translator(new['FD2'], origin='FD2').translate_formula(address)) == normalized(new[address])
    sid = sheet['sheetId']
    requests = [
        {'updateCells': {'range': {'sheetId': sid, 'startRowIndex': 1, 'endRowIndex': 2, 'startColumnIndex': 159, 'endColumnIndex': 160}, 'rows': [{'values': [{'userEnteredValue': {'formulaValue': new['FD2']}}]}], 'fields': 'userEnteredValue'}},
        {'copyPaste': {'source': {'sheetId': sid, 'startRowIndex': 1, 'endRowIndex': 2, 'startColumnIndex': 159, 'endColumnIndex': 160}, 'destination': {'sheetId': sid, 'startRowIndex': 2, 'endRowIndex': 250, 'startColumnIndex': 159, 'endColumnIndex': 160}, 'pasteType': 'PASTE_FORMULA'}}
    ]
    # Rollback usa as fórmulas exatas do snapshot, nunca a fonte truncada.
    rollback = [{'updateCells': {'range': {'sheetId': sid, 'startRowIndex': 1, 'endRowIndex': 250, 'startColumnIndex': 159, 'endColumnIndex': 160}, 'rows': [{'values': [{'userEnteredValue': {'formulaValue': old[f'FD{r}']}}]} for r in range(2, 251)], 'fields': 'userEnteredValue'}}]
    manifest = {'target': TARGET, 'controle': CONTROL, 'oficial': OFFICIAL, 'range': 'DADOS_R53!FD2:FD250', 'formulas': len(new), 'pedidos': len(requests), 'bytes_antes': sum(len(f.encode()) for f in old.values()), 'bytes_candidatos': sum(len(f.encode()) for f in new.values()), 'antes_sha256': digest(old), 'candidato_sha256': digest(new), 'pedidos_sha256': digest(requests), 'dependencias_iguais': True, 'originais_equivalentes_ao_modelo': True, 'translacao_249_linhas_conferida': True}
    for name, data in [('pedidos', requests), ('rollback', rollback), ('esperado', new), ('manifesto', manifest)]:
        (PRIVATE/f'{name}.json').write_text(json.dumps(data, ensure_ascii=False)+'\n')
    print(json.dumps(manifest, ensure_ascii=False))


if __name__ == '__main__':
    main()
