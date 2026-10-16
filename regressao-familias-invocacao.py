#!/usr/bin/env python3
"""Nove Famílias e autonomia: fórmulas reais, cenários esperados independentes.

O artefato de cálculo fica local. Não é importado no Google Sheets.
"""
import gzip
import json
import shutil
import subprocess
import tempfile
import sys
from pathlib import Path
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parent
EVID = ROOT.parent/'evidencias-3b'
pack = EVID/'layout-autonomia9-pacote.json'
subprocess.run(['python3', str(ROOT/'ficha-v01/monta.py'), '--exportar-progressao', str(pack)],
               check=True, stdout=subprocess.DEVNULL)
d = json.loads(pack.read_text())['layout']
abas = {a['nome']: a for a in d['abas']}
baseline = json.loads(gzip.decompress((ROOT/'ficha-v01/progressao-base3a.json.gz').read_bytes()))
old = {a['nome']: a for a in baseline}
raw = {n: {v[0]: v[1] for v in a['celulas']} for n, a in abas.items()}
sys.path[:0]=[str(ROOT/'ficha')]
import emitir_gs

def col(n):
    out = ''
    while n:
        n, r = divmod(n-1, 26)
        out = chr(65+r)+out
    return out

def num(s):
    return sum((ord(c)-64)*26**i for i, c in enumerate(s[::-1]))

generated={a['nome']:{col(c)+str(r):v for r,c,v,*_ in a['vals']}
           for a in emitir_gs.abas_do_script(str(ROOT/'apps-script/Ficha.gs'))}

old_raw = {name: {col(c)+str(r): v for r, c, v, *_ in a['vals']} for name, a in old.items()}

headers = {num(c[:-1]): v for c, v in raw['DADOS_INVOC'].items() if c.endswith('1') and c[:-1].isalpha()}
find = lambda text: min(c for c, v in headers.items() if v == text)
core = find('f0')
card = find('conta de habilidade')
fields = {v: c for c, v in headers.items() if card <= c <= 463}
nfam, rep = col(find('nfam'))+'2', col(find('rep'))+'2'
forma = find('forma')
cura = next(int(c[len(col(forma)):]) for c, v in raw['DADOS_INVOC'].items()
            if c.startswith(col(forma)) and c[len(col(forma)):].isdigit() and v == 'Cura')
peca = find('peça')
amparo = next(r for r in range(2, 85) if raw['DADOS_INVOC'].get(col(peca+1)+str(r)) == 'Amparo')
menu_f = col(find('menu de forma da ficha 1'))+str(cura)
menu_m = col(find('menu de melhoria da ficha 1'))+str(amparo)
x0, x1 = col(fields['x0'])+'2', col(fields['x1'])+'2'
targets = [nfam, rep, menu_f, menu_m, x0, x1, 'HF2', 'RA2', 'NT2', 'HF3']
familias = ['O48', 'P48', 'Q48', 'R48', 'S48', 'O50', 'P50', 'Q50', 'R50']

checks = []
norm = lambda v: None if v == '' else v
def check(name, result):
    print(('OK ' if result else 'FALHA ')+name, flush=True)
    checks.append(result)

check('108 vagas, sem deslocar as cinco primeiras de cada entidade', all(
    all(norm(raw['INVOCAÇÕES'].get(col(c)+str(48+198*j))) ==
        norm(old_raw['INVOCAÇÕES'].get(col(c)+str(48+198*j)))
        for c in range(15+19*k, 20+19*k)) and
    all(col(c)+str(50+198*j) in raw['INVOCAÇÕES'] for c in range(15+19*k,19+19*k))
    for k in range(2) for j in range(6)))
keys = [v for c, v in raw['ANALISE'].items() if c.startswith('A') and c[1:].isdigit() and c != 'A1']
check('7485 chaves únicas, incluindo autonomia, limites e 48 Famílias extras', len(keys) == 7485 and len(set(keys)) == 7485 and
      all(f'invocacao.{i}.familia.{j}' in keys for i in range(1, 13) for j in (6, 7, 8, 9)) and
      all(f'invocacao.{i}.autonomia' in keys for i in range(1,13)))
check('autonomia individual padrão Autônoma, sem seletor global',
      all(raw['INVOCAÇÕES'].get(col(11+19*k)+str(50+198*j))=='Autônoma' for k in range(2) for j in range(6)) and
      not any(v in ('Autônoma','Não autônoma') for v in raw['CONFIGURAÇÃO'].values()))
check('linha de respiro antes da Definição, sem deslocar habilidades',all(
      raw['INVOCAÇÕES'].get(col(9+19*k)+str(51+198*j)) is None and
      raw['INVOCAÇÕES'].get(col(9+19*k)+str(52+198*j))=='DEFINIÇÃO' and
      col(9+19*k)+str(53+198*j)+':'+col(25+19*k)+str(54+198*j) in abas['INVOCAÇÕES']['mescladas']
      for k in range(2) for j in range(6)))
check('todas as notas de autonomia citam o futuro Guia de Criação',all(
      'Guia de Criação' in abas['INVOCAÇÕES']['notas'][col(11+19*k)+str(50+198*j)]
      for k in range(2) for j in range(6)))
check('rótulos de Livre leem o limite da DADOS_INVOC, não células vazias da própria ficha',all(
      str(raw['INVOCAÇÕES'].get(col(c)+str(r))).startswith('=IF(DADOS_INVOC!$RA$')
      for k in range(2) for j in range(6)
      for r,c in [(47+198*j,c) for c in range(15+19*k,20+19*k)]+
                 [(49+198*j,c) for c in range(15+19*k,19+19*k)]))
check('108 rótulos gerados nos .gs iguais às fórmulas finais do gerador',all(
      generated['INVOCAÇÕES'].get(col(c)+str(r))==raw['INVOCAÇÕES'].get(col(c)+str(r))
      for k in range(2) for j in range(6)
      for r,c in [(47+198*j,c) for c in range(15+19*k,20+19*k)]+
                 [(49+198*j,c) for c in range(15+19*k,19+19*k)]))

scenarios = [
    # nome, escolhas, modo, Trilha, limite, Livres, contagem, duplicatas,
    # Amparo disponível, código de preço Médio de Amparo, campanha
    ('nona abre Amparo', ['Alcance','Área','Mira','Controle','Auxiliares','Castigo','','','Amparo'], 'Autônoma','',3,1,7,0,True,2,{}),
    ('oitava repete primeira', ['Alcance','','','','','','','Alcance','Amparo'], 'Autônoma','',3,1,3,2,True,2,{}),
    ('duplicata entre oitava e nona', ['','','','','','','','Amparo','Amparo'], 'Autônoma','',3,1,2,2,True,2,{}),
    ('família ausente', ['Alcance']+['']*8, 'Autônoma','',3,1,1,0,False,2,{}),
    ('tudo vazio', ['']*9, 'Autônoma','',3,1,0,0,False,2,{}),
    ('não autônoma padrão: apenas uma Livre nova', ['Alcance','Amparo']+['']*7, 'Não autônoma','',4,2,2,0,True,5,{}),
    ('autônoma Principal: sexta comum', ['Alcance','Área','','','','Amparo','','',''], 'Autônoma','Invocação Principal',5,2,3,0,True,2,{}),
    ('não autônoma Principal: sexta Livre', ['Alcance','Área','','','','Amparo','','',''], 'Não autônoma','Invocação Principal',6,3,3,0,True,5,{}),
    ('não autônoma Parceria', ['Alcance','Amparo']+['']*7, 'Não autônoma','Parceria',5,2,2,0,True,5,{}),
    ('não autônoma Múltiplas', ['Alcance','Amparo']+['']*7, 'Não autônoma','Múltiplas Invocações',5,2,2,0,True,5,{}),
    ('campanha permite nove, nona Livre', ['Alcance','','','','','','','','Amparo'], 'Autônoma','',9,9,2,0,True,5,{139:9,140:9}),
    ('campanha altera ganho adicional para dois', ['Alcance','Área','','','','Amparo','','',''], 'Não autônoma','',5,3,3,0,True,5,{145:2}),
    ('campo apagado usa Autônoma', ['Amparo']+['']*8, '', '',3,1,1,0,True,5,{}),
]
with tempfile.TemporaryDirectory(prefix='b47-fam7-') as tmp:
    tmp = Path(tmp)
    out = tmp/'out'
    out.mkdir()
    inputs = []
    for i, (_, values, mode, trail, *_expected, overrides) in enumerate(scenarios):
        wb = Workbook()
        wb.remove(wb.active)
        ws = wb.create_sheet('DADOS_INVOC')
        inv = wb.create_sheet('INVOCAÇÕES')
        cfg = wb.create_sheet('CONFIGURAÇÃO')
        for coord, value in raw['DADOS_INVOC'].items():
            if value is not None and not (isinstance(value, str) and value.startswith('=')):
                ws[coord] = value
        for j in range(5):
            ws[col(core+j)+'2'] = raw['DADOS_INVOC'][col(core+j)+'2']
        for coord in tuple(col(c)+'2' for c in range(464,479))+('QZ3',)+tuple(targets):
            ws[coord] = raw['DADOS_INVOC'][coord]
        ws['DB4'] = trail
        for row in range(139,146):
            cfg['E'+str(row)] = raw['CONFIGURAÇÃO']['E'+str(row)]
            cfg['G'+str(row)] = raw['CONFIGURAÇÃO']['G'+str(row)]
        for row,value in overrides.items():
            cfg['F'+str(row)] = value
        inv['K50'] = mode
        inv['K248'] = 'Autônoma'
        ws[col(fields['ficha'])+'2'] = 1
        ws[col(fields['lf'])+'2'] = cura-1
        ws[col(fields['w1'])+'2'] = 2
        ws[col(fields['f1'])+'2'] = 'Amparo'
        ws[col(fields['classe'])+'2'] = 1
        for coord, value in zip(familias, values):
            inv[coord] = value
        path = tmp/f'caso-{i}.xlsx'
        wb.save(path)
        inputs.append(path)
    subprocess.run([shutil.which('libreoffice') or 'soffice', '-env:UserInstallation='+ (tmp/'profile').as_uri(),
                    '--headless', '--convert-to', 'xlsx', '--outdir', str(out), *map(str, inputs)],
                   check=True, capture_output=True, timeout=120)
    for i, (name, _, mode, trail, maximum, livres, count, duplicates, opened, price, overrides) in enumerate(scenarios):
        ws = load_workbook(out/f'caso-{i}.xlsx', data_only=True)['DADOS_INVOC']
        values = [ws[c].value for c in targets]
        expected = [count, duplicates, 'Cura' if opened else None,
                    raw['DADOS_INVOC'][col(peca)+str(amparo)] if opened else None,
                    0 if opened else 1, 0 if opened else 1, maximum, livres, price,
                    maximum-(overrides.get(145,1) if mode=='Não autônoma' else 0)]
        # O recalculador pode representar a string vazia como None.
        normal = lambda v: None if v == '' else v
        check(name+': contagem, duplicatas, menus, desconto, limites e segunda entidade independente',
              list(map(normal, values)) == expected)
        if list(map(normal, values)) != expected:
            print('obtido', values, 'esperado', expected)
if not all(checks):
    raise SystemExit(1)
