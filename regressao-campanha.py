#!/usr/bin/env python3
"""Recalcula a base B47 e a nova ficha em processos seriais, com contas independentes."""
import hashlib,json,math,re,shutil,subprocess,sys,tempfile
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parent
sys.path.insert(0,str(ROOT/'capacidade-b47'))
from contratos_historicos import transformar,classificar_vazio
FALHAS=[]
def check(nome,cond,det=''):
 print(('OK ' if cond else 'FALHA ')+nome+(' '+str(det) if not cond else ''),flush=True)
 if not cond:FALHAS.append(nome)
def vals(w,s):return {c.coordinate:c.value for row in w[s].iter_rows() for c in row if c.value is not None}
def formula(v):return re.sub(r'(?<![A-Z_.])IFS\(','_xlfn.IFS(',v) if isinstance(v,str) and v.startswith('=') else v
base=ROOT/'nuvem-b47/fixtures/base-parte1.xlsx'
novo=ROOT/'ficha-v01/ficha-projeto-m-0.1.xlsx'
if not base.exists() or not (shutil.which('libreoffice') or shutil.which('soffice')):sys.exit('Falta base parte 1 ou LibreOffice; teste não rodou')
wb=load_workbook(novo)
sys.path.insert(0,str(ROOT/'ficha-v01'));import indice_ficha as ix
idx={wb['DADOS'].cell(r,53).value:ix.endereco(wb['DADOS'].cell(r,54).value) for r in range(5,wb['DADOS'].max_row+1) if wb['DADOS'].cell(r,53).value}
linhas=[[c.value for c in row] for row in wb['ANALISE'].iter_rows(min_row=2,max_col=4) if row[0].value]
oldkeys=[]
for r in linhas[:5156]:
 if r[0].startswith(('regra.','vida.rolagem.','aviso.regras_campanha','aviso.vida_rolada','sequela.extras','custom.','aviso.custom.')):continue
 oldkeys.append(transformar(r,versao=3,reverter=True))
snap=json.loads((ROOT/'ficha-v01/analise-v3-preservada.json').read_text())
check('5.156 chaves preservadas, com as mudanças exatas aprovadas no formato 7',len(oldkeys)==snap['linhas'] and hashlib.sha256(json.dumps(oldkeys,ensure_ascii=False,separators=(',',':')).encode()).hexdigest()==snap['sha256'])
check('ANALISE formato 11 (C)',linhas[0][1]=='11')
check('morte permanente não é exportada',not any('morte_permanente' in r[0] for r in linhas))
check('chaves únicas',len(linhas)==len({r[0] for r in linhas}))
wb.close()
if '--estrutura' in sys.argv:sys.exit(bool(FALHAS))
with tempfile.TemporaryDirectory(prefix='b47-campanha-') as tmp:
 tmp=Path(tmp);inputdir=tmp/'in';out=tmp/'out';inputdir.mkdir();out.mkdir()
 cache={}
 def preparar(src,nome,edit=None,nivel=10):
  if src not in cache:
   cache[src]=load_workbook(src)
   for ss in cache[src]:
    ss._images=[]
    for cc in tuple(ss._cells.values()):
     if isinstance(cc.value,str) and cc.value.startswith('='):cc.value=formula(cc.value)
  w=cache[src];f=w['FICHA'];f[idx['caminho']]='Vanguarda';f[idx['nivel']]=nivel
  for k,v in [('atr_base_Força',3),('atr_base_Destreza',2),('atr_base_Constituição',3),('atr_base_Inteligência',1),('atr_base_Essência',2),('buff de vida',5),('buff de deslocamento',.5)]:f[idx[k]]=v
  for s in w:
   s._images=[]
   for c in tuple(s._cells.values()):
    if isinstance(c.value,str) and c.value.startswith('='):c.value=formula(c.value)
  restore=edit(w) if edit else []
  p=inputdir/(nome+'.xlsx');w.save(p)
  for ss,cc,vv in restore:w[ss][cc]=vv
  return p
 def edits(d):
  def aplica(w):
   antes=[(s,c,w[s][c].value) for s,c in d]
   for (s,c),v in d.items():w[s][c]=v
   return antes
  return aplica
 casos=[('anterior',base,{},10),('vazia',novo,{},10),
  ('atributos',novo,{('CONFIGURAÇÃO','E47'):12,('CONFIGURAÇÃO','E48'):4,('CONFIGURAÇÃO','E49'):7},10),
  ('treinos',novo,{('CONFIGURAÇÃO','E50'):11,('CONFIGURAÇÃO','E51'):3,('CONFIGURAÇÃO','E52'):13},10),
  ('deslocamento',novo,{('CONFIGURAÇÃO','E53'):6},10),
  ('morrendo',novo,{('CONFIGURAÇÃO','E54'):.5,('CONFIGURAÇÃO','E55'):.25,('CONFIGURAÇÃO','E56'):.75,('CONFIGURAÇÃO','E57'):2,('CONFIGURAÇÃO','E58'):5,('FICHA',idx['sequela 1']):True,('FICHA',idx['sequela extras']):2},10),
  ('xp',novo,{('CONFIGURAÇÃO','E46'):5,('CONFIGURAÇÃO','E69'):700},10),
  ('xp_zero',novo,{('CONFIGURAÇÃO','E46'):5,('CONFIGURAÇÃO','E69'):0},10),
  ('xp_um_pendente',novo,{('CONFIGURAÇÃO','E46'):1},10),
  ('rolada',novo,{('CONFIGURAÇÃO','E59'):'Sim',('CONFIGURAÇÃO','E101'):5,('CONFIGURAÇÃO','E102'):4,('CONFIGURAÇÃO','E103'):6},4),
  ('rolada_falta',novo,{('CONFIGURAÇÃO','E59'):'Sim'},4),
  ('invalida',novo,{('CONFIGURAÇÃO','E47'):'doze'},10),
  ('troca_zero',novo,{('CONFIGURAÇÃO','E59'):'Sim',('CONFIGURAÇÃO','B11'):True,('CONFIGURAÇÃO','C11'):'Vida máxima',('CONFIGURAÇÃO','D11'):'trocar por',('CONFIGURAÇÃO','E11'):0,('CONFIGURAÇÃO','F11'):'Decisão do mestre',('CONFIGURAÇÃO','G11'):'Teste'},4)]
 paths=[preparar(src,nome,edits(d),nv) for nome,src,d,nv in casos]
 for ww in cache.values():ww.close()
 cache.clear()
 perfil=tmp/'perfil'
 for p in paths:
  r=subprocess.run([shutil.which('libreoffice') or 'soffice',f'-env:UserInstallation={perfil.as_uri()}','--headless','--convert-to','xlsx','--outdir',str(out),str(p)],capture_output=True,text=True,timeout=600)
  check('recalcular '+p.stem,r.returncode==0 and (out/p.name).exists(),r.stderr[-200:])
 if any(not (out/p.name).exists() for p in paths):sys.exit(1)
 dados={}
 for nome,*_ in casos:
  w=load_workbook(out/(nome+'.xlsx'),read_only=True,data_only=True)
  dados[nome]={s:vals(w,s) for s in w.sheetnames if s not in ['CONFIGURAÇÃO','DADOS_CONFIG','ANALISE']}
  dados[nome]['DADOS_CONFIG']=vals(w,'DADOS_CONFIG') if 'DADOS_CONFIG' in w.sheetnames else {}
  w.close()
 def get(nome,s,k):return dados[nome][s].get(k)
 # Todas as células pré-existentes, não apenas uma seleção de resultados.
 dif=[]
 for s,cs in dados['anterior'].items():
  if s=='DADOS_CONFIG':continue
  for c,v in cs.items():
   atual=get('vazia',s,c)
   bate=math.isclose(v,atual,rel_tol=1e-10,abs_tol=1e-9) if isinstance(v,(int,float)) and isinstance(atual,(int,float)) else v==atual
   if not bate:dif.append((s,c,v,atual))
 print('DIFERENÇAS COMPLETAS CAMPANHA VAZIA: '+json.dumps(dif,ensure_ascii=False,default=str),flush=True)
 aprovadas,pendentes=classificar_vazio(dif,lambda s,c:get('vazia',s,c))
 print('CLASSIFICAÇÃO CAMPANHA VAZIA: '+json.dumps({'aprovadas':aprovadas,'pendentes':pendentes},ensure_ascii=False,default=str),flush=True)
 check('CAMPANHA vazia: valores preservados, com rótulos de C e cabeçalhos conferidos no novo endereço',not pendentes,pendentes)
 v=lambda caso,k:get(caso,'FICHA',idx[k])
 check('pontos iniciais somam sem apagar marcos',get('atributos','DADOS','AK6')==get('vazia','DADOS','AK6')+3)
 check('treinos com ofícios e alternativa mudam só seus totais',get('treinos','DADOS','AK19')==11 and get('treinos','DADOS','AK20')==3 and get('treinos','DADOS','AK27')==13)
 check('deslocamento base mantém Buff',str(v('deslocamento','deslocamento')).replace(',','.')=='6.5 m',v('deslocamento','deslocamento'))  # Calc em pt-BR escreve 6,5
 vida=v('morrendo','vida_max');cl=v('morrendo','classe máxima')
 check('Aguentar usa fração e arredonda para cima',v('morrendo','morrendo aguentar')==math.ceil(vida*.5))
 check('Insistir mantém teto configurado',v('morrendo','morrendo insistir 1')==min(math.ceil(vida*.25),2*cl) and v('morrendo','morrendo insistir 2')==min(math.ceil(vida*.75),2*cl))
 check('prazo soma três caixas e Sequelas extras',v('morrendo','morrendo prazo')=='2 rodadas')
 check('XP zero no nível inicial e custo alternativo',get('xp','DADOS','AR8')==0 and get('xp','DADOS','AR9')==700)
 check('custo XP zero não é tratado como vazio',get('xp_zero','DADOS','AR9')==0)
 check('nível 1 sem custo de avanço fica pendente',get('xp_um_pendente','DADOS_CONFIG','Z3')=='Pendente' and bool(get('xp_um_pendente','DADOS_CONFIG','Z33')))
 # Vida inicial da tabela do livro, não a fórmula nova que está sendo julgada.
 wb0=load_workbook(base,read_only=True,data_only=False);pv=None
 for r in wb0['DADOS'].iter_rows(min_row=5,max_row=10,min_col=14,max_col=18):
  if r[0].value=='Vanguarda':pv=r[2].value
 wb0.close();esperada=pv+5+4+6+3*4+5
 check('dados guardados + vida inicial + CON por nível + Buff',v('rolada','vida_max')==esperada,(v('rolada','vida_max'),esperada))
 check('rolagem faltando não inventa vida',v('rolada_falta','vida_max')=='#N/A' and bool(get('rolada_falta','DADOS_CONFIG','AF32')))
 check('regra inválida fica visível',get('invalida','DADOS','AK6')=='#N/A')
 check('trocar por zero domina a nova conta',v('troca_zero','vida_max')==0)
 # Regras sem relação não podem alterar recursos, ataque ou Defesa.
 for caso in ['atributos','treinos','xp','xp_zero','xp_um_pendente']:
  check(caso+': recursos e combate preservados',all(v(caso,k)==v('vazia',k) for k in ['vida_max','energia_max','defesa','proteção','corpo a corpo','à distância']))
print('CAMPANHA CONFERE' if not FALHAS else f'{len(FALHAS)} FALHAS')
sys.exit(bool(FALHAS))
