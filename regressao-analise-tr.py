"""Regressão da ampliação de leitura: contrato, XLSX e Apps Script devem concordar.

O manifesto registra a base antes da ampliação. Uma alteração intencional futura
no desenho exige revisar esse manifesto; não se atualizam hashes só para passar.
"""
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile
import shutil

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent


class AnaliseTR(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.base = json.loads((RAIZ / "regressao-analise-tr-base.json").read_text())
        with tempfile.TemporaryDirectory() as tmp:
            destino = Path(tmp) / "analise.json"
            subprocess.run([sys.executable, str(RAIZ / "ficha-v01/monta.py"), "--exportar-analise", str(destino)],
                           cwd=RAIZ, check=True, stdout=subprocess.DEVNULL)
            cls.linhas = json.loads(destino.read_text())["linhas"]

    def test_chaves_antigas_preservadas_com_mesma_formula_e_origem(self):
        antigas = [list(l) for l in self.linhas[:self.base["linhas_antigas"]]]
        self.assertEqual(antigas[0][1], "4")
        antigas[0][1] = "2"  # única alteração autorizada no prefixo antigo
        h = hashlib.sha256(json.dumps(antigas, ensure_ascii=False, separators=(",", ":")).encode()).hexdigest()
        self.assertEqual(h, self.base["contrato_v2_sha256"])
        self.assertEqual(len(self.linhas), 5381)
        self.assertEqual(len({l[0] for l in self.linhas}), 5381)

    def test_entradas_e_resultados_sao_distintos(self):
        linhas = {k: (v, c, p) for k, v, c, p in self.linhas}
        for k in ("tr.fisico.ajuste", "pericia.atletismo.atributo", "item.1.quantidade", "carga.ajuste_limite"):
            self.assertEqual(linhas[k][2], "entrada", k)
        for k in ("tr.fisico.bonus", "pericia.atletismo.bonus", "invocacao.1.pericia.1.bonus", "carga.total"):
            self.assertEqual(linhas[k][2], "resultado", k)
        self.assertEqual(linhas["invocacao.1.buff.pericia.1"][1], linhas["invocacao.1.pericia.1.bonus"][1])
        self.assertNotIn("configuracao.versao", linhas)

    def test_xlsx_mantem_as_outras_nove_abas_estilos_e_imagens(self):
        with zipfile.ZipFile(RAIZ / "ficha-v01/ficha-projeto-m-0.1.xlsx") as z:
            for nome, h in self.base["xlsx_preservado_sha256"].items():
                self.assertEqual(hashlib.sha256(z.read(nome)).hexdigest(), h, nome)

    def test_xlsx_e_apps_script_correspondem_ao_exportador(self):
        w = load_workbook(RAIZ / "ficha-v01/ficha-projeto-m-0.1.xlsx", read_only=True, data_only=False)
        try:
            self.assertEqual(w["ANALISE"].sheet_state, "hidden")
            linhas = [tuple(l) for l in w["ANALISE"].iter_rows(min_row=2, max_col=4, values_only=True) if l[0]]
        finally:
            w.close()
        gs = (RAIZ / "apps-script/Analise.gs").read_text()
        # A única declaração é JSON gerado; não executa o Apps Script.
        objetos = json.JSONDecoder().raw_decode(gs[gs.index("[", gs.index("var ABAS_DA_ANALISE")):])[0]
        mapa = {(l, c): v for l, c, v, _ in objetos[0]["vals"]}
        extra = (RAIZ / "apps-script/AnaliseTestes.gs").read_text()
        objetos_extra = json.JSONDecoder().raw_decode(extra[extra.index("[", extra.index("var VALORES_DOS_TESTES_")):])[0]
        mapa.update({(l, c): v for l, c, v, _ in objetos_extra[0]["vals"]})
        self.assertEqual(len(linhas), len(self.linhas))
        for i, (k, v, c, p) in enumerate(linhas, 2):
            gs_v = mapa[(i, 2)]
            # Sheets recebe apóstrofo explícito para conservar número como texto;
            # o XLSX grava esse texto diretamente.
            if isinstance(gs_v, str) and gs_v.startswith("'"):
                gs_v = gs_v[1:]
            self.assertEqual((mapa[(i, 1)], gs_v, mapa[(i, 3)], mapa[(i, 4)]), (k, v, c or "", p))
            self.assertEqual((k, p), (self.linhas[i - 2][0], self.linhas[i - 2][3]))

    def test_juncao_em_todas_as_ordens_idempotente_e_falta_detectada(self):
        self.assertIsNotNone(shutil.which("node"), "Node é necessário para validar o transporte Apps Script")
        codigo = r"""
const fs=require('fs'),vm=require('vm'),assert=require('assert');
const nomes=['Ficha.gs','Analise.gs','AnaliseTestes.gs'];
function perm(a){return a.length ? a.flatMap((x,i)=>perm(a.filter((_,j)=>i!==j)).map(b=>[x,...b])):[[]];}
const resultados=[];
for(const ordem of perm(nomes)) {
 const c=vm.createContext({});
 for(const n of ['Origens.gs',...ordem,...ordem]) vm.runInContext(fs.readFileSync('apps-script/'+n,'utf8'),c);
 const vals=vm.runInContext('ABAS.filter(s=>s.nome==="ANALISE")[0].vals',c);
 assert.equal(vals.length,5382*4);
 assert.equal(new Set(vals.map(v=>v[0]+','+v[1])).size,vals.length);
 resultados.push(JSON.stringify(vals));
}
assert(resultados.every(x=>x===resultados[0]));
const c=vm.createContext({});
for(const n of ['Origens.gs','Ficha.gs','Analise.gs']) vm.runInContext(fs.readFileSync('apps-script/'+n,'utf8'),c);
assert.throws(()=>vm.runInContext('montar_()',c),/Falta o AnaliseTestes.gs/);
console.log('6 ordens, repetição e falta de complemento: OK');
"""
        p = subprocess.run(["node", "-e", codigo], cwd=RAIZ, capture_output=True, text=True)
        self.assertEqual(p.returncode, 0, p.stdout + p.stderr)


if __name__ == "__main__":
    unittest.main()
