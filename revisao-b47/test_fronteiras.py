# -*- coding: utf-8 -*-
"""Primeira e última vaga dos catálogos de Família e de treino, recalculadas no LibreOffice.

As células de cada registro vêm da própria ANALISE do modelo (chaves custom.<tipo>.N.<campo> e
aviso.custom.<tipo>.N); as tabelas ocultas DADOS_CRIACAO/DADOS_PROGRESSAO entram inteiras com
suas fórmulas e a CONFIGURAÇÃO só com valores. Cada caso é um arquivo; um único recálculo.
Calculadora local: nunca importar o XLSX no Google.

    python revisao-b47/test_fronteiras.py --perfil simulado-alvo
"""
import json
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formula import Tokenizer

import comum

ap = comum.argumentos(__doc__)
args = ap.parse_args()
p = comum.perfil(args.perfil)
P = comum.Placar('fronteiras ' + args.perfil)
modelo = json.loads(comum.exportar(args.perfil, 'pecas').read_text(encoding='utf-8'))
abas = {a['nome']: a for a in modelo['layout']['abas']}
A = {l[0]: l for l in comum._linhas_da_aba(abas['ANALISE'])}
print(f'SHA {comum.sha()} · perfil {args.perfil}\n')


def celula(formula):
    for t in Tokenizer(formula).items:
        if t.type == 'OPERAND' and t.subtype == 'RANGE':
            dono, end = t.value.rsplit('!', 1) if '!' in t.value else ('ANALISE', t.value)
            return dono.strip("'"), end.replace('$', '')


def campo(cat, n, nome):
    return celula(A[f'custom.{cat}.{n}.{nome}'][1])[1]          # sempre na CONFIGURAÇÃO


def aviso(cat, n):
    return celula(A[f'aviso.custom.{cat}.{n}'][1])              # (DADOS_*, célula do problema)


TABELA = {'familia': 'DADOS_CRIACAO', 'treino': 'DADOS_PROGRESSAO', 'pericia': 'DADOS_PROGRESSAO',
          'oficio': 'DADOS_PROGRESSAO'}
PADRAO = {'familia': {'estado': 'Neutra'}, 'treino': {'tipo': 'Perícia', 'atributo': 'Força'},
          'pericia': {'tipo': 'Perícia', 'atributo': 'Força'}, 'oficio': {'tipo': 'Ofício', 'atributo': 'Inteligência'}}
ult = {c: p['catalogos'][c]['capacidade'] for c in TABELA if c in p['catalogos']}
casos = []   # (nome, cat, {(n, campo): valor}, verificação)
for cat, N in ult.items():
    extra = PADRAO[cat]
    def reg(n, nome, **mais):
        return {**{(n, k): v for k, v in extra.items()}, (n, 'nome'): nome, **{(n, k): v for k, v in mais.items()}}
    casos += [
        (f'{cat}: primeira e última vaga válidas e no catálogo', cat, {**reg(1, 'Primeira própria'), **reg(N, 'Última própria')},
         dict(validos=[1, N], no_catalogo=['Primeira própria', 'Última própria'])),
        (f'{cat}: última repete a primeira sem diferenciar maiúsculas', cat, {**reg(1, 'Eco Duplo'), **reg(N, 'eco duplo')},
         dict(invalidos=[1, N])),
        (f'{cat}: curinga na última é literal (Lua* não colide com Lua X)', cat, {**reg(1, 'Lua X'), **reg(N, 'Lua*')},
         dict(validos=[1, N], no_catalogo=['Lua X', 'Lua*'])),
    ]
    if cat in ('pericia', 'oficio'):
        outro = 'Ofício' if cat == 'pericia' else 'Perícia'
        casos.append((f'{cat}: TIPO apagado num cartão vazio não acende aviso', cat, {(N, 'tipo'): None},
                      dict(validos=[N])))
        casos.append((f'{cat}: TIPO de outro catálogo colado na última vaga é recusado', cat,
                      {**reg(N, 'Tipo trocado'), (N, 'tipo'): outro}, dict(invalidos=[N])))
    if cat == 'treino':
        casos.append((f'{cat}: última vaga com nome e sem TIPO é inválida', cat,
                      {(N, 'nome'): 'Sem tipo', (N, 'atributo'): 'Força'}, dict(invalidos=[N])))
        casos.append((f'{cat}: Ofício na última vaga entra no catálogo de ofícios', cat,
                      {(N, 'tipo'): 'Ofício', (N, 'nome'): 'Ofício próprio', (N, 'atributo'): 'Inteligência'},
                      dict(validos=[N], no_catalogo=['Ofício próprio'])))

if 'pericia' in ult and 'oficio' in ult:
    casos.append(('pericia+oficio: mesmo nome nos dois tipos continua inválido (como no catálogo único)', 'pericia',
                  {(1, 'nome'): 'Ofuscar', (1, 'tipo'): 'Perícia', (1, 'atributo'): 'Força'},
                  dict(invalidos=[1], extra_oficio={(1, 'nome'): 'Ofuscar', (1, 'tipo'): 'Ofício', (1, 'atributo'): 'Força'})))
vazias = {s: {} for s in abas}
tmp = Path(tempfile.mkdtemp(prefix='revisao-b47-fronteiras-'))
try:
    out = tmp / 'out'
    out.mkdir()
    caminhos = []
    for i, (nome, cat, entradas, verif) in enumerate(casos):
        wb = Workbook()
        wb.remove(wb.active)
        for s in abas:
            wb.create_sheet(s)          # abas citadas existem, mesmo vazias
        for c, v, *_ in abas['CONFIGURAÇÃO']['celulas']:
            if v is not None and not (isinstance(v, str) and v.startswith('=')):
                wb['CONFIGURAÇÃO'][c] = v
        for s in set(TABELA.values()):
            for c, v, *_ in abas[s]['celulas']:
                if v is None:
                    continue
                if isinstance(v, str) and v.startswith('='):
                    v = re.sub(r'(?<![A-Z_.])(TEXTJOIN|IFS|SWITCH)\(', r'_xlfn.\1(', v)
                wb[s][c] = v
        for (n, k), v in entradas.items():
            wb['CONFIGURAÇÃO'][campo(cat, n, k)] = v
        for (n, k), v in verif.get('extra_oficio', {}).items():
            wb['CONFIGURAÇÃO'][campo('oficio', n, k)] = v
        caminho = tmp / f'caso{i:02d}.xlsx'
        wb.save(caminho)
        caminhos.append(caminho)
    r = subprocess.run([shutil.which('libreoffice') or 'soffice', '-env:UserInstallation=' + (tmp / 'perfil').as_uri(),
                        '--headless', '--convert-to', 'xlsx', '--outdir', str(out), *map(str, caminhos)],
                       capture_output=True, text=True, timeout=1200)
    P.check('LibreOffice recalculou todos os casos', all((out / c.name).exists() for c in caminhos), r.stderr[-300:])
    for (nome, cat, entradas, esperado), caminho in zip(casos, caminhos):
        wb = load_workbook(out / caminho.name, data_only=True, read_only=True)
        s = TABELA[cat]
        valores = {c.coordinate: c.value for row in wb[s] for c in row if c.value is not None}
        wb.close()
        prob = {n: valores.get(aviso(cat, n)[1], '') or '' for n in esperado.get('validos', []) + esperado.get('invalidos', [])}
        ok = all(prob[n] == '' for n in esperado.get('validos', [])) and all(prob[n] != '' for n in esperado.get('invalidos', []))
        contagem = {x: sum(1 for v in valores.values() if v == x) for x in esperado.get('no_catalogo', [])}
        ok_cat = all(q >= 2 for q in contagem.values())        # cadastro cru + catálogo
        nivel = 'risco' if ('curinga' in nome and cat == 'treino') or 'TIPO apagado' in nome else 'bloqueante'
        P.check(nome, ok and ok_cat, f'avisos {prob}; ocorrências {contagem}', nivel=nivel)
finally:
    shutil.rmtree(tmp, ignore_errors=True)
P.fim(dict(perfil=args.perfil, ultimas=ult))
