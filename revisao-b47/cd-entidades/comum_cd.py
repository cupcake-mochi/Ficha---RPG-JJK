# -*- coding: utf-8 -*-
"""Utilidades da frente de entidades (C+D). Só lê o núcleo; saídas em nuvem-b47/saidas/cd-entidades/."""
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parent.parent
SAIDAS = RAIZ / 'nuvem-b47' / 'saidas' / 'cd-entidades'
sys.path[:0] = [str(RAIZ / 'ficha-v01'), str(RAIZ / 'ficha'), str(RAIZ), str(AQUI)]


def python():
    venv = RAIZ / '.venv' / 'bin' / 'python'
    return str(venv) if venv.exists() else sys.executable


def sha():
    return subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=RAIZ, text=True).strip()


def nucleo_limpo():
    return not subprocess.check_output(['git', 'status', '--porcelain', '--', '.', ':!revisao-b47'], cwd=RAIZ,
                                       text=True).strip()


def _marca():
    fonte = (AQUI / 'entidades_custom.py').read_bytes() + (AQUI / 'rodar_com_entidades.py').read_bytes()
    import hashlib
    return sha()[:12] + '-' + hashlib.sha256(fonte).hexdigest()[:8] + ('' if nucleo_limpo() else '-sujo')


def exportar(tipos=None, talentos=None, raiz=RAIZ):
    """Modelo completo (--exportar-pecas). tipos=None: gerador sem o módulo (base de comparação)."""
    SAIDAS.mkdir(parents=True, exist_ok=True)
    marca = _marca()
    nome = 'base' if tipos is None else f'ent-{tipos}-{talentos}'
    modelo, tr = SAIDAS / f'{nome}-{marca}.json', SAIDAS / f'{nome}-{marca}-tr.json'
    if modelo.exists() and not marca.endswith('-sujo'):
        return modelo, (tr if tr.exists() else None)
    if tipos is None:
        cmd = [python(), str(RAIZ / 'revisao-b47' / 'rodar_gerador.py'), str(raiz), '{}', '--exportar-pecas', str(modelo)]
    else:
        cmd = [python(), str(AQUI / 'rodar_com_entidades.py'), str(raiz), str(tr), str(tipos), str(talentos),
               '--exportar-pecas', str(modelo)]
    r = subprocess.run(cmd, cwd=raiz, capture_output=True, text=True)
    if r.returncode or not modelo.exists():
        raise SystemExit((r.stdout + r.stderr)[-3000:])
    return modelo, (tr if tr.exists() else None)


def carregar(caminho):
    return json.loads(Path(caminho).read_text(encoding='utf-8'))


def celulas(modelo):
    return {a['nome']: {v[0]: v[1] for v in a['celulas']} for a in modelo['layout']['abas']}


def recalcular(modelo, cenarios, leitura, prefixo='cd'):
    """Recalcula no LibreOffice um XLSX por cenário e devolve {cenário: {(aba, a1): valor}}.

    cenarios: {nome: {(aba, a1): valor}}; leitura: [(aba, a1)]. Todas as abas do modelo entram com
    fórmulas e valores; o XLSX é calculadora local, nunca vai ao Google.
    """
    from openpyxl import Workbook, load_workbook
    calc = shutil.which('libreoffice') or shutil.which('soffice')
    if not calc:
        raise RuntimeError('IMPEDIDO: LibreOffice Calc ausente')
    base = celulas(modelo)
    tmp = Path(tempfile.mkdtemp(prefix=f'{prefixo}-'))
    try:
        out = tmp / 'out'
        out.mkdir()
        nomes = []
        for i, (nome, entradas) in enumerate(cenarios.items()):
            wb = Workbook()
            wb.remove(wb.active)
            for aba, cels in base.items():
                ws = wb.create_sheet(aba)
                for a1, v in cels.items():
                    if v is None:
                        continue
                    if isinstance(v, str) and v.startswith('='):
                        v = re.sub(r'(?<![A-Z_.])(IFS|SWITCH|TEXTJOIN|IFNA|XLOOKUP|CONCAT)\(', r'_xlfn.\1(', v)
                    ws[a1] = v
            for (aba, a1), v in entradas.items():
                wb[aba][a1] = v
            p = tmp / f'c{i:02d}.xlsx'
            wb.save(p)
            nomes.append((nome, p))
        r = subprocess.run([calc, '-env:UserInstallation=' + (tmp / 'perfil').as_uri(), '--headless', '--convert-to',
                            'xlsx', '--outdir', str(out), *[str(p) for _, p in nomes]],
                           capture_output=True, text=True, timeout=3600)
        res = {}
        for nome, p in nomes:
            f = out / p.name
            if not f.exists():
                raise RuntimeError(f'LibreOffice não converteu {nome}: {r.stderr[-500:]}')
            wb = load_workbook(f, data_only=True, read_only=False)
            res[nome] = {(aba, a1): wb[aba][a1].value for aba, a1 in leitura}
            wb.close()
        return res
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


class Placar:
    def __init__(self, titulo):
        self.titulo, self.itens = titulo, []

    def check(self, nome, ok, detalhe='', nivel='bloqueante'):
        self.itens.append(dict(nome=nome, ok=bool(ok), detalhe=str(detalhe)[:800], nivel=nivel))
        marca = 'OK ' if ok else ('FALHA' if nivel == 'bloqueante' else 'RISCO')
        print(f'{marca:5} {nome}' + ('' if ok or not detalhe else f' — {str(detalhe)[:500]}'), flush=True)

    def impedido(self, nome, motivo):
        self.itens.append(dict(nome=nome, ok=False, detalhe=motivo, nivel='impedido'))
        print(f'IMPED {nome} — {motivo}', flush=True)

    def fim(self, extra=None):
        falhas = [i for i in self.itens if not i['ok'] and i['nivel'] == 'bloqueante']
        imp = [i for i in self.itens if i['nivel'] == 'impedido']
        riscos = [i for i in self.itens if not i['ok'] and i['nivel'] == 'risco']
        SAIDAS.mkdir(parents=True, exist_ok=True)
        nome = re.sub(r'[^a-z0-9]+', '-', self.titulo.lower()).strip('-')
        (SAIDAS / f'{nome}.json').write_text(json.dumps(dict(teste=self.titulo, sha=sha(), nucleo_limpo=nucleo_limpo(),
            itens=self.itens, **(extra or {})), ensure_ascii=False, indent=1), encoding='utf-8')
        ok = len(self.itens) - len(falhas) - len(imp) - len(riscos)
        print(f'\n{self.titulo}: {ok} ok · {len(falhas)} falha(s) · {len(riscos)} risco(s) · {len(imp)} impedido(s)')
        sys.exit(1 if falhas else (2 if imp and not ok else 0))
