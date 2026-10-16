# -*- coding: utf-8 -*-
"""B3: recalcula o XLSX versionado no LibreOffice e mostra quais componentes do aviso 3a viram #NAME?.

Uso (~2,5 min, pesado; rodar sozinho): python revisao-b47/sonda_libreoffice_name.py
"""
import os,re,shutil,subprocess,tempfile,sys
from openpyxl import load_workbook
src='ficha-v01/ficha-projeto-m-0.1.xlsx'
wb=load_workbook(src)
ws=wb['DADOS_CUSTOM']
for c in ['S2','U2','V2','W2','X2']:
    f=str(ws[c].value); print(c, sorted(set(re.findall(r'([A-Z_.][A-Z0-9_.]*)\(',f))))
pasta=tempfile.mkdtemp(); copia=os.path.join(pasta,'f.xlsx')
for w in wb:
    for cell in tuple(w._cells.values()):
        if isinstance(cell.value,str) and cell.value.startswith('='):
            cell.value=re.sub(r"(?<![A-Z_.])(IFS|SWITCH)\(", r"_xlfn.\1(", cell.value)
wb.save(copia)
subprocess.run(['libreoffice','--headless','--convert-to','xlsx','--outdir',pasta+'/o',copia],capture_output=True,timeout=900)
r=load_workbook(pasta+'/o/f.xlsx',data_only=True)['DADOS_CUSTOM']
for c in ['S2','T2','U2','V2','W2','X2','S3','U3','V3','W3','X3']: print(c,repr(r[c].value))
shutil.rmtree(pasta,ignore_errors=True)
